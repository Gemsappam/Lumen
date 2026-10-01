"""Excel invoices (.xlsx / .xls) -> plain text table for the AI (models don't read spreadsheets directly)."""
import io
import re


def _fmt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def excel_to_text(data: bytes, filename: str = "") -> str:
    rows_by_sheet = []
    name = (filename or "").lower()
    try:
        if name.endswith(".xls") and not data[:4] == b"PK\x03\x04":
            import xlrd
            book = xlrd.open_workbook(file_contents=data)
            for sh in book.sheets():
                rows = [[_fmt(sh.cell_value(r, c)) for c in range(sh.ncols)] for r in range(sh.nrows)]
                rows_by_sheet.append((sh.name, rows))
        else:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
            for ws in wb.worksheets:
                rows = [[_fmt(v) for v in r] for r in ws.iter_rows(values_only=True)]
                rows_by_sheet.append((ws.title, rows))
    except Exception:
        # some «.xls» are really HTML tables
        txt = data.decode("utf-8", errors="ignore")
        txt = re.sub(r"<\s*(br|/tr|/p)\s*/?>", "\n", txt, flags=re.I)
        txt = re.sub(r"<\s*/t[dh]\s*>", " | ", txt, flags=re.I)
        return re.sub(r"<[^>]+>", "", txt)
    out = []
    for title, rows in rows_by_sheet:
        rows = [r for r in rows if any(x for x in r)]
        if not rows:
            continue
        out.append(f"=== Лист: {title} ===")
        for r in rows[:400]:
            while r and not r[-1]:
                r = r[:-1]
            out.append(" | ".join(r))
    return "\n".join(out)


def looks_like_master(data: bytes) -> bool:
    """Our учёт file: has «Пополнения» or «Пополнение dd.mm» sheets."""
    try:
        from openpyxl import load_workbook
        names = load_workbook(io.BytesIO(data), read_only=True).sheetnames
    except Exception:
        return False
    return "Пополнения" in names or any(n.startswith("Пополнение ") or n.startswith("Поставка ") for n in names)
