import json
import re
import asyncio
import shutil
import time
from datetime import datetime, timedelta

from aiogram import Bot, Dispatcher, F, Router
from aiogram.filters import Command, CommandStart
from aiogram.types import (CallbackQuery, InlineKeyboardButton, InlineKeyboardMarkup, MenuButtonWebApp, Message,
                           ReplyKeyboardRemove, WebAppInfo)
from sqlmodel import select

from . import ai
from .api import (DRAFTS, money_of, active_topup, awb_spread, breakdown_from_text, export_topups, book_document, import_floratrack, payments_for, split_by_farm, fill_mawb, money_from_text, save_draft,
                  set_active_topup, store_breakdown, topup_from_text)
from .config import ALLOWED_IDS, BOT_TOKEN, DATA_DIR, MASTER_XLSX, WEBAPP_URL
from .models import Farm, Line, session

LAST_DRAFT: dict[int, tuple[list, float]] = {}   # user -> ([draft ids], time)
bot = Bot(BOT_TOKEN) if BOT_TOKEN else None
dp = Dispatcher()
pv = F.chat.type == "private"          # every dp.message handler is private-only; groups -> `grp` router
from . import roles
ops = F.from_user.id.func(lambda i: roles.role_of(i) is not None)          # any registered user
wr = F.from_user.id.func(lambda i: roles.can_write(roles.role_of(i)))       # sys + super
sysf = F.from_user.id.func(lambda i: roles.role_of(i) == "sys")

# Only INLINE buttons and the menu button pass initData (login) to the Mini App.
# Reply-keyboard buttons open it without initData -> "bad initData".
KB = InlineKeyboardMarkup(inline_keyboard=[[
    InlineKeyboardButton(text="📒 Открыть учёт", web_app=WebAppInfo(url=WEBAPP_URL))]])


async def setup_menu_button():
    """Blue button next to the message field opens the Mini App."""
    if bot and WEBAPP_URL:
        await bot.set_chat_menu_button(menu_button=MenuButtonWebApp(text="Учёт", web_app=WebAppInfo(url=WEBAPP_URL)))


@dp.errors()
async def on_error(event):
    import logging
    logging.exception("handler error", exc_info=event.exception)
    try:
        msg = event.update.message
        if msg:
            await msg.answer(f"❌ Ошибка: {type(event.exception).__name__}: {str(event.exception)[:300]}")
    except Exception:
        pass
    return True


TRUCK_TEXT = F.text.func(lambda t: bool(re.search(r"товар\s+забран|едет на склад|время прибытия|прошла границу", t or "", re.I)))


@dp.message(pv, wr, TRUCK_TEXT)
async def truck_forwarded(m: Message):
    """Messages forwarded (or pasted) from the TK MSK chat: original date is used for «+1 hour»."""
    from datetime import timedelta
    origin = getattr(m, "forward_origin", None)
    sent = (origin.date if origin else m.date).replace(tzinfo=None) + timedelta(hours=3)
    text = m.text
    # pasted (not forwarded) text may start with the original date: «29.09 15:08 …» / «29.09.2026 15:08 …»
    d = re.match(r"\s*(\d{1,2})\.(\d{1,2})(?:\.(\d{2,4}))?[ ,]+(\d{1,2}):(\d{2})\s*", text or "")
    if d and not origin:
        y = int(d.group(3)) if d.group(3) else _msk_now().year
        y = y + 2000 if y < 100 else y
        try:
            sent = datetime(y, int(d.group(2)), int(d.group(1)), int(d.group(4)), int(d.group(5)))
            text = text[d.end():]
        except ValueError:
            pass
    await _truck_event(None, text, sent, notify_uid=m.from_user.id)


@dp.message(pv, CommandStart(), ops)
async def start(m: Message):
    r = roles.role_of(m.from_user.id)
    if not roles.can_write(r):
        await m.answer("Учёт поставок Люмен · роль «1С оператор» (просмотр).\n"
                       "• «Учёт» — смотреть пополнения, инвойсы, себестоимость.\n• /excel — прислать актуальный Excel.",
                       reply_markup=KB if WEBAPP_URL.startswith("https://") else None)
        return
    kb = KB if WEBAPP_URL.startswith("https://") else None
    await m.answer("Обновил кнопки 👇", reply_markup=ReplyKeyboardRemove())   # remove the old reply keyboard
    await m.answer("Учёт поставок Люмен.\n\n"
                   "• Кнопка ниже или «Учёт» слева от поля ввода — пополнения, инвойсы, логистика, выгрузка в Excel.\n"
                   "• Кидай PDF/фото инвойса или счёта за фрахт. С суммами в подписи («1198$ 105472₽») — внесу сразу,"
                   " без подписи — спрошу суммы.\n"
                   "• Документы идут в активное пополнение (последнее). Сменить: /pop, разово — дата в подписи.\n"
                   "• Разбивку кг кидай с MAWB в подписи.\n"
                   "• Отчёт Floratrack (.xlsx) — привяжу машины к MAWB и посчитаю ₽ по вашим оплатам из «Баланса».\n"
                   "• Любой другой .xlsx станет мастер-файлом учёта."
                   + ("" if kb else "\n\n⚠️ WEBAPP_URL не https — кнопка приложения отключена."),
                   reply_markup=kb)


# ---------- /ai (/intelect): несколько файлов + объяснение словами -------------------------
AISESSION: dict = {}   # user -> {"files": [(data, mime, name, file_id)], "note": str, "status": Message}


def _ai_kb():
    return _kb([[("✅ Готово", "ai:go"), ("✖️ Отмена", "ai:cancel")]])


@dp.message(pv, wr, Command("ai", "intelect", "intellect", "intel"))
async def ai_start(m: Message):
    text = (m.text or "").split(maxsplit=1)
    AISESSION[m.from_user.id] = {"files": [], "note": text[1] if len(text) > 1 else "", "status": None}
    st = await m.answer(
        "🧠 Режим «объясню словами». Кидай файлы и пиши, что с ними, например:\n"
        "• «тут 2 файла: в одном инвойс Tambuzi, в другом пакинг»\n"
        "• «Expolanka потеряла коробку, это кредит — минусни со старого MAWB»\n"
        "• «тут инвойс и кредит-нота, кредит вычли из нового инвойса»\n"
        "Когда всё скинул — «✅ Готово».", reply_markup=_ai_kb())
    AISESSION[m.from_user.id]["status"] = st


async def _ai_status(uid: int):
    p = AISESSION.get(uid)
    if not p:
        return
    txt = (f"🧠 Файлов: {len(p['files'])}"
           + ("".join(f"\n📎 {n}" for _d, _m, n, _f in p["files"]))
           + (f"\n📝 {p['note'][:600]}" if p["note"] else "\n📝 (пояснения пока нет — напиши словами, что с файлами)")
           + "\n\nЕщё файлы / текст — или «✅ Готово».")
    try:
        if p.get("status"):
            await p["status"].delete()
    except Exception:
        pass
    p["status"] = await bot.send_message(uid, txt, reply_markup=_ai_kb())


@dp.message(pv, wr, F.from_user.id.func(lambda i: i in AISESSION), F.document | F.photo)
async def ai_collect_file(m: Message):
    p = AISESSION[m.from_user.id]
    if m.photo:
        f = await bot.download(m.photo[-1])
        data, mime, name, fid = f.read(), "image/jpeg", "фото", m.photo[-1].file_id
    else:
        f = await bot.download(m.document)
        data, name, fid = f.read(), m.document.file_name or "файл", m.document.file_id
        mime = m.document.mime_type or ""
        low = name.lower()
        if low.endswith((".xls", ".xlsx")):
            from .xltext import excel_to_text
            data, mime = excel_to_text(data, name).encode(), "text/plain"
        elif low.endswith(".docx"):
            from .xltext import docx_to_text
            data, mime = docx_to_text(data).encode(), "text/plain"
        elif mime not in ("application/pdf", "image/jpeg", "image/png"):
            await m.answer(f"«{name}» — такой формат не читаю. Нужны PDF, фото, Excel или Word.")
            return
    p["files"].append((data, mime, name, fid))
    if m.caption:
        p["note"] = (p["note"] + "\n" + m.caption).strip()
    await _ai_status(m.from_user.id)


@dp.message(pv, wr, F.from_user.id.func(lambda i: i in AISESSION), F.text.func(lambda t: not t.startswith("/")))
async def ai_collect_text(m: Message):
    p = AISESSION[m.from_user.id]
    p["note"] = (p["note"] + "\n" + m.text).strip()
    await _ai_status(m.from_user.id)


@dp.callback_query(F.data.startswith("ai:"), wr)
async def ai_buttons(c: CallbackQuery):
    import asyncio
    uid, what = c.from_user.id, c.data.split(":", 1)[1]
    p = AISESSION.get(uid)
    await c.answer()
    if what == "cancel" or not p:
        AISESSION.pop(uid, None)
        await c.message.edit_text("Ок, режим «объясню словами» закрыт." if what == "cancel" else "Устарело — начни /ai заново.")
        return
    if what == "go":
        if not p["files"]:
            await c.message.edit_text("Файлов пока нет — скинь хотя бы один, потом «✅ Готово».", reply_markup=_ai_kb())
            return
        if len(p["files"]) == 1:
            AISESSION.pop(uid, None)
            await c.message.edit_text("🧠 Читаю с твоим пояснением…")
            asyncio.create_task(_ai_guard(uid, p["files"], p["note"], c.message))
            return
        await c.message.edit_text(
            f"🧠 Файлов: {len(p['files'])}. Это один документ или разные?",
            reply_markup=_kb([[("🧩 Один документ — склеить (инвойс + пакинг)", "ai:one")],
                              [("📑 Разные документы — по очереди", "ai:sep")],
                              [("✖️ Отмена", "ai:cancel")]]))
        return
    AISESSION.pop(uid, None)
    if what == "one":
        await c.message.edit_text("🧠 Склеиваю файлы в один документ и читаю…")
        asyncio.create_task(_ai_guard(uid, p["files"], p["note"], c.message))
    elif what == "sep":
        await c.message.edit_text(f"🧠 Разбираю {len(p['files'])} документа по очереди — на каждый задам обычные вопросы.")
        asyncio.create_task(_ai_queue(uid, p["files"], p["note"]))


def _busy(uid: int) -> bool:
    return uid in FLOW or uid in CORR or uid in NEWFARM or uid in FIXWAIT or uid in FIXPEND


async def _ai_queue(uid: int, files: list, note: str):
    import asyncio, time as _t
    for n, f in enumerate(files, 1):
        t0 = _t.time()
        while _busy(uid) and _t.time() - t0 < 1800:       # wait until the previous document's questions are answered
            await asyncio.sleep(3)
        msg = await bot.send_message(uid, f"🧠 Документ {n} из {len(files)}: {f[2]} — читаю…")
        await _ai_guard(uid, [f], note, msg)


async def _ai_guard(uid: int, files: list, note: str, msg):
    import asyncio, logging, traceback
    try:
        await asyncio.wait_for(_ai_run(uid, files, note, msg), timeout=480)
    except asyncio.TimeoutError:
        await _send(uid, "⏱ AI не ответил за 8 минут. Попробуй ещё раз: /ai", None, msg)
    except Exception:
        logging.exception("ai run failed")
        tb = traceback.format_exc().strip().splitlines()[-1]
        await _send(uid, f"❌ Не смог разобрать: {tb[:300]}", None, msg)


async def _ai_run(uid: int, files: list, note: str, msg):
    from .api import merge_lines, record_upload, update_upload
    with session() as s:
        fs = [f.model_dump() for f in s.exec(select(Farm)).all()]
        catalog = sorted({l.name for l in s.exec(select(Line)).all()})
    ups = []
    for data, mime, name, fid in files:
        try:
            ups.append(record_upload(uid, "", name, mime, tg_file_id=fid, kind="ai"))
        except Exception:
            ups.append(0)
    (d0, m0, _n0, _f0), rest = files[0], files[1:]
    ai_note = ("ПОЯСНЕНИЕ ОПЕРАТОРА (главнее документа, делай как он говорит): " + note) if note.strip() else ""
    out = await ai.parse_document(d0, m0, fs, catalog, note=ai_note, more=[(d, m) for d, m, _n, _f in rest])
    out = merge_lines(out)
    up_id = ups[0] if ups else 0
    out["upload_id"], out["source_file"] = up_id, f"up:{up_id}"
    update_upload(up_id, kind=out.get("doc_type", ""),
                  summary=f"{out.get('farm') or ''} {out.get('invoice_no') or ''} MAWB {out.get('awb') or '?'} (/ai)")
    if len(files) == 1:
        _fix_remember(uid, up_id, d0, m0, note, msg)
    if out.get("doc_type") not in ("farm_invoice", "freight_invoice"):
        await _send(uid, f"Прочитал как «{out.get('doc_type')}». В режиме /ai пока разбираю инвойсы ферм и счета ТК "
                         "(в т.ч. кредит-ноты) — такой документ пришли обычным сообщением.", None, msg)
        return
    fr = out.get("freight") or {}
    pre = ""
    if out.get("doc_type") == "freight_invoice" and (fr.get("total_usd") or 0) < 0:
        pre = (f"🧾 Кредит-нота {out.get('invoice_no') or ''}: −${abs(fr['total_usd']):g} по MAWB {out.get('awb') or '?'} — "
               "запишу минусом на этот MAWB (без кг), скидка разойдётся по его фермам.\n\n")
    if len(files) > 1:
        pre = f"🧩 Склеил {len(files)} файла в один документ.\n" + pre
    warn = ("\n⚠️ " + "\n⚠️ ".join(out["warnings"])) if out.get("warnings") else ""
    # the explanation doubles as a caption («оплачено 2353.95$ 23.10»); a credit note's $ comes from the document only
    await _route_invoice(uid, out, "" if pre.startswith("🧾") or "🧾" in pre else note, msg, warn, pre=pre)


@dp.message(pv, wr, F.document.file_name.lower().endswith(".xls"))
async def excel_xls(m: Message):
    """Old-format Excel = a farm invoice / freight bill. A «.xls» that is really .xlsx (BiFlorica sends such) -> xlsx router."""
    from .xltext import excel_to_text
    f = await bot.download(m.document)
    data = f.read()
    if data[:2] == b"PK":
        await _xlsx_router(m, data)
        return
    text = excel_to_text(data, m.document.file_name)
    if not text.strip():
        await m.answer("Файл пустой или не читается. Пришли PDF или фото инвойса.")
        return
    await _parse_and_reply(m, text.encode(), "text/plain")


@dp.message(pv, wr, F.document.file_name.lower().regexp(r"\.docx?$"))
async def word_doc(m: Message):
    """Word invoice (.docx reliable, old .doc best effort) -> text -> AI."""
    from .xltext import doc_to_text, docx_to_text
    f = await bot.download(m.document)
    data = f.read()
    name = m.document.file_name.lower()
    try:
        text = docx_to_text(data) if name.endswith(".docx") else doc_to_text(data)
    except Exception:
        text = ""
    if len(text.strip()) < 20:
        await m.answer("Не смог прочитать Word-файл. Пересохрани его как .docx или PDF и пришли ещё раз.")
        return
    await _parse_and_reply(m, text.encode(), "text/plain")


@dp.message(pv, wr, F.document.file_name.lower().endswith(".xlsx"))
async def master(m: Message):
    f = await bot.download(m.document)
    await _xlsx_router(m, f.read())


async def _xlsx_router(m: Message, data: bytes):
    from io import BytesIO
    from openpyxl import load_workbook
    from . import floratrack as ft
    try:
        is_ft = ft.is_floratrack(load_workbook(BytesIO(data), read_only=True))
    except Exception:
        is_ft = False
    if is_ft:
        _record(m, "floratrack")
        await _floratrack(m, data)
        return
    from . import biflorica
    if biflorica.is_statement(data):
        _record(m, "broker_statement")
        await _broker_statement(m, data)
        return
    from . import kbreak
    if kbreak.is_breakdown(data):
        _record(m, "kenya_breakdown")
        await _kenya_breakdown(m, data)
        return
    if kbreak.is_weight_report(data):
        _record(m, "consolidation")
        info = kbreak.parse_weight_report(data)
        awb = ai.find_mawb(m.caption or "")
        if not awb:
            _CONS_WAIT[m.from_user.id] = info
            await m.answer(f"⚖️ Отчёт о весе: {sum(r['packs'] for r in info['rows'])} кор., "
                           f"{len(info['rows'])} ферм ({info['country']}). AWB в файле нет — напиши MAWB этого отчёта, "
                           "например `369-99583094`. В следующий раз можно сразу в подписи к файлу.", parse_mode="Markdown")
            return
        await _consolidation(m, awb, info["country"], info["rows"], info["eta"])
        return
    if kbreak.is_prealert(data):
        _record(m, "consolidation")
        info = kbreak.parse_prealert(data)
        await _consolidation(m, info["awb"], info["country"], info["rows"], info["eta"])
        return
    from .xltext import excel_to_text, looks_like_master
    if not looks_like_master(data):                    # not our учёт file -> it's an invoice in Excel
        text = excel_to_text(data, m.document.file_name)
        await _parse_and_reply(m, text.encode(), "text/plain")
        return
    _record(m, "master")
    if MASTER_XLSX.exists():
        shutil.copy(MASTER_XLSX, DATA_DIR / f"учет_backup_{datetime.now():%Y%m%d_%H%M}.xlsx")
    MASTER_XLSX.write_bytes(data)
    from .backup import mark_dirty
    mark_dirty()
    await m.answer("Мастер-файл обновлён ✅ (старый сохранён в бэкап)")


async def _broker_statement(m: Message, data: bytes):
    """BiFlorica statement: deposits + purchases of Tessa / Plazoleta -> books, balance, cost per stem."""
    from . import biflorica
    from .api import farm_balance_text, transit_view
    from .models import BrokerDeposit, Invoice, TopUp
    try:
        r = biflorica.import_statement(data)
    except Exception as e:
        await m.answer(f"Не смог разобрать выписку брокера: {e}")
        return
    fmt = lambda x: f"{x:,.2f}".replace(",", " ")
    deps = [o for o in r["ops"] if o["kind"] == "dep"]
    buys = [o for o in r["ops"] if o["kind"] == "buy"]
    lines = [f"🏦 Выписка брокера BiFlorica: {len(deps)} пополн., {len(buys)} закупок"
             + (f" (новых: {len(r['new_dep'])} пополн., {len(r['new_buy']) + len(r['linked'])} закупок)" if r["new_dep"] or r["new_buy"] or r["linked"] else " — всё уже было внесено")]
    rows = []
    with session() as s:
        tdate = {t.id: (t.date, t.rub / t.usd if t.usd else 0) for t in s.exec(select(TopUp)).all()}
        for did in r["new_dep"]:
            d = s.get(BrokerDeposit, did)
            td = tdate.get(d.topup_id, ("?", 0))
            lines.append(f"💵 {d.date[8:10]}.{d.date[5:7]}: на биржу ${d.usd_credited:g} (отправлено ${d.usd_sent:g}) — из пополнения {td[0]}"
                         + (f" → {d.usd_sent * td[1] / d.usd_credited:.2f} ₽ за $ на бирже" if td[1] else ""))
            rows.append([(f"↔️ Другое пополнение для ${d.usd_credited:g} от {d.date[8:10]}.{d.date[5:7]}", f"bfd:{d.id}")])
    by_farm = {}
    for o in buys:
        f = by_farm.setdefault(o["farm"], [0, 0.0])
        f[0] += o["stems"]; f[1] += o["amount"] + o["fee"]
    for farm, (st, usd) in by_farm.items():
        fee = sum(o["fee"] for o in buys if o["farm"] == farm); base = sum(o["amount"] for o in buys if o["farm"] == farm)
        lines.append(f"🌸 {farm}: {st:g} ст на ${fmt(usd)} (с комиссией {fee / base * 100 if base else 0:.0f}%)")
    tv = {i["id"]: i for i in transit_view()["invoices"]}
    new = [tv[i] for i in r["new_buy"] + r["linked"] if i in tv]
    if new:
        lines.append("\nСебестоимость цветка (без логистики):")
        for i in new[:8]:
            l = i["lines"][0] if i["lines"] else None
            if l:
                lines.append(f"• {i['farm']} {i['invoice_date']}: {l['name'][:40]} — {l['price_rub']:.2f} ₽/ст")
    if r.get("new_farms"):
        lines.append("\n🆕 Новые фермы у посредника: " + ", ".join(r["new_farms"])
                     + " — добавил в плантации как «через посредника». Проверь название и страну в «Ещё» → «Плантации».")
    if r.get("fee_fix"):
        lines.append("ℹ️ Комиссия посредника по выписке: " + "; ".join(r["fee_fix"]))
    lines.append("\n" + farm_balance_text("Tessa"))
    await send_packing_lists()
    await m.answer("\n".join(lines)[:4000], reply_markup=_kb(rows) if rows else None)


@dp.callback_query(F.data.startswith("bfd:"), wr)
async def broker_dep_topup(c: CallbackQuery):
    from .models import TopUp
    did = int(c.data.split(":")[1])
    with session() as s:
        tops = s.exec(select(TopUp).order_by(TopUp.id.desc()).limit(8)).all()
    await c.answer()
    await c.message.answer("Из какого пополнения ушли эти доллары на биржу?",
                           reply_markup=_kb([[(f"{t.date} · курс {t.rub / t.usd:.2f}", f"bft:{did}:{t.id}")] for t in tops]))


@dp.callback_query(F.data.startswith("bft:"), wr)
async def broker_dep_topup_set(c: CallbackQuery):
    from .models import BrokerDeposit, TopUp
    _, did, tid = c.data.split(":")
    with session() as s:
        d, t = s.get(BrokerDeposit, int(did)), s.get(TopUp, int(tid))
        d.topup_id = t.id
        s.add(d); s.commit()
        txt = f"✅ ${d.usd_credited:g} от {d.date} — из пополнения {t.date}: {d.usd_sent * t.rub / t.usd / d.usd_credited:.2f} ₽ за $ на бирже"
    from .backup import mark_dirty
    mark_dirty()
    await c.answer()
    await c.message.edit_text(txt + "\nСебестоимость Tessa / Plazoleta пересчитана.")


def _farm_label(raw: str) -> str:
    """Unknown shipper on a consolidation list: «Buitron Martin Stefania Renee (B&M Fiori)» -> «B&M Fiori»,
    «FLORICULTORA X S. A.» -> «X»."""
    m = re.search(r"\(([^)]+)\)", raw or "")
    if m:
        return m.group(1).strip()
    n = re.sub(r"[,.]?\s*\b(s\.?\s?a\.?\s?s?\.?|ltd|limited|llc|inc|b\.?v\.?|cia\.?|ltda\.?)\s*$", "", raw or "", flags=re.I)
    n = re.sub(r"^\s*(floricultora|florícola|floricola|agricola|agrícola|hacienda)\s+", "", n, flags=re.I).strip(" ,.-")
    return n.title() if n.isupper() or n.islower() else n


async def _consolidation(m: Message, awb_raw: str, country: str, rows: list, eta=None):
    """Ecuador / Colombia consolidation list: boxes per farm (+kg if given) -> packing chats get it before packings."""
    from . import kbreak
    from .ai import find_mawb
    from .calc import norm_awb
    from .api import resolve_farm
    awb = find_mawb(awb_raw or "") or awb_raw
    if not awb or not rows:
        await m.answer("Не нашёл в листе AWB и строки по фермам.")
        return
    with session() as s:
        for r in rows:
            f = resolve_farm(s, r["farm_raw"])
            r["farm"] = f.name if f else _farm_label(r["farm_raw"])
    farms = [{"farm": r["farm"], "packs": int(r.get("packs") or 0), "kg": float(r.get("kg") or 0),
              "hawb": (r.get("hawb") or "").strip().upper() or None} for r in rows]
    from .api import assign_awb_by_farms
    assign_awb_by_farms(awb, [f["farm"] for f in farms])
    w = None
    if any(f["kg"] for f in farms):
        w = store_breakdown(awb, [{"farm": f["farm"], "kg": f["kg"]} for f in farms if f["kg"]])
    was_sent = kbreak.is_sent(norm_awb(awb))
    kbreak.store_rows(norm_awb(awb), awb, country, farms, eta)
    lines = "\n".join(f"• {f['farm']}: {f['packs']} кор." + (f" · {f['kg']:g} кг" if f["kg"] else "") for f in farms)
    arrive = kbreak._state().get(norm_awb(awb), {}).get("arrive")
    await m.answer(f"📋 Консолидация {kbreak.flag(country)} · MAWB {awb}: {sum(f['packs'] for f in farms)} кор.\n{lines}"
                   + (("\n\n" + _applied_text(w)) if w else "")
                   + (f"\n🛬 Прибытие ориентировочно {arrive}" if arrive else "")
                   + "\n\n📦 В чаты пакингов уйдёт ОДНИМ файлом: детализация + пакинг-листы инвойсов этого MAWB.")
    from . import kbreak as _kb_
    await send_packing_lists()
    await missing_reminder([m.from_user.id])
    if was_sent:                      # list changed after the shipment went out -> send the whole shipment again
        from . import kbreak as _kb2
        if norm_awb(awb) not in {x["awb_key"] for x in _kb2.missing_invoices()}:   # only a COMPLETE shipment
            await push_packing_all(norm_awb(awb))
    await send_packing_lists()


async def _kenya_breakdown(m: Message, data: bytes):
    """TK Kenya box breakdown: kg per farm for the MAWB + the file goes to packing chats before the packings."""
    from . import kbreak
    from .ai import find_mawb
    from .calc import norm_awb
    from .api import resolve_farm
    info = kbreak.parse(data)
    awb = find_mawb(info["awb"] or "") or info["awb"]
    if not awb or not info["rows"]:
        await m.answer("Не нашёл в файле AWB и строки по фермам.")
        return
    with session() as s:
        for r in info["rows"]:
            f = resolve_farm(s, r["farm_raw"])
            r["farm"] = f.name if f else _farm_label(r["farm_raw"])
    col = info["use"]
    per = [{"farm": r["farm"], "kg": round(r[col], 2), "boxes": r["packs"]} for r in info["rows"]]
    w = store_breakdown(awb, per)
    was_sent = kbreak.is_sent(norm_awb(awb))
    kbreak.store(data, info, norm_awb(awb), [{"farm": r["farm"], "packs": r["packs"], "kg": r[col]} for r in info["rows"]])
    if was_sent:                      # list changed after the shipment went out -> send the whole shipment again
        from . import kbreak as _kb2
        if norm_awb(awb) not in {x["awb_key"] for x in _kb2.missing_invoices()}:   # only a COMPLETE shipment
            await push_packing_all(norm_awb(awb))
    lines = "\n".join(f"• {r['farm']}: {r['packs']} кор. · {r[col]:g} кг" for r in info["rows"])
    txt = (f"📋 Детализация MAWB {awb}: {info['packs']} кор., {info[col]:g} кг "
           f"({'реальный вес' if col == 'weight' else 'объёмный вес'} — он больше: {info['weight']:g} / VW {info['vw']:g})\n"
           + lines + "\n\n" + _applied_text(w)
           + "\n\n📦 В чаты пакингов уйдёт ОДНИМ файлом: детализация (без ETD/ETA) + пакинг-листы кенийских инвойсов этого MAWB."
           + (f"\n🛬 Прибытие ориентировочно {kbreak._state().get(norm_awb(awb), {}).get('arrive')}" if info.get("eta") else ""))
    await m.answer(txt)
    await send_packing_lists()
    await missing_reminder([m.from_user.id])


async def _floratrack(m: Message, data: bytes):
    note = await m.answer("Отчёт Floratrack — разбираю машины и оплаты…")
    try:
        r = import_floratrack(data)
    except Exception as e:
        await note.edit_text(f"Не смог разобрать отчёт: {e}")
        return
    fmt = lambda x: f"{x:,.0f}".replace(",", " ")
    kinds = {"import": "Кения АМС-МСК", "ecuador": "Эквадор", "colombia": "Колумбия"}
    recent = sorted(r["matched"], key=lambda x: x[1].date or datetime.min, reverse=True)[:12]
    lines = [f"• {awb} · {kinds[c.kind]} · {c.kg:g} кг · ${c.usd:,.2f} × {c.rate:.2f} = {fmt(c.rub)} ₽"
             + (" (курс предв.)" if c.provisional else "") for awb, c in recent]
    txt = (f"🚚 Floratrack: привязано {len(r['matched'])} AWB (новых {r['added']}, обновлено {r['updated']}).\n"
           + (f"Баланс у Floratrack: ${r['balance_usd']:,.2f}\n" if roles.sees_money(roles.role_of(m.from_user.id)) else "")
           + "\n" + "\n".join(lines))
    fresh = [c for c in r["unmatched"] if c.date and (datetime.now() - c.date).days <= 21]
    if fresh:
        txt += "\n\nНет наших инвойсов/разбивки с этими AWB (последние 3 недели):\n" + "\n".join(
            f"• …{c.last4} · {kinds[c.kind]} · {c.kg:g} кг · {c.sheet}" for c in fresh[:15])
        txt += "\nВнеси их и перекинь отчёт — привяжутся."
    if r.get("offsets") and roles.sees_money(roles.role_of(m.from_user.id)):
        txt += ("\n\n⚠️ Floratrack списал ваши оплаты не на машины:\n" + "\n".join(
            f"• {d:%d.%m} · ${u:,.2f} · «{n or 'без пояснения'}»" for d, u, n in r["offsets"])
            + "\nЭти деньги на ваши машины не пошли — поэтому последние машины «курс предв.» и долг перед Floratrack.")
    if r["ambiguous"]:
        txt += "\n\n⚠️ Несколько MAWB с такими 4 цифрами, взял по весу: " + ", ".join(r["ambiguous"][:10])
    affected = sorted({x["topup_id"] for awb, _c in r["matched"] for x in awb_spread(awb)})
    warn = [w for w in r["warnings"] if "MAWB" in w][:8]
    if warn:
        txt += "\n\n⚠️ " + "\n⚠️ ".join(warn)
    await note.edit_text(txt[:4000])
    if affected:
        await export_topups(affected, m.from_user.id, "🚚 Floratrack разнесён по пополнениям, где лежит этот товар.")


async def _parse_and_reply(m: Message, data: bytes, mime: str):
    """Guard: never leave the user staring at 'Читаю документ…'."""
    import asyncio, logging, traceback
    note_holder = {}
    async def _ticker():
        for n, txt in enumerate(["Читаю документ… (большой инвойс — раскладываю коробки по сортам)",
                                 "Ещё читаю… много позиций, обычно до 3–5 минут",
                                 "Почти… проверяю, что сорта сходятся с итогом"]):
            await asyncio.sleep(60)
            if note_holder.get("note") and not note_holder.get("done"):
                try:
                    await note_holder["note"].edit_text(txt)
                except Exception:
                    pass
    tick = asyncio.create_task(_ticker())
    try:
        await asyncio.wait_for(_parse_and_reply_inner(m, data, mime, note_holder), timeout=480)
    except asyncio.TimeoutError:
        await _say(m, note_holder, "⏱ AI не ответил за 8 минут. Пришли документ ещё раз — обычно со второго раза проходит.")
    except Exception as e:
        logging.exception("parse failed")
        tb = traceback.format_exc().strip().splitlines()[-1]
        await _say(m, note_holder, f"❌ Ошибка при разборе: {tb[:300]}\nПерешли это сообщение мне (разработчику) — починю.")
    finally:
        note_holder["done"] = True
        tick.cancel()


async def _say(m: Message, holder: dict, text: str):
    try:
        if holder.get("note"):
            await holder["note"].edit_text(text)
            return
    except Exception:
        pass
    await m.answer(text)


def _record(m: Message, kind: str = "") -> int:
    """Keep the original file (Telegram file_id) with who/when — visible to the system admin forever."""
    from .api import record_upload
    doc = m.document
    fid = doc.file_id if doc else (m.photo[-1].file_id if m.photo else None)
    name = (doc.file_name if doc else "") or ("фото" if m.photo else "")
    mime = (doc.mime_type if doc else "image/jpeg") or ""
    u = m.from_user
    who = (u.full_name or "") + (f" (@{u.username})" if u.username else "")
    try:
        return record_upload(u.id, who, name, mime, tg_file_id=fid, kind=kind)
    except Exception as e:
        print(f"[lumen] archive: {e}", flush=True)
        return 0


async def _parse_and_reply_inner(m: Message, data: bytes, mime: str, holder: dict):
    if m.from_user.id in BOXWAIT:
        iid = BOXWAIT.pop(m.from_user.id)
        note = await m.answer("Читаю детализацию по коробкам…")
        holder["note"] = note
        from .api import set_box_detail, _boxes_from_raw
        out = await ai.parse_document(data, mime, [], [], note="Это детализация по коробкам (packing list) к уже внесённому "
                                      "инвойсу: заполни boxes_detail — что лежит в каждой коробке.")
        boxes = out.get("boxes_detail") or _boxes_from_raw(out.get("lines") or [])
        if not boxes:
            await _send(m.from_user.id, "Не нашёл в документе раскладку по коробкам. Пришли другой файл.", None, note)
            BOXWAIT[m.from_user.id] = iid
            return
        txt, key = set_box_detail(iid, boxes)
        edited = await refresh_shipment(key, "добавлена детализация по коробкам") if key else 0
        await _send(m.from_user.id, txt + ("\n🔄 Пакинг в чатах обновлён." if edited else ""), None, note)
        return
    note = await m.answer("Читаю документ…")
    holder["note"] = note
    up_id = _record(m)
    with session() as s:
        fs = [f.model_dump() for f in s.exec(select(Farm)).all()]
        catalog = sorted({l.name for l in s.exec(select(Line)).all()})
    try:
        out = await ai.parse_document(data, mime, fs, catalog, note=m.caption or "")
        from .api import merge_lines
        out = merge_lines(out)
    except Exception as e:
        await note.edit_text(f"Не смог прочитать: {e}")
        return
    warn = ("\n⚠️ " + "\n⚠️ ".join(out["warnings"])) if out.get("warnings") else ""
    if warn and out.get("doc_type") == "farm_invoice":
        warn += "\n✏️ Если бот ошибся — нажми «Исправить текстом» или ответь на это сообщение: «там 840 ст и 24 коробки»."
    from .api import update_upload
    update_upload(up_id, kind=out.get("doc_type", ""),
                  summary=f"{out.get('farm') or ''} {out.get('invoice_no') or ''} MAWB {out.get('awb') or '?'} ${out.get('invoice_total_usd') or ''}")
    out["upload_id"], out["source_file"] = up_id, f"up:{up_id}"

    if out.get("doc_type") == "consolidation":
        from datetime import datetime as _dt
        cons = out.get("consolidation") or {}
        country = {"BOG": "Колумбия", "MDE": "Колумбия", "UIO": "Эквадор", "GYE": "Эквадор", "NBO": "Кения"}.get(
            (cons.get("origin") or "").upper(), out.get("country") or "Колумбия")
        try:
            eta = _dt.strptime(cons.get("eta") or "", "%Y-%m-%d")
        except ValueError:
            eta = None
        rows = [{"farm_raw": r.get("shipper") or "", "packs": r.get("boxes") or 0, "kg": r.get("weight") or 0,
                 "hawb": r.get("hawb")}
                for r in cons.get("rows") or [] if r.get("shipper")]
        try:
            await note.delete()
        except Exception:
            pass
        await _consolidation(m, cons.get("awb") or "", country, rows, eta)
        return

    if out.get("doc_type") == "topup_receipt":
        await note.edit_text(_register_topup(out.get("topup") or {}, m.caption or ""))
        return

    if out.get("doc_type") == "kg_breakdown":
        kg = out.get("per_farm_kg") or []
        head = (f"⚖️ Разбивка кг · MAWB {out.get('awb') or '?'}\n"
                + "\n".join(f"• {x.get('farm')}: {x.get('kg'):g} кг" for x in kg)
                + f"\nИтого {sum(x.get('kg') or 0 for x in kg):g} кг" + warn)
        if out.get("awb"):
            await note.edit_text(head + "\n\n" + _applied_text(store_breakdown(out["awb"], kg, out.get("source_file"))))
            return
        did = save_draft(out)
        LAST_DRAFT[m.from_user.id] = ([did], time.time())
        await note.edit_text(head + "\n\nНапиши MAWB следующим сообщением — сохраню разбивку.")
        return

    # farm invoice / freight bill -> «оплачен?» flow
    _fix_remember(m.from_user.id, up_id, data, mime, m.caption or "", note)
    await _route_invoice(m.from_user.id, out, m.caption or "", note, warn)


async def _route_invoice(uid: int, out: dict, caption: str, note, warn: str, pre: str = ""):
    """Parsed farm invoice / freight bill -> broker packing, new farm, correction question or the «оплачен?» flow."""
    subs = split_by_farm(out) if out.get("doc_type") == "farm_invoice" else [out]
    if out.get("doc_type") == "farm_invoice":
        from .api import attach_broker_packing, is_broker_farm
        broker = [d for d in subs if is_broker_farm(d.get("farm") or "")]
        if broker:
            for d in broker:
                fill_mawb(d, None)
            msg = "\n\n".join(attach_broker_packing(d) for d in broker)
            for k in {d.get("_refresh_awb") for d in broker if d.get("_refresh_awb")}:
                if await refresh_shipment(k, "обновлён пакинг " + ", ".join(d.get("farm") or "" for d in broker)):
                    msg += f"\n🔄 Пакинг MAWB {k} в чатах обновлён."
            subs = [d for d in subs if d not in broker]
            if not subs:
                await _send(uid, pre + msg, _fix_kb(out), note)
                await send_packing_lists()
                return
            await bot.send_message(uid, pre + msg, reply_markup=_fix_kb(out))
    if out.get("doc_type") == "farm_invoice" and any(not d.get("farm") for d in subs):
        raw = out.get("supplier") or next((mm.group(1) for w in out.get("warnings", [])
                                           for mm in [re.search(r"[Пп]оставщик\s+['«\"]([^'»\"]+)", w)] if mm), "")
        if raw:
            NEWFARM[uid] = {"raw": raw, "subs": subs, "out": out, "caption": caption,
                                       "head": _flow_head(subs)}
            await _send(uid, f"🌱 Новая ферма: «{raw}» — её нет в списке плантаций.\nДобавить? Выбери страну:",
                        _kb([[("🇰🇪 Кения", "nf:Кения"), ("🇪🇨 Эквадор", "nf:Эквадор")],
                             [("🇨🇴 Колумбия", "nf:Колумбия"), ("🇳🇱 Нидерланды", "nf:Нидерланды")],
                             [("Не добавлять", "nf:skip")]]), note)
            return
    if out.get("doc_type") == "farm_invoice":
        from .api import find_same_invoice
        same = [(d, find_same_invoice(d)) for d in subs]
        same = [(d, inv) for d, inv in same if inv]
        if same and "новый" not in caption.lower():
            d, inv = same[0]
            CORR[uid] = {"pairs": [(save_draft(dd), ii["id"]) for dd, ii in same],
                                    "subs": subs, "out": out, "caption": caption, "head": pre + _flow_head(subs) + warn}
            who = ", ".join(ii["farm"] for _dd, ii in same)
            await _send(uid,
                        f"📝 Инвойс {who} № {inv['invoice_no']} уже есть в учёте ({inv['where']}).\nЭто КОРРЕКТИРОВКА инвойса?",
                        _kb([[("✅ Да, корректировка — заменить", "cor:yes")],
                             [("➕ Нет, это новый инвойс", "cor:no")]] + _fix_rows(out)), note)
            return
    await _start_flow(uid, subs, out, caption, note, pre + _flow_head(subs) + warn)



# ---------- «✏️ бот ошибся — исправить текстом» ------------------------------------------
FIXWAIT: dict = {}    # user -> upload id: the next text from him is a correction of that invoice
MSG2UP: dict = {}     # (user, bot message id) -> upload id: a reply to that message = correction
LASTUP: dict = {}     # user -> last invoice upload id (for /fix)


def _fixdir():
    from .api import DRAFTS
    d = DRAFTS.parent / "fixsrc"
    d.mkdir(exist_ok=True)
    return d


def _fix_rows(out) -> list:
    up = out.get("upload_id") if isinstance(out, dict) else out
    return [[("✏️ Бот ошибся — исправить текстом", f"fx:{up}")]] if up else []


def _fix_kb(out):
    rows = _fix_rows(out)
    return _kb(rows) if rows else None


def _fix_remember(uid: int, up_id: int, data: bytes, mime: str, caption: str, note=None):
    """Keep what the AI read (bytes / extracted text) so a typed correction can re-read the same document."""
    if not up_id:
        return
    LASTUP[uid] = up_id
    if note is not None:
        MSG2UP[(uid, note.message_id)] = up_id
    try:
        d = _fixdir()
        (d / f"{up_id}.bin").write_bytes(data)
        (d / f"{up_id}.json").write_text(json.dumps({"mime": mime, "caption": caption}, ensure_ascii=False))
        old = sorted(d.glob("*.bin"), key=lambda p: p.stat().st_mtime)[:-80]
        for p in old:
            p.unlink(missing_ok=True)
            p.with_suffix(".json").unlink(missing_ok=True)
    except Exception as e:
        print(f"[lumen] fixsrc: {e}", flush=True)


async def _fix_load(up_id: int):
    """-> (data, mime, caption) of the original document, or None."""
    d = _fixdir()
    if (d / f"{up_id}.bin").exists():
        meta = json.loads((d / f"{up_id}.json").read_text()) if (d / f"{up_id}.json").exists() else {}
        return (d / f"{up_id}.bin").read_bytes(), meta.get("mime") or "application/pdf", meta.get("caption") or ""
    from .models import Upload
    with session() as s:
        u = s.get(Upload, up_id)
        fid, name, mime = (u.tg_file_id, (u.filename or "").lower(), u.mime or "") if u else (None, "", "")
    if not fid:
        return None
    f = await bot.download(fid)
    data = f.read()
    if name.endswith((".xls", ".xlsx")):
        from .xltext import excel_to_text
        return excel_to_text(data, name).encode(), "text/plain", ""
    if name.endswith(".docx"):
        from .xltext import docx_to_text
        return docx_to_text(data).encode(), "text/plain", ""
    if name.endswith(".doc"):
        from .xltext import doc_to_text
        return doc_to_text(data).encode(), "text/plain", ""
    return data, ("image/jpeg" if mime.startswith("image") or name == "фото" else mime or "application/pdf"), ""


def _boxes_text(boxes: list, limit: int = 30) -> str:
    """boxes_detail -> «• 3 кор: Light Blue 25/кор» lines for checking the packing before it is applied."""
    rows = []
    for b in boxes or []:
        parts = []
        for c in b.get("content") or []:
            n = c.get("stems_per_box")
            parts.append(f"{c.get('name')} {n:g}/кор" if isinstance(n, (int, float)) else
                         f"{c.get('name')} {float(c.get('stems_total') or c.get('stems') or 0):g} ст")
        rows.append(f"• {int(b.get('qty') or 1)} кор{(' ' + b['pack']) if b.get('pack') else ''}: " + ", ".join(parts))
    if len(rows) > limit:
        rows = rows[:limit] + [f"… ещё {len(rows) - limit} строк"]
    return "\n".join(rows)


def _d_totals(d: dict):
    stems = sum(l.get("stems") or 0 for l in d.get("lines") or [])
    boxes = sum(int(b.get("qty") or 1) for b in d.get("boxes_detail") or []) or \
        sum(l.get("boxes") or 0 for l in d.get("lines") or []) or (d.get("boxes_total") or 0)
    return stems, boxes


def _inv_totals(inv_id: int):
    from .models import Invoice
    with session() as s:
        inv = s.get(Invoice, inv_id)
        if inv.via_broker and inv.packing_lines_json:
            lines = json.loads(inv.packing_lines_json)
        else:
            lines = [{"stems": l.stems, "boxes": l.boxes} for l in s.exec(select(Line).where(Line.invoice_id == inv_id)).all()]
        bx = json.loads(inv.boxes_json) if inv.boxes_json else []
    return _d_totals({"lines": lines, "boxes_detail": bx})


def _booked_for(up_id: int, subs: list) -> list:
    """[(sub, invoice id | None)] — which parts of this document are already in the books."""
    from .models import Invoice
    from .api import find_same_invoice, _norm_name
    with session() as s:
        invs = [(i.id, i.farm) for i in s.exec(select(Invoice).where(Invoice.source_file == f"up:{up_id}")).all()]
    res = []
    for d in subs:
        hit = next((iid for iid, f in invs if _norm_name(f) == _norm_name(d.get("farm") or "")), None)
        if hit is None and len(invs) == 1 and len(subs) == 1:
            hit = invs[0][0]
        if hit is None:
            same = find_same_invoice(d)
            hit = same["id"] if same else None
        res.append((d, hit))
    return res


def _fmt_n(x) -> str:
    return f"{x:g}" if isinstance(x, (int, float)) else str(x)


def _want_numbers(text: str):
    t = text.lower()
    st = re.search(r"(\d[\d\s]*)\s*(?:ст\b|стеб|stem)", t)
    bx = re.search(r"(\d+)\s*(?:кор|box|кб\b)", t)
    return (int(re.sub(r"\D", "", st.group(1))) if st else None), (int(bx.group(1)) if bx else None)


def _drop_pending(uid: int, up_id: int):
    """The old reading of this document is still waiting in a flow / question -> throw it away."""
    from .api import DRAFTS
    fl = FLOW.get(uid)
    if fl:
        docs = _flow_docs(fl)
        if docs and any(d.get("upload_id") == up_id for _p, d in docs):
            for p, _d in docs:
                p.unlink(missing_ok=True)
            FLOW.pop(uid, None)
    p = CORR.get(uid)
    if p and (p.get("out") or {}).get("upload_id") == up_id:
        for did, _iid in p.get("pairs") or []:
            (DRAFTS / f"{did}.json").unlink(missing_ok=True)
        CORR.pop(uid, None)
    p = NEWFARM.get(uid)
    if p and (p.get("out") or {}).get("upload_id") == up_id:
        NEWFARM.pop(uid, None)


async def _apply_fix(uid: int, up_id: int, text: str):
    import asyncio, logging, traceback
    note = await bot.send_message(uid, "✏️ Перечитываю инвойс с твоим исправлением… (до пары минут)")
    MSG2UP[(uid, note.message_id)] = up_id
    try:
        await asyncio.wait_for(_apply_fix_inner(uid, up_id, text, note), timeout=480)
    except asyncio.TimeoutError:
        await _send(uid, "⏱ AI не ответил за 8 минут. Напиши исправление ещё раз.", _fix_kb(up_id), note)
    except Exception:
        logging.exception("fix failed")
        tb = traceback.format_exc().strip().splitlines()[-1]
        await _send(uid, f"❌ Ошибка при исправлении: {tb[:300]}", _fix_kb(up_id), note)


async def _apply_fix_inner(uid: int, up_id: int, text: str, note):
    from .api import merge_lines, correct_invoice, invoice_costs, DRAFTS
    src = await _fix_load(up_id)
    if not src:
        await _send(uid, "Не нашёл исходный файл этого инвойса. Пришли его заново и напиши исправление в подписи к файлу.",
                    None, note)
        return
    data, mime, caption = src
    # what we read last time — from the books or from the drafts still waiting
    prev = []
    with session() as s:
        from .models import Invoice
        for i in s.exec(select(Invoice).where(Invoice.source_file == f"up:{up_id}")).all():
            st, bx = _inv_totals(i.id)
            prev.append(f"{i.farm}: {st:g} стеблей, {bx:g} коробок")
    if not prev:
        for p in DRAFTS.glob("*.json"):
            try:
                d = json.loads(p.read_text())
            except ValueError:
                continue
            if d.get("upload_id") == up_id:
                st, bx = _d_totals(d)
                prev.append(f"{d.get('farm') or '?'}: {st:g} стеблей, {bx:g} коробок")
    ai_note = ((caption + "\n\n") if caption else "") + (
        "ИСПРАВЛЕНИЕ ОТ ОПЕРАТОРА — оно главнее твоего прошлого прочтения: «" + text.strip() + "».\n"
        + (("В прошлый раз ты прочитал так: " + "; ".join(prev) + " — это неверно.\n") if prev else "")
        + "Перечитай документ заново, строку за строкой, ничего не пропуская (строки на стыке страниц, "
          "последнюю строку перед TOTAL, коробки с одной строкой). Итоговые стебли и коробки должны совпасть "
          "с исправлением оператора и со строкой TOTAL. boxes_total и stems_total — как в исправлении, если оно их называет.\n"
          "НЕ подгоняй цифры: не увеличивай стебли существующих строк, чтобы сойтись с итогом. Пропущенная строка — это "
          "отдельная строка документа со своим сортом и своей коробкой. Если оператор сам описал строку или коробку "
          "(сорт, длина, стебли) — добавь её ровно так. Если строку в документе так и не нашёл — не выдумывай, "
          "а напиши в warnings, чего не нашёл. boxes_detail — что лежит в каждой коробке, как в документе.")
    with session() as s:
        fs = [f.model_dump() for f in s.exec(select(Farm)).all()]
        catalog = sorted({l.name for l in s.exec(select(Line)).all()})
    out = merge_lines(await ai.parse_document(data, mime, fs, catalog, note=ai_note))
    out["upload_id"], out["source_file"] = up_id, f"up:{up_id}"
    if out.get("doc_type") != "farm_invoice":
        await _send(uid, "Перечитал, но это не похоже на инвойс фермы — исправлять текстом можно только инвойсы ферм.",
                    None, note)
        return
    want_st, want_bx = _want_numbers(text)
    new_st, new_bx = _d_totals(out)
    check = ""
    if (want_st and abs(new_st - want_st) >= 1) or (want_bx and new_bx != want_bx):
        check = (f"\n\n⚠️ Ты написал {(_fmt_n(want_st) + ' ст') if want_st else ''}"
                 f"{' и ' if want_st and want_bx else ''}{(_fmt_n(want_bx) + ' кор') if want_bx else ''}, "
                 f"а у меня сейчас {new_st:g} ст / {new_bx:g} кор. Напиши подробнее, какой строки не хватает "
                 "(сорт, длина, стеблей, в какой коробке) — ответь на это сообщение.")
    subs = split_by_farm(out)
    pairs = _booked_for(up_id, subs)
    booked = [(d, iid) for d, iid in pairs if iid]
    if not booked:
        # not in the books yet: replace the old reading and continue the usual questions with the new one
        _drop_pending(uid, up_id)
        _fix_remember(uid, up_id, data, mime, caption, note)
        pre = "✏️ Перечитал с исправлением.\n" + ("Было: " + "; ".join(prev) + "\n" if prev else "") + \
              f"Стало: {new_st:g} стеблей, {new_bx:g} коробок.{check}\n" + \
              (("По коробкам (так пойдёт в пакинг):\n" + _boxes_text(out.get("boxes_detail")) + "\n") if out.get("boxes_detail") else "") + "\n"
        warn = ("\n⚠️ " + "\n⚠️ ".join(out["warnings"])) if out.get("warnings") else ""
        await _route_invoice(uid, out, caption, note, warn, pre=pre)
        return
    rest = [d for d, iid in pairs if not iid]
    FIXPEND[uid] = {"up": up_id, "pairs": [(save_draft(d), iid) for d, iid in booked]}
    txt = "✏️ Перечитал инвойс с исправлением. Проверь, прежде чем применю:\n"
    for d, iid in booked:
        b_st, b_bx = _inv_totals(iid)
        a_st, a_bx = _d_totals(d)
        with session() as s:
            from .models import Invoice
            inv = s.get(Invoice, iid)
            old = {l.name: l.stems for l in s.exec(select(Line).where(Line.invoice_id == iid)).all()}
            if inv.via_broker and inv.packing_lines_json:
                old = {x.get("name"): x.get("stems") for x in json.loads(inv.packing_lines_json)}
            head = f"{inv.farm} № {inv.invoice_no}"
        newl = {}
        for l in d.get("lines") or []:
            newl[l.get("name")] = newl.get(l.get("name"), 0) + (l.get("stems") or 0)
        ch = [f"{n}: {_fmt_n(old.get(n, '—'))} → {_fmt_n(newl[n]) if n in newl else 'убран'}"
              for n in sorted(set(old) | set(newl)) if old.get(n) != newl.get(n)]
        txt += (f"\n📄 {head}: стеблей {b_st:g} → {a_st:g}, коробок {b_bx:g} → {a_bx:g}\n"
                + ("\n".join("• " + x for x in ch[:25]) + "\n" if ch else "")
                + (("По коробкам (так пойдёт в пакинг):\n" + _boxes_text(d.get("boxes_detail")) + "\n")
                   if d.get("boxes_detail") else ""))
    if rest:
        txt += "\nНе нашёл в учёте: " + ", ".join(d.get("farm") or "?" for d in rest) + " — пришли их инвойс отдельно."
    if out.get("warnings"):
        txt += "\n⚠️ " + "\n⚠️ ".join(out["warnings"])
    await _send(uid, (txt + check)[:4000],
                _kb([[("✅ Применить", "fxok"), ("✏️ Ещё поправить", f"fx:{up_id}")], [("✖️ Не менять", "fxno")]]), note)


FIXPEND: dict = {}    # user -> corrected reading waiting for «✅ Применить»


@dp.callback_query(F.data.in_({"fxok", "fxno"}), wr)
async def fix_apply(c: CallbackQuery):
    from .api import correct_invoice, invoice_costs, DRAFTS
    p = FIXPEND.pop(c.from_user.id, None)
    await c.answer()
    if not p:
        await c.message.edit_text("Устарело — напиши исправление ещё раз.")
        return
    if c.data == "fxno":
        for did, _iid in p["pairs"]:
            (DRAFTS / f"{did}.json").unlink(missing_ok=True)
        await c.message.edit_text(c.message.text.split("\n")[0] + "\n✖️ Ничего не менял.", reply_markup=_fix_kb(p["up"]))
        return
    txt, done, keys = "", [], set()
    for did, iid in p["pairs"]:
        f = DRAFTS / f"{did}.json"
        if not f.exists():
            continue
        d = json.loads(f.read_text())
        b_st, b_bx = _inv_totals(iid)
        before = invoice_costs(iid)
        r = correct_invoice(iid, d)
        f.unlink(missing_ok=True)
        a_st, a_bx = _inv_totals(iid)
        done.append((iid, before, r))
        txt += f"✅ {r['farm']} № {r['invoice_no']}: стеблей {b_st:g} → {a_st:g}, коробок {b_bx:g} → {a_bx:g}\n"
        if r.get("awb_key"):
            keys.add(r["awb_key"])
    txt += "Оплата, пополнение, MAWB и баланс фермы — как были."
    from . import kbreak
    for k in keys:
        edited = await refresh_shipment(k, "исправлен инвойс (" + ", ".join(r["farm"] for _i, _b, r in done) + ")")
        if edited:
            txt += f"\n🔄 Пакинг MAWB {k} в чатах обновлён ({edited} сообщ.)."
        elif kbreak.is_sent(k):
            txt += f"\n⚠️ Пакинг MAWB {k} уже ушёл старым сообщением — нажми «📦 Пакинги в чаты» → этот MAWB."
    await c.message.edit_text(txt[:4000], reply_markup=_fix_kb(p["up"]))
    for iid, before, r in done:
        if r["changes"]:
            await send_correction(iid, before, f"Исправление инвойса фермы № {r['invoice_no']}")
    await send_packing_lists()


@dp.callback_query(F.data.startswith("fx:"), wr)
async def fix_button(c: CallbackQuery):
    up = int(c.data.split(":")[1])
    FIXWAIT[c.from_user.id] = up
    await c.answer()
    await bot.send_message(c.from_user.id,
                           "✏️ Напиши одним сообщением, что не так. Например:\n"
                           "• «там 840 стеблей и 24 коробки»\n"
                           "• «не хватает коробки: Freedom 50 см 300 ст»\n"
                           "• «Explorer не 60, а 70 см»\n"
                           "Перечитаю инвойс с этим и покажу, что изменилось. /cancel — передумал.")


@dp.message(pv, wr, F.from_user.id.func(lambda i: i in FIXWAIT),
            F.text.func(lambda t: not t.startswith("/") or t.startswith("/cancel")))
async def fix_text(m: Message):
    up = FIXWAIT.pop(m.from_user.id)
    if m.text.strip().lower() in ("/cancel", "отмена", "отменить"):
        await m.answer("Ок, ничего не меняю.")
        return
    await _apply_fix(m.from_user.id, up, m.text)


@dp.message(pv, wr, F.reply_to_message.func(lambda r: r is not None), F.text,
            F.func(lambda m: (m.from_user.id, m.reply_to_message.message_id) in MSG2UP and not (m.text or "").startswith("/")
                   and not any(money_from_text(m.text or "")) and not ai.find_mawb(m.text or "")))
async def fix_reply(m: Message):
    await _apply_fix(m.from_user.id, MSG2UP[(m.from_user.id, m.reply_to_message.message_id)], m.text)


@dp.message(pv, wr, Command("fix"))
async def fix_cmd(m: Message):
    """/fix там 840 стеблей и 24 коробки — исправить последний присланный инвойс."""
    from .models import Upload
    uid = m.from_user.id
    up = LASTUP.get(uid)
    if not up:
        with session() as s:
            u = s.exec(select(Upload).where(Upload.uid == uid, Upload.kind == "farm_invoice")
                       .order_by(Upload.id.desc())).first()
            up = u.id if u else None
    if not up:
        await m.answer("Не нашёл твоего последнего инвойса. Пришли файл, потом нажми «✏️ Исправить текстом».")
        return
    text = (m.text or "").split(maxsplit=1)
    if len(text) < 2:
        FIXWAIT[uid] = up
        await m.answer("✏️ Напиши, что не так в последнем инвойсе (например «там 840 стеблей и 24 коробки»).")
        return
    await _apply_fix(uid, up, text[1])

CORR: dict = {}       # user -> pending «is it a correction?»
NEWFARM: dict = {}    # user -> invoice from a farm we don't know yet


def _clean_farm_name(raw: str) -> str:
    """«FLORA DELIGHT LTD» -> «Flora Delight»."""
    n = re.sub(r"[,.]?\s*\b(ltd|limited|s\.?a\.?s\.?|s\.?a\.?|llc|inc|b\.?v\.?|co)\b\.?", "", raw, flags=re.I).strip(" ,.-")
    return " ".join(w.capitalize() if w.isupper() or w.islower() else w for w in n.split()) or raw.strip()


@dp.callback_query(F.data.startswith("nf:"), wr)
async def new_farm_answer(c: CallbackQuery):
    from .api import resolve_farm
    p = NEWFARM.pop(c.from_user.id, None)
    await c.answer()
    if not p:
        await c.message.edit_text("Устарело — пришли инвойс ещё раз.")
        return
    country = c.data.split(":", 1)[1]
    head = p["head"]
    if country != "skip" and not p.get("country"):
        p["country"] = country
        NEWFARM[c.from_user.id] = p
        await c.message.edit_text(f"🌱 «{_clean_farm_name(p['raw'])}» ({country}) — как закупаем?",
                                  reply_markup=_kb([[("Напрямую у фермы", f"nf:{country}")],
                                                    [("Через посредника (BiFlorica)", "nfb:" + country)]]))
        return
    if country != "skip":
        name = _clean_farm_name(p["raw"])
        with session() as s:
            f = resolve_farm(s, name, create_country=country)
            if p["raw"].upper() not in (f.aliases or "").upper():
                f.aliases = ",".join(x for x in [f.aliases, p["raw"].upper()] if x)
            s.add(f); s.commit()
            name = f.name
        for d in p["subs"]:
            if not d.get("farm"):
                d["farm"], d["country"] = name, d.get("country") or country
                d["warnings"] = [w for w in d.get("warnings", []) if "не найден" not in w]
        head = f"🌱 Ферма «{name}» ({country}) добавлена в плантации — закупка напрямую.\n\n" + _flow_head(p["subs"])
    await _start_flow(c.from_user.id, p["subs"], p["out"], p["caption"], c.message, head)


@dp.callback_query(F.data.startswith("nfb:"), wr)
async def new_farm_broker(c: CallbackQuery):
    """New farm bought through the broker: same scheme as Tessa/Plazoleta — money from the BiFlorica statement,
    this invoice only gives the packing list."""
    from .api import resolve_farm, attach_broker_packing
    p = NEWFARM.pop(c.from_user.id, None)
    await c.answer()
    if not p:
        await c.message.edit_text("Устарело — пришли инвойс ещё раз.")
        return
    country = c.data.split(":", 1)[1]
    name = _clean_farm_name(p["raw"])
    with session() as s:
        f = resolve_farm(s, name, create_country=country)
        if p["raw"].upper() not in (f.aliases or "").upper():
            f.aliases = ",".join(x for x in [f.aliases, p["raw"].upper()] if x)
        f.in_fee_pct, f.markup_pct, f.account = f.in_fee_pct or 3, f.markup_pct or 7, "Брокер"
        s.add(f); s.commit()
        name = f.name
    msgs = []
    for d in p["subs"]:
        if not d.get("farm"):
            d["farm"], d["country"] = name, d.get("country") or country
        fill_mawb(d, None)
        msgs.append(attach_broker_packing(d))
    await c.message.edit_text(f"🌱 Ферма «{name}» ({country}) добавлена — закупка через посредника "
                              "(−3% на входе, комиссия по выписке BiFlorica).\n\n" + "\n\n".join(msgs))
    await send_packing_lists()
BOXWAIT: dict = {}    # user -> invoice id waiting for its box breakdown file


@dp.callback_query(F.data.startswith("bx:"), wr)
async def box_detail_answer(c: CallbackQuery):
    _, what, iid = c.data.split(":")
    await c.answer()
    if what == "yes":
        BOXWAIT[c.from_user.id] = int(iid)
        await c.message.edit_text("Жду детализацию по коробкам: пришли файл, фото или PDF (packing list фермы, "
                                  "раскладку по коробкам). Следующий документ от тебя возьму как детализацию к этому инвойсу.")
    else:
        await c.message.edit_text("Ок, без детализации. В пакинге коробки будут по строкам инвойса.")


async def _start_flow(uid: int, subs: list, out: dict, caption: str, note, head: str):
    t = topup_from_text(caption)
    if out.get("doc_type") == "farm_invoice":
        for d in subs:
            fill_mawb(d, t.id if t else None)     # MAWB from the Expolanka breakdown right away
    cu, _cr = money_from_text(caption)
    flow = {"ids": [save_draft(d) for d in subs], "paid": None, "topup": t.id if t else None,
            "pays": None, "kind": out.get("doc_type"), "ts": time.time(), "note": note,
            "up": out.get("upload_id") if out.get("doc_type") == "farm_invoice" else None}
    if "не оплач" in caption.lower():
        flow["paid"] = False
    elif cu or t:
        flow["paid"] = True                       # sums or a top-up date in the caption = paid
    if cu:
        flow["pays"] = payments_for(caption, subs)
        if flow["paid"] is False:                 # «не оплачен 1100$» = approximate $
            flow["approx"] = True
    FLOW[uid] = flow
    LAST_DRAFT[uid] = (flow["ids"], time.time())
    await _step(uid, head=head)


@dp.callback_query(F.data.startswith("cor:"), wr)
async def correction_answer(c: CallbackQuery):
    from .api import correct_invoice, invoice_costs
    p = CORR.pop(c.from_user.id, None)
    await c.answer()
    if not p:
        await c.message.edit_text("Устарело — пришли инвойс ещё раз.")
        return
    if c.data == "cor:no":
        await _start_flow(c.from_user.id, p["subs"], p["out"], p["caption"], c.message, p["head"])
        return
    txt, done = "", []
    for did, inv_id in p["pairs"]:
        d = json.loads((DRAFTS / f"{did}.json").read_text())
        before = invoice_costs(inv_id)
        r = correct_invoice(inv_id, d)
        (DRAFTS / f"{did}.json").unlink(missing_ok=True)
        done.append((inv_id, before, r))
        txt += (f"✅ Корректировка внесена: {r['farm']} № {r['invoice_no']}.\n"
                + ("\n".join("• " + x for x in r["changes"]) or "Строки не изменились, обновлена разбивка по коробкам.") + "\n\n")
    txt += "Оплата, пополнение, MAWB и баланс фермы — как были."
    edited = 0
    if r.get("awb_key"):
        edited = await refresh_shipment(r["awb_key"], f"{r['farm']} — корректировка инвойса № {r['invoice_no']}")
        from . import kbreak
        if edited:
            txt += f"\n\n🔄 Пакинг в чатах ОБНОВЛЁН (отредактировал {edited} сообщ.), новый файл не слал."
        elif kbreak.is_sent(r["awb_key"]):
            txt += ("\n\n⚠️ Пакинг этой поставки ушёл в чаты до обновления бота — то сообщение я отредактировать не могу. "
                    "Нажми «📦 Пакинги в чаты» → этот MAWB: уйдёт исправленный. Дальше буду править старые сообщения сам.")
        else:
            txt += "\n\nВ чаты эта поставка ещё не уходила — уйдёт уже исправленной."
    await c.message.edit_text(txt[:4000])
    for inv_id, before, r in done:
        if r["changes"]:
            await send_correction(inv_id, before, f"Корректировка инвойса фермы № {r['invoice_no']}")


# ---------- the «оплачен? → откуп → $» conversation --------------------------------------
FLOW: dict[int, dict] = {}   # user -> current document in progress
_CONS_WAIT: dict = {}        # user -> parsed weight report waiting for its MAWB


def _flow_docs(flow):
    docs = []
    for i in flow["ids"]:
        p = DRAFTS / f"{i}.json"
        if p.exists():
            docs.append((p, json.loads(p.read_text())))
    return docs


def _flow_head(subs) -> str:
    if len(subs) > 1:
        return (f"📄 Общий инвойс {subs[0].get('invoice_no') or ''} на {len(subs)} плантации — делю:\n" + "\n".join(
            f"• {d['farm']}: {sum(l.get('stems') or 0 for l in d['lines']):g} ст, ${d['invoice_total_usd']:g}"
            + (f", MAWB {d['awb']}" if d.get("awb") else ", MAWB ?") for d in subs))
    return _doc_head(subs[0])


def _kb(rows):
    return InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text=t, callback_data=d) for t, d in r] for r in rows])


def _kbb(rows=None):
    """Flow keyboard: always has «↩️ Назад» (undo the last answer)."""
    return _kb((rows or []) + [[("↩️ Назад", "fl:back")]])


RESET = {"paid": None, "topup": None, "awb_done": False, "pays": None, "farm": None, "bal_done": False}


def _mark(flow: dict, key: str):
    flow.setdefault("hist", []).append(key)


async def _send(uid: int, text: str, kb=None, note=None):
    if note is not None:
        try:
            await note.edit_text(text, reply_markup=kb)
            return
        except Exception:
            pass
    await bot.send_message(uid, text, reply_markup=kb)


async def _step(uid: int, head: str = ""):
    """Ask the next missing thing, or book when everything is known."""
    from .models import TopUp
    flow = FLOW.get(uid)
    if not flow:
        return
    docs = _flow_docs(flow)
    if not docs:
        FLOW.pop(uid, None)
        return
    note, flow["note"] = flow.get("note"), None
    head = (head + "\n\n") if head else ""
    freight = flow["kind"] == "freight_invoice"
    if flow["paid"] is None:
        await _send(uid, head + "Оплачен?", _kb([[("✅ Оплачен", "fl:paid"), ("🚚 Не оплачен — в пути", "fl:unpaid")]] + ([[("↩️ Назад", "fl:back")]] if flow.get("hist") else []) + _fix_rows(flow.get("up")) + [[("✖️ Отменить (останется в черновиках)", "fl:cancel")]]), note)
        return
    if flow["paid"] and not flow["topup"]:
        with session() as s:
            tops = s.exec(select(TopUp).order_by(TopUp.id.desc()).limit(8)).all()
        if not tops:
            await _send(uid, head + "Нет ни одного пополнения — создай его (скрин покупки USDT или «Учёт»).", None, note)
            return
        rows = [[(f"{x.date} · курс {x.rub / x.usd:.2f}", f"fl:tp:{x.id}")] for x in tops]
        await _send(uid, head + "Из какого пополнения оплачен?", _kbb(rows), note)
        return
    if not flow.get("awb_done"):
        for p, d in docs:
            if not d.get("awb") and d.get("doc_type") == "farm_invoice":
                fill_mawb(d, flow.get("topup"))
                if d.get("awb"):
                    p.write_text(json.dumps(d, ensure_ascii=False))
        awbs = sorted({d.get("awb") for _p, d in docs if d.get("awb")})
        if awbs:
            await _send(uid, head + f"MAWB: {', '.join(awbs)}" + (" (из разбивки)" if any(d.get('mawb_note') for _p, d in docs) else "") + " — верно?",
                        _kbb([[("✅ Верно", "fl:awb:ok"), ("✏️ Другой", "fl:awb:edit")], [("Ещё не знаем", "fl:awb:none")]]), note)
        else:
            await _send(uid, head + "Номер MAWB уже известен?",
                        _kbb([[("✏️ Да, напишу", "fl:awb:edit"), ("Ещё не знаем", "fl:awb:none")]]), note)
        return
    if freight and not flow["pays"]:
        usd = (docs[0][1].get("freight") or {}).get("total_usd")
        if usd:
            flow["pays"] = [(usd, None)]           # the bill's own $
    if not flow["pays"]:
        many = len(docs) > 1
        if flow["paid"]:
            t = None
            with session() as s:
                t = s.get(TopUp, flow["topup"])
            ask = (f"Сколько $ оплатили{' — по строке на ферму' if many else ''}? ₽ посчитаю по курсу "
                   f"{t.date} ({t.rub / t.usd:.2f}). Если списали иначе — добавь ₽.")
        else:
            ask = (f"Примерно сколько $ будет оплата{' — по строке на ферму' if many else ''}? "
                   f"Себестоимость посчитаю заранее (±10%).")
        ex = "\n".join(f"{d['farm']} {round(d.get('invoice_total_usd') or 1000)}$" for _p, d in docs) if many else \
            f"{round(docs[0][1].get('invoice_total_usd') or 1000)}$"
        await _send(uid, head + ask + f"\nНапример:\n{ex}", _kbb(), note)
        return
    if flow["paid"] and not freight and flow.get("farm") is None:
        from .api import farm_balance_text
        tot = [round(d.get("invoice_total_usd") or sum((l.get("stems") or 0) * (l.get("price_usd") or 0) for l in d.get("lines", [])), 2)
               for _p, d in docs]
        bal = "\n".join(farm_balance_text(d.get("farm") or "") for _p, d in docs)
        label = f"Ровно по инвойсу ${tot[0]:g}" if len(docs) == 1 else "Ровно по инвойсам"
        from . import calc
        rule = calc.RULES.get((docs[0][1].get("farm") or "").strip().lower(), {})
        if rule.get("in_fee") and len(docs) == 1:          # broker: 97 % of the dollars reach the exchange
            paid_usd = (flow["pays"][0] or (0, 0))[0] if flow.get("pays") else 0
            label = f"Как обычно: на биржу ${paid_usd:g} − {rule['in_fee']:g}% = ${paid_usd * (1 - rule['in_fee'] / 100):,.2f}".replace(",", " ")
            bal += f"\nℹ️ {docs[0][1].get('farm')}: закупка через брокера — к инвойсу +{rule.get('markup', 0):g}%, на биржу заходит {100 - rule['in_fee']:g}% долларов."
        await _send(uid, head + bal + "\n\nСколько дошло до фермы?",
                    _kbb([[(label, "fl:farm:exact")], [("Другая сумма (переплата / аванс / недоплата)", "fl:farm:other")],
                          [("✏️ У фермы был другой баланс", "fl:bal:edit")]]), note)
        return
    if not flow["paid"] and not freight and not flow.get("bal_done"):
        from .api import farm_balance_projection
        lines = [farm_balance_projection(d.get("farm") or "", d.get("invoice_total_usd") or
                                         sum((l.get("stems") or 0) * (l.get("price_usd") or 0) for l in d.get("lines", [])))
                 for _p, d in docs]
        await _send(uid, head + "\n".join(lines) + "\n\nВерно?",
                    _kbb([[("✅ Верно", "fl:bal:ok"), ("✏️ Другой баланс", "fl:bal:edit")]]), note)
        return
    await _book_flow(uid, head, note)


async def _book_flow(uid: int, head: str, note):
    from .models import TopUp
    flow = FLOW.pop(uid, None)
    docs = _flow_docs(flow)
    t = None
    if flow["paid"]:
        with session() as s:
            t = s.get(TopUp, flow["topup"])
    pays = flow["pays"] + [None] * (len(docs) - len(flow["pays"]))
    farm_usd = (flow.get("farm") or []) + [None] * len(docs)
    out, ask_boxes = [], []
    for n, ((p, d), pr) in enumerate(zip(docs, pays)):
        usd, rub = pr or (None, None)
        if d.get("doc_type") == "farm_invoice":
            fill_mawb(d, t.id if t else None)
        snap, kind = book_document(d, t.id if t else 0, usd, rub, uid, paid=bool(flow["paid"]), farm_usd=farm_usd[n])
        if not snap:
            out.append(f"❌ {d.get('farm') or ''}: не внёс — {kind}. Черновик в «Учёт».")
            continue
        p.unlink(missing_ok=True)
        pre = f"{d['farm']}: " if len(docs) > 1 else ""
        out.append(pre + (_booked_text(snap, d, kind, t, uid) if t else _transit_text(snap, d, kind)))
        if d.get("doc_type") == "farm_invoice":
            from .api import farm_balance_text, find_invoice_id, needs_box_detail
            out[-1] += "\n" + farm_balance_text(d.get("farm") or "")
            if needs_box_detail(d):
                iid = find_invoice_id(d)
                if iid:
                    ask_boxes.append((iid, d.get("farm") or ""))
    await _send(uid, head + "\n\n".join(out), None, note)
    for iid, farm in ask_boxes:
        await bot.send_message(uid, f"📦 В инвойсе {farm} НЕТ детализации по коробкам (что в какой коробке).\n"
                                    "Помоги складу — скинешь детализацию по коробкам?",
                               reply_markup=_kb([[("📎 Да, сейчас скину", f"bx:yes:{iid}")],
                                                 [("Нет, без неё", f"bx:no:{iid}")]]))
    await _flush_export(uid, "📊 Логистика легла на товар из прошлого пополнения.")


def _transit_text(v: dict, d: dict, kind: str) -> str:
    fmt = lambda x: f"{x:,.0f}".replace(",", " ")
    if kind == "freight":
        lg = max(v["unpaid_logistics"], key=lambda x: x["id"]) if v["unpaid_logistics"] else None
        return (f"🚚 Фрахт в пути, не оплачен: ${lg['usd']:g} ≈ {fmt(lg['rub_est'])} ₽ по курсу {v['est_rate']:.2f} "
                f"(последнее пополнение {v['est_topup']})") if lg else "🚚 Фрахт внесён как неоплаченный"
    inv = max(v["invoices"], key=lambda x: x["id"])
    return (f"🚚 В пути, не оплачен: ≈${inv['est_usd']:g} ≈ {fmt(inv['rub_goods'])} ₽ (курс {v['est_rate']:.2f}, "
            f"пополнение {v['est_topup']})\n"
            f"≈ Себестоимость {inv['avg_rub']:.2f} ₽/ст ±10% (логистика: {inv['logi_source']})\n"
            "Оплатишь — нажми «Оплачен» в «В пути» или в напоминании во вторник/среду.")


@dp.callback_query(F.data.startswith("fl:"), wr)
async def flow_button(c: CallbackQuery):
    flow = FLOW.get(c.from_user.id)
    if not flow:
        await c.answer("Документ уже обработан или устарел", show_alert=True)
        return
    part = c.data.split(":")
    if part[1] == "back":
        for w in ("awb_wait", "farm_wait", "bal_wait"):
            flow[w] = False
        if flow.get("hist"):
            flow[flow["hist"].pop()] = None
            for k, v in RESET.items():
                if flow.get(k) is None:
                    flow[k] = v
    elif part[1] == "cancel":
        FLOW.pop(c.from_user.id, None)
        await c.answer()
        await c.message.edit_text("Отменено. Документ лежит в «Учёт» → Черновики, можно разобрать позже.")
        return
    elif part[1] == "paid":
        flow["paid"] = True; _mark(flow, "paid")
    elif part[1] == "unpaid":
        flow["paid"] = False; _mark(flow, "paid")
    elif part[1] == "tp":
        flow["topup"] = int(part[2]); _mark(flow, "topup")
    elif part[1] == "bal" and part[2] == "ok":
        flow["bal_done"] = True; _mark(flow, "bal_done")
    elif part[1] == "bal" and part[2] == "edit":
        flow["bal_wait"] = True
        await c.answer()
        many = len(flow["ids"]) > 1
        await c.message.edit_text("Напиши баланс фермы ДО этой поставки: `+48` — у фермы аванс $48, `-100` — мы должны $100, `0` — ровно."
                                  + ("\nНесколько ферм — по строке: `Agriflora +48`" if many else ""),
                                  parse_mode="Markdown", reply_markup=_kbb())
        return
    elif part[1] == "awb" and part[2] in ("ok", "none"):
        flow["awb_done"] = True; _mark(flow, "awb_done")
        if part[2] == "none":
            for p, d in _flow_docs(flow):                   # «ещё не знаем» -> no guessed MAWB either
                d["awb"], d["mawb_note"], d["weight_kg"] = None, None, None
                p.write_text(json.dumps(d, ensure_ascii=False))
    elif part[1] == "awb" and part[2] == "edit":
        flow["awb_wait"] = True
        await c.answer()
        await c.message.edit_text("Напиши номер MAWB, например `074-48014901`.", parse_mode="Markdown", reply_markup=_kbb())
        return
    elif part[1] == "farm" and part[2] == "exact":
        flow["farm"] = []; _mark(flow, "farm")              # [] = exactly the invoice for every farm
    elif part[1] == "farm" and part[2] == "other":
        flow["farm_wait"] = True
        await c.answer()
        many = len(flow["ids"]) > 1
        await c.message.edit_text("Напиши, сколько $ дошло до фермы" + (" — по строке на ферму" if many else "") +
                                  ", например `520$`. Остаток станет авансом у фермы, нехватка — долгом.",
                                  parse_mode="Markdown", reply_markup=_kbb())
        return
    flow["note"] = c.message
    await c.answer()
    await _step(c.from_user.id)


_AUTO_EXPORT: set = set()


async def _flush_export(uid: int, caption: str):
    """Freight paid from a new top-up for goods from an older one -> resend both sheets."""
    if _AUTO_EXPORT:
        ids = sorted(_AUTO_EXPORT)
        _AUTO_EXPORT.clear()
        try:
            await export_topups(ids, uid, caption)
        except Exception as e:
            await bot.send_message(uid, f"Не смог обновить Excel: {e}")


async def _book_or_draft(m: Message, note, out: dict, t, usd, rub):
    """Both sums known -> straight into the top-up. Otherwise a draft waiting for '1198$ 105472₽'."""
    warn = ("\n⚠️ " + "\n⚠️ ".join(out["warnings"])) if out.get("warnings") else ""
    head = _doc_head(out)
    if usd:
        snap, kind = book_document(out, t.id, usd, rub, m.from_user.id)
        if snap:
            await note.edit_text(head + warn + "\n\n" + _booked_text(snap, out, kind, t, m.from_user.id))
            await _flush_export(m.from_user.id, "📊 Логистика легла на товар из прошлого пополнения.")
            return
        warn += f"\n⚠️ Не внёс сразу: {kind}"
    did = save_draft(out)
    LAST_DRAFT[m.from_user.id] = ([did], time.time())
    await _edit(note, head + warn + f"\n\n→ пополнение {t.date}. Ответь суммой оплаты в $, например "
                      f"`{round(out.get('invoice_total_usd') or 1000)}$` — ₽ посчитаю по курсу пополнения "
                      f"({t.rub / t.usd:.2f}). Если списали иначе — добавь ₽: `1115$ 98000₽`.")


async def _book_or_draft_multi(m: Message, note, subs: list, t):
    """Trader invoice (NextWave): one file, several farms -> one invoice per farm, each paid separately."""
    for d in subs:
        d["topup_id"] = t.id
        fill_mawb(d, t.id)
    head = f"📄 Общий инвойс {subs[0].get('invoice_no') or ''} на {len(subs)} плантации — делю:\n" + "\n".join(
        f"• {d['farm']}: {sum(l.get('stems') or 0 for l in d['lines']):g} ст, ${d['invoice_total_usd']:g}"
        + (f", MAWB {d['awb']}" if d.get("awb") else ", MAWB ?") for d in subs)
    pays = payments_for(m.caption or "", subs)
    res, left = [], []
    for d, pr in zip(subs, pays or [None] * len(subs)):
        if pr and pr[0]:
            snap, kind = book_document(d, t.id, pr[0], pr[1], m.from_user.id)
            if snap:
                res.append(f"{d['farm']}: " + _booked_text(snap, d, kind, t, m.from_user.id).split("\n")[0].replace("✅ ", "✅ "))
                continue
        left.append(d)
    if left:
        ids = [save_draft(d) for d in left]
        LAST_DRAFT[m.from_user.id] = (ids, time.time())
        ex = "\n".join(f"{d['farm']} {round(d['invoice_total_usd'])}$" for d in left)
        res.append("Ответь оплатой в $ по каждой плантации, по строке на каждую (₽ посчитаю по курсу пополнения):\n`" + ex + "`\n"
                   "Одна сумма на всех — разделю пропорционально инвойсу.")
    await _edit(note, head + "\n\n" + "\n".join(res))


async def _edit(note, text: str):
    try:
        await note.edit_text(text, parse_mode="Markdown")
    except Exception:
        await note.edit_text(text.replace("`", ""))


def _doc_head(out: dict) -> str:
    if out.get("doc_type") == "freight_invoice":
        fr = out.get("freight") or {}
        return f"✈️ Фрахт · {fr.get('provider') or out.get('farm') or '?'} · MAWB {out.get('awb') or '?'}" + (
            f" · ${fr['total_usd']:g}" if fr.get("total_usd") else "")
    stems, boxes = _d_totals(out)
    return (f"📄 {out.get('farm') or '?'} · MAWB {out.get('awb') or '?'}\n"
            f"Строк: {len(out.get('lines', []))}, стеблей: {stems:g}, коробок: {boxes:g}, итог: ${out.get('invoice_total_usd') or '?'}"
            + (f"\n🔗 {out['mawb_note']}" if out.get("mawb_note") else ""))


def _booked_text(snap: dict, out: dict, kind: str, t, uid: int | None = None) -> str:
    fmt = lambda x: f"{x:,.0f}".replace(",", " ")
    if kind == "freight":
        lg = max((l for l in snap["logistics"] if l["topup_id"] == t.id), key=lambda l: l["id"])
        txt = (f"✅ Фрахт внесён в пополнение {t.date}: ${lg['usd'] or 0:g} / {fmt(lg['rub'] or 0)} ₽"
               + (f" · {lg['weight_kg']:g} кг по счёту" if lg.get("weight_kg") else "")
               + (f" · {lg['rub_per_kg']:.2f} ₽/кг" if lg.get("rub_per_kg") else "")
               + (f" · разбивка {lg['kg_total']:g} кг ✓" if lg.get("kg_total") else " · разбивки кг ещё нет"))
        spread = awb_spread(lg["awb"])
        if spread:
            col = "air" if lg["leg"] == "air" else "msk"
            txt += "\nЛегло на:\n" + "\n".join(
                f"• {x['farm']} (пополнение {x['topup']}): {x[col]:.2f} ₽/ст" for x in spread)
            other = sorted({x["topup"] for x in spread if x["topup_id"] != t.id})
            if other:
                txt += f"\n↪️ Товар оплачен из другого пополнения ({', '.join(other)}) — его себестоимость обновлена, Excel пришлю."
                _AUTO_EXPORT.add(t.id)
                _AUTO_EXPORT.update(x["topup_id"] for x in spread)
    else:
        inv = max(snap["invoices"], key=lambda i: i["id"])
        st = sum(l["stems"] for l in inv["lines"]) or 1
        flower = sum(l["price_rub"] * l["stems"] for l in inv["lines"]) / st
        logi = sum((l["air_rub"] + l["msk_rub"]) * l["stems"] for l in inv["lines"]) / st
        txt = (f"✅ Внесено в пополнение {t.date}: ${inv['usd_paid']:g} / {fmt(inv['rub_paid'])} ₽"
               + (f" (по курсу {t.rub / t.usd:.2f})" if out.get("_rub_auto") else "") + "\n"
               f"Себестоимость в среднем {flower + logi:.2f} ₽/стебель (цветок {flower:.2f} + логистика {logi:.2f})")
    rel = [w for w in snap["warnings"] if (out.get("farm") or "~").lower()[:5] in w.lower()
           or (out.get("awb") or "~").replace("-", "")[:6] in w.replace("-", "")]
    txt += ("\n⚠️ " + "\n⚠️ ".join(rel) if rel else "")
    if uid and roles.sees_money(roles.role_of(uid)):
        txt += f"\nОстаток пополнения ${money_of(t.id)['usd_left']:g}"
    return txt


def _register_topup(tp: dict, caption: str) -> str:
    """Crypto purchase screenshot -> new top-up dated today, and it becomes the active one."""
    from .models import TopUp
    rub = tp.get("rub")
    usd = tp.get("usd_withdrawn") or tp.get("usd_bought")     # what actually left for payments
    cu, cr = money_from_text(caption)                          # caption «1728.86$ 152000₽» overrides
    usd, rub = cu or usd, cr or rub
    if not (rub and usd):
        return "Не нашёл на скрине обе суммы (₽ и USDT). Пришли ещё раз или подпиши: «1728.86$ 152000₽»."
    order = (tp.get("order_no") or "").lstrip("#")
    with session() as s:
        if order:
            dup = s.exec(select(TopUp).where(TopUp.note.contains(f"#{order}"))).first()
            if dup:
                return f"Это пополнение уже внесено: {dup.date}, {_n(dup.rub, 2)} ₽ → ${_n(dup.usd, 4)} (заявка #{order})."
        note = f"заявка #{order}" if order else ""
        if tp.get("usd_withdrawn") and tp.get("usd_bought"):
            note += f"; куплено {tp['usd_bought']} USDT, выведено {tp['usd_withdrawn']}"
        t = TopUp(date=datetime.now().strftime("%d.%m.%Y"), rub=rub, usd=usd, note=note.strip("; "))
        s.add(t); s.commit(); s.refresh(t)
    set_active_topup(t.id)
    fee = ""
    if tp.get("usd_withdrawn") and tp.get("usd_bought") and not cu:
        fee = f"\n(куплено {tp['usd_bought']:g} USDT, за вывод ушло {tp['usd_bought'] - tp['usd_withdrawn']:.4f} — считаю по выведенным)"
    return (f"💰 Новое пополнение {t.date}\n{_n(rub, 2)} ₽ → ${_n(usd, 4)}\nКурс {rub / usd:.4f} ₽/$" + fee +
            "\n\nДокументы из чата теперь идут в него.")


def _n(x: float, d: int) -> str:
    return f"{x:,.{d}f}".replace(",", " ")


def _applied_text(w: dict) -> str:
    on = {i["farm"] for i in w["invoices"] if i.get("weight_kg")}
    missing = [f for f in w["farm_kg"] if not any(f.lower()[:4] == x.lower()[:4] for x in on)]
    txt = "✅ Разбивка сохранена"
    if on:
        txt += f", вес проставлен: {', '.join(sorted(on))}"
    if missing:
        txt += f"\nЕщё без инвойсов: {', '.join(missing)} — вес встанет сам, когда их внесёшь"
    return txt


@dp.message(pv, wr, F.document.mime_type.in_({"application/pdf", "image/jpeg", "image/png"}))
async def doc(m: Message):
    f = await bot.download(m.document)
    await _parse_and_reply(m, f.read(), m.document.mime_type)


@dp.message(pv, wr, F.photo)
async def photo(m: Message):
    f = await bot.download(m.photo[-1])
    await _parse_and_reply(m, f.read(), "image/jpeg")


@dp.channel_post(F.text.startswith("/id"))
async def channel_id(m: Message):
    """Post /id in the backup channel -> bot replies with the channel id for BACKUP_CHAT_ID."""
    await m.answer(f"BACKUP_CHAT_ID={m.chat.id}")


@dp.message(pv, wr, F.forward_origin.chat)
async def forwarded_from_channel(m: Message):
    """Forward any post from the channel to the bot -> it tells the channel id. Works even if
    the bot isn't admin yet (but it must be admin for backups to work)."""
    ch = m.forward_origin.chat
    await m.answer(f"Канал «{ch.title}»\nBACKUP_CHAT_ID={ch.id}\n\nВставь эту строку в .env и перезапусти. "
                   "Бот должен быть админом канала с правом публиковать и закреплять.")


@dp.message(pv, sysf, Command("backup"))
async def manual_backup(m: Message):
    from .backup import BACKUP_CHAT_ID, backup_now
    if not BACKUP_CHAT_ID:
        await m.answer("BACKUP_CHAT_ID не задан в .env")
        return
    await backup_now(bot, "вручную", force=True)
    await m.answer("💾 Бэкап отправлен в канал и закреплён")


@dp.message(pv, sysf, Command("backup_cleanup"))
async def backup_cleanup(m: Message):
    """Delete the pile of old backups in the channel, keep the newest 10."""
    from .backup import cleanup_channel
    note = await m.answer("🧹 Удаляю старые бэкапы из канала (оставлю 10 последних)…")
    n = await cleanup_channel(bot)
    await note.edit_text(f"🧹 Готово: удалено ~{n} старых сообщений, оставлены 10 последних бэкапов (закреплённый — последний).")


@dp.message(pv, ~ops)
async def stranger(m: Message):
    u = m.from_user
    await m.answer("Запрос на доступ отправлен администратору. Как только он выдаст роль — напиши /start.")
    who = f"{u.full_name}" + (f" (@{u.username})" if u.username else "") + f", ID {u.id}"
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="Супер-админ", callback_data=f"role:{u.id}:super"),
         InlineKeyboardButton(text="1С оператор", callback_data=f"role:{u.id}:viewer")],
        [InlineKeyboardButton(text="Системный супер-админ", callback_data=f"role:{u.id}:sys")],
        [InlineKeyboardButton(text="Отклонить", callback_data=f"role:{u.id}:no")]])
    for sid in roles.sys_ids():
        try:
            await bot.send_message(sid, f"🔑 Просит доступ: {who}", reply_markup=kb)
        except Exception:
            pass


@dp.callback_query(F.data.startswith("role:"), sysf)
async def give_role(c: CallbackQuery):
    _, uid, role = c.data.split(":")
    uid = int(uid)
    if role == "no":
        await c.message.edit_text(c.message.text + "\n❌ Отклонено")
        await c.answer()
        return
    name = c.message.text.split(": ", 1)[-1].split(",")[0]
    try:
        roles.set_role(uid, role, name)
    except ValueError as e:
        await c.answer(str(e), show_alert=True)
        return
    from .backup import mark_dirty
    mark_dirty()
    await c.message.edit_text(c.message.text + f"\n✅ Роль: {roles.ROLE_NAMES[role]}")
    try:
        await bot.send_message(uid, f"Доступ выдан: {roles.ROLE_NAMES[role]}. Нажми /start.")
    except Exception:
        pass
    await c.answer()


@dp.message(pv, sysf, Command("users"))
async def list_users_cmd(m: Message):
    lines = [f"• {u['name'] or '—'} · ID {u['tg_id']} · {u['role_name']}" for u in roles.users()]
    await m.answer("Пользователи:\n" + "\n".join(lines) + "\n\nМенять роли — «Учёт» → Пользователи.")


@dp.message(pv, ops, Command("excel"))
async def excel_cmd(m: Message):
    """Everyone (including 1С) can get the current report."""
    from .models import TopUp
    with session() as s:
        ids = [t.id for t in s.exec(select(TopUp)).all()]
    if not ids:
        await m.answer("Пока нет ни одного пополнения.")
        return
    await m.answer("Собираю учёт…")
    aud = "owner" if roles.can_write(roles.role_of(m.from_user.id)) else "operator"
    await export_topups(ids, m.from_user.id, "📊 Актуальный учёт", audience=aud)


@dp.message(pv, wr, F.document)
async def unknown_file(m: Message):
    await m.answer("Этот формат не читаю. Пришли инвойс как PDF, фото (jpg/png), Excel (.xlsx/.xls) или Word (.docx/.doc).")


@dp.message(pv, ops, F.document | F.photo)
async def viewer_upload(m: Message):
    await m.answer("У тебя роль «1С оператор» — только просмотр. Отчёт: /excel или «Учёт» → «Excel в чат».")


@dp.message(pv, wr, F.from_user.id.func(lambda i: bool(FLOW.get(i, {}).get("bal_wait"))), F.text)
async def balance_answer(m: Message):
    """Farm balance BEFORE this shipment, typed: «+48», «-100», «0», or «Agriflora +48» per line."""
    from .api import adjust_farm_balance
    flow = FLOW[m.from_user.id]
    docs = _flow_docs(flow)
    vals = re.findall(r"([A-Za-zА-Яа-яЁё][^\n+\-\d]*)?\s*([+\-−]?\s*\d+(?:[.,]\d+)?)", m.text)
    if not vals:
        await m.answer("Напиши число: `+48` (аванс у фермы), `-100` (наш долг) или `0`.", parse_mode="Markdown")
        return
    done = []
    for n, (name, num) in enumerate(vals):
        v = float(num.replace(" ", "").replace("−", "-").replace(",", "."))
        name = (name or "").strip()
        d = next((d for _p, d in docs if name and (d.get("farm") or "").lower().startswith(name.lower()[:4])), None)
        d = d or (docs[n][1] if n < len(docs) else docs[0][1])
        adjust_farm_balance(d.get("farm") or "", v)
        done.append(f"{d.get('farm')}: {'аванс' if v > 0 else 'долг' if v < 0 else 'ровно'} ${abs(v):g}")
    flow["bal_wait"] = False
    if not flow["paid"]:
        flow["bal_done"] = True
        _mark(flow, "bal_done")
    await _step(m.from_user.id, head="✓ Баланс до поставки: " + "; ".join(done))


@dp.message(pv, wr, F.text.func(lambda t: bool(re.search(r"(?:9\d|1[0-4]\d)\s*[.,x×*]\s*\d{2}\s*[.,x×*]\s*\d{2}(?!\d)", t or ""))))
async def weights_message(m: Message):
    """Expolanka weights from WhatsApp ('Zeeflora - 11 - 100.48.25 …') -> volumetric kg per farm."""
    from .api import learn_dims, volumetric_breakdown
    rows = volumetric_breakdown(m.text)
    if not rows:
        await m.answer("Не понял размеры. Формат как у Expolanka: `Zeeflora - 11 - 100.48.25`", parse_mode="Markdown")
        return
    learn_dims([r for r in rows if r["farm"] != "?"])
    fmt = lambda x: f"{x:g}"
    lines = []
    for r in rows:
        parts = " + ".join(f"{n}×{d} ({fmt(k)} кг)" + (" — обычная коробка" if len(rest) else "")
                           for n, d, k, *rest in r["dims"])
        lines.append(f"• {r['farm']}: {r['boxes']} кор. = {fmt(r['kg'])} кг  [{parts}]")
    total = sum(r["kg"] for r in rows)
    txt = "📦 Объёмный вес (Д×Ш×В / 6000):\n" + "\n".join(lines) + f"\nИтого {fmt(round(total, 1))} кг"
    if any(r["farm"] == "?" for r in rows):
        txt += "\n⚠️ У одной строки нет фермы — допиши имя и пришли ещё раз"
    per = [{"farm": r["farm"], "kg": round(r["kg"], 1), "boxes": r["boxes"]} for r in rows if r["farm"] != "?"]
    awb = ai.find_mawb(m.text)
    if awb:
        await m.answer(txt + "\n\n" + _applied_text(store_breakdown(awb, per)))
        return
    did = save_draft({"doc_type": "kg_breakdown", "awb": None, "per_farm_kg": per, "lines": [], "warnings": [],
                      "source": "volumetric"})
    LAST_DRAFT[m.from_user.id] = ([did], time.time())
    await m.answer(txt + "\n\nНапиши MAWB этой отправки — сохраню как разбивку.")


@dp.message(pv, wr, F.text.func(lambda t: any(money_from_text(t)) and not ai.find_mawb(t)))
async def money_followup(m: Message):
    flow = FLOW.get(m.from_user.id)
    if flow and time.time() - flow["ts"] < 60 * 60:
        docs = _flow_docs(flow)
        pays = payments_for(m.text, [d for _p, d in docs]) or []
        if flow.get("farm_wait"):
            if not any(pr and pr[0] for pr in pays):
                await m.answer("Нужна сумма в $, например `520$`", parse_mode="Markdown")
                return
            flow["farm"], flow["farm_wait"] = [pr[0] if pr else None for pr in pays], False
            _mark(flow, "farm")
            await _step(m.from_user.id)
            return
        if not any(pr and pr[0] for pr in pays):
            await m.answer("Не вижу суммы в $ — напиши, например, `1115$`", parse_mode="Markdown")
            return
        flow["pays"] = pays
        _mark(flow, "pays")
        await _step(m.from_user.id)
        return
    pay = PAY.get(m.from_user.id)
    if pay and pay.get("topup"):
        usd, rub = money_from_text(m.text)
        if not usd:
            await m.answer("Нужна сумма в $, например `1115$`", parse_mode="Markdown")
            return
        await _finish_pay(m.from_user.id, usd, rub)
        return
    await _money_followup_legacy(m)


async def _money_followup_legacy(m: Message):
    """'1198$ 105472₽' right after a document -> books the last draft(s).
    For a split trader invoice: one line per farm ('Agriflora 301$ 26500₽'), or one sum for all."""
    ids, ts = LAST_DRAFT.get(m.from_user.id, ([], 0))
    paths = [DRAFTS / f"{i}.json" for i in ids if (DRAFTS / f"{i}.json").exists()]
    if not paths or time.time() - ts > 30 * 60:
        await m.answer("Не к чему привязать суммы: сначала пришли документ (суммы можно прямо в подписи).")
        return
    docs = [json.loads(p.read_text()) for p in paths]
    pays = payments_for(m.text, docs) or []
    from .models import TopUp
    out, left = [], []
    for p, d, pr in zip(paths, docs, pays + [None] * (len(docs) - len(pays))):
        usd, rub = pr or (None, None)
        if d.get("doc_type") == "freight_invoice" and usd is None:
            usd = (d.get("freight") or {}).get("total_usd")
        if not usd:
            left.append(d.get("farm") or "документ")
            continue
        with session() as s:
            t = s.get(TopUp, d.get("topup_id") or 0)
        t = t or active_topup()
        snap, kind = book_document(d, t.id, usd, rub, m.from_user.id)
        if not snap:
            out.append(f"Не внёс {d.get('farm') or ''}: {kind}")
            continue
        p.unlink()
        out.append((f"{d['farm']}: " if len(docs) > 1 else "") + _booked_text(snap, d, kind, t, m.from_user.id))
        await _flush_export(m.from_user.id, "📊 Логистика легла на товар из прошлого пополнения.")
    if left:
        out.append("Без сумм остались: " + ", ".join(left) + " — пришли `$` для них.")
        LAST_DRAFT[m.from_user.id] = ([p.stem for p in paths if p.exists()], time.time())
    await m.answer("\n\n".join(out), parse_mode="Markdown")


@dp.message(pv, wr, Command("pop"))
async def choose_topup(m: Message):
    """Pick the top-up that chat documents go to."""
    from .models import TopUp
    act = active_topup()
    with session() as s:
        tops = s.exec(select(TopUp).order_by(TopUp.id.desc()).limit(8)).all()
    if not tops:
        await m.answer("Пополнений пока нет — создай в «Учёт».")
        return
    kb = InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(
        text=("✅ " if act and t.id == act.id else "") + f"{t.date} · ${t.usd:g}", callback_data=f"pop:{t.id}")] for t in tops])
    await m.answer("Куда вносить документы из чата:", reply_markup=kb)


@dp.callback_query(F.data.startswith("pop:"), wr)
async def set_topup(c: CallbackQuery):
    from .models import TopUp
    tid = int(c.data.split(":")[1])
    set_active_topup(tid)
    with session() as s:
        t = s.get(TopUp, tid)
    await c.message.edit_text(f"Документы из чата теперь идут в пополнение {t.date} ✅\n"
                              f"Разово другое — напиши дату в подписи: «23.09 1198$ 105472₽»")
    await c.answer()


@dp.message(pv, wr, F.text.func(lambda t: bool(ai.find_mawb(t))))
async def mawb_followup(m: Message):
    """MAWB sent as a separate message right after a document -> goes into that draft.
    MAWB + 'Farm kg' pairs in one message -> a kg breakdown for that MAWB."""
    info = _CONS_WAIT.pop(m.from_user.id, None)
    if info:
        await _consolidation(m, ai.find_mawb(m.text), info["country"], info["rows"], info["eta"])
        return
    flow = FLOW.get(m.from_user.id)
    if flow and flow.get("awb_wait"):
        awb = ai.find_mawb(m.text)
        for p, d in _flow_docs(flow):
            d["awb"], d["mawb_note"] = awb, None
            d["weight_kg"] = None
            p.write_text(json.dumps(d, ensure_ascii=False))
        flow["awb_wait"], flow["awb_done"] = False, True
        _mark(flow, "awb_done")
        await _step(m.from_user.id, head=f"MAWB {awb} ✓")
        return
    awb_t, pairs = breakdown_from_text(m.text)
    if len(pairs) >= 1:
        w = store_breakdown(awb_t, pairs)
        await m.answer(f"⚖️ Разбивка MAWB {awb_t}: " + ", ".join(f"{p['farm']} {p['kg']:g} кг" for p in pairs)
                       + "\n" + _applied_text(w))
        return
    ids, ts = LAST_DRAFT.get(m.from_user.id, ([], 0))
    did = ids[0] if ids else None
    path = DRAFTS / f"{did}.json" if did else None
    if not path or not path.exists() or time.time() - ts > 15 * 60:
        await m.answer("Не к чему привязать этот MAWB: сначала пришли документ (можно MAWB прямо в подписи к нему).")
        return
    d = json.loads(path.read_text())
    if d.get("doc_type") == "kg_breakdown":
        w = store_breakdown(ai.find_mawb(m.text), d.get("per_farm_kg"), d.get("source_file"))
        path.unlink()
        await m.answer(f"MAWB {ai.find_mawb(m.text)}\n" + _applied_text(w))
        return
    old, d["awb"] = d.get("awb"), ai.find_mawb(m.text)
    d["warnings"] = [w for w in d.get("warnings", []) if "MAWB" not in w and "awb" not in w.lower()]
    path.write_text(json.dumps(d, ensure_ascii=False))
    from .backup import mark_dirty
    mark_dirty()
    await m.answer(f"MAWB {d['awb']} записан в черновик #{did}" + (f" (было {old})" if old and old != d["awb"] else ""))


# ---------- payments from the reminder / «В пути» -----------------------------------------
PAY: dict[int, dict] = {}    # user -> {"type": "i"|"l", "id": .., "topup": ..}


async def _ask_pay_topup(uid: int, kind: str, oid: int, note=None):
    from .models import TopUp
    PAY[uid] = {"type": kind, "id": oid, "topup": None}
    with session() as s:
        tops = s.exec(select(TopUp).order_by(TopUp.id.desc()).limit(8)).all()
    rows = [[(f"{x.date} · курс {x.rub / x.usd:.2f}", f"pt:{x.id}")] for x in tops]
    await _send(uid, "Из какого пополнения оплачен?", _kb(rows), note)


@dp.callback_query(F.data.startswith("pay:"), wr)
async def pay_button(c: CallbackQuery):
    _, kind, oid = c.data.split(":")
    await c.answer()
    await _ask_pay_topup(c.from_user.id, kind, int(oid))


@dp.callback_query(F.data.startswith("pt:"), wr)
async def pay_topup_button(c: CallbackQuery):
    from .models import Invoice, Logistics, TopUp
    pay = PAY.get(c.from_user.id)
    if not pay:
        await c.answer("Устарело — нажми «Оплачен» ещё раз", show_alert=True)
        return
    pay["topup"] = int(c.data.split(":")[1])
    await c.answer()
    with session() as s:
        t = s.get(TopUp, pay["topup"])
        o = s.get(Invoice if pay["type"] == "i" else Logistics, pay["id"])
        hint = (o.est_usd if pay["type"] == "i" else o.usd) if o else None
        name = (o.farm if pay["type"] == "i" else o.provider) if o else "?"
    kb = _kb([[(f"Как было: ${hint:g}", f"pu:{hint}")]]) if hint else None
    await c.message.edit_text(f"{name} → пополнение {t.date}. Сколько $ оплатили? ₽ посчитаю по курсу {t.rub / t.usd:.2f}"
                              f" (если списали иначе — добавь ₽: `1115$ 98000₽`).", reply_markup=kb)


@dp.callback_query(F.data.startswith("pu:"), wr)
async def pay_same_usd(c: CallbackQuery):
    await c.answer()
    await _finish_pay(c.from_user.id, float(c.data.split(":")[1]), None, c.message)


async def _finish_pay(uid: int, usd: float, rub, note=None):
    from .api import pay_invoice, pay_logistics
    pay = PAY.pop(uid, None)
    if not pay:
        return
    try:
        r = pay_invoice(pay["id"], pay["topup"], usd, rub) if pay["type"] == "i" else \
            pay_logistics(pay["id"], pay["topup"], usd, rub)
    except ValueError as e:
        await _send(uid, f"Не вышло: {e}", None, note)
        return
    who = r.get("farm") or r.get("provider")
    await _send(uid, f"✅ {who}: оплачен из пополнения {r['topup']} — ${r['usd']:g} / {r['rub']:,.0f} ₽".replace(",", " ")
                + (" (по курсу)" if rub is None else "") + ". Себестоимость теперь точная.", None, note)
    if pay["type"] == "i":
        from .api import farm_balance_text
        await _send(uid, farm_balance_text(r.get("farm") or ""))
    _AUTO_EXPORT.add(pay["topup"])
    await _flush_export(uid, "📊 Инвойс оплачен — лист пополнения обновлён.")


# ---------- scheduler: arrivals + payment reminders (Tue 10:00 & Wed 10:00 MSK) ---------------
def _msk_now():
    from datetime import datetime, timedelta, timezone
    return datetime.now(timezone.utc).replace(tzinfo=None) + timedelta(hours=3)


def _money_people():
    return [u["tg_id"] for u in roles.users() if u["role"] in ("sys", "super")]


async def payment_reminder(uid_list=None):
    from .api import unpaid_items
    u = unpaid_items()
    fmt = lambda x: f"{x:,.0f}".replace(",", " ")
    if not u["invoices"] and not u["logistics"]:
        text, kb = "💸 Платежи: неоплаченных инвойсов нет.", None
    else:
        rate = u["rate"] or 0
        lines, rows, tot_usd, tot_rub = [], [], 0.0, 0.0
        for x in u["invoices"]:
            usd = x.get("est_usd") or 0
            lines.append(f"• {x['farm']} · инв {x.get('invoice_no') or '—'} · MAWB {x.get('awb') or '?'} · ≈${usd:g} ≈ {fmt(usd * rate)} ₽")
            rows.append([(f"✅ Оплачен: {x['farm']} ${usd:g}", f"pay:i:{x['id']}")])
            tot_usd += usd; tot_rub += usd * rate
        for x in u["logistics"]:
            usd = x.get("usd") or 0
            lines.append(f"• {x['provider'] or 'Фрахт'} · MAWB {x['awb']} · ${usd:g} ≈ {fmt(usd * rate)} ₽")
            rows.append([(f"✅ Оплачен: {x['provider'] or 'фрахт'} ${usd:g}", f"pay:l:{x['id']}")])
            tot_usd += usd; tot_rub += usd * rate
        text = (f"💸 Платежи на среду ({len(lines)}):\n" + "\n".join(lines) +
                f"\n\nИтого ≈ ${fmt(tot_usd)} ≈ {fmt(tot_rub)} ₽ (курс {rate:.2f}, последнее пополнение {u['topup']})")
        kb = _kb(rows[:20])
    for uid in uid_list or _money_people():
        try:
            await bot.send_message(uid, text, reply_markup=kb)
        except Exception as e:
            print(f"[lumen] reminder to {uid} failed: {e}", flush=True)


@dp.message(pv, wr, Command("mark"))
async def mark_cmd(m: Message):
    """/mark — show the default marking; /mark ABC — set it for new invoices."""
    from .api import SETTINGS, _settings, marking
    arg = (m.text or "").split(maxsplit=1)[1:] 
    if not arg:
        await m.answer(f"Маркировка по умолчанию: {marking()}\nСменить: `/mark НОВАЯ`", parse_mode="Markdown")
        return
    new = arg[0].strip().upper()
    SETTINGS.write_text(json.dumps({**_settings(), "marking": new}))
    from .backup import mark_dirty
    mark_dirty()
    await m.answer(f"✅ Новые инвойсы идут с маркировкой {new}. У отдельного инвойса её можно поменять в приложении.")


@dp.message(pv, wr, Command("packing"))
async def packing_cmd(m: Message):
    from . import packing
    ts = packing.targets()
    txt = ("Пакинг-листы уходят в:\n" + "\n".join(f"• {t['title'] or t['chat_id']}" + (" (тема)" if t.get('thread_id') else "")
                                                  for t in ts)) if ts else "Чаты для пакинг-листов не заданы."
    from . import reader
    who = "бота-читателя" if reader.RBOT else "бота-читателя (сейчас он выключен — нужен READER_BOT_TOKEN)"
    await m.answer(txt + f"\n\nПакинг-листы отправляет {who}, финансовый бот в чатах не нужен.\n"
                         "Добавить чат: добавь бота-читателя в группу и напиши там /packing_here "
                         "(в супергруппе — внутри нужной темы). Убрать: /packing_off там же.")


@dp.message(pv, sysf, Command("status"))
async def status_cmd(m: Message):
    """Why didn't a truck message reach the staff chat? Shows what the reader really sees and where it posts."""
    from . import reader, staffnotify, packing
    from .api import _settings
    ft = _settings().get("ft_chats") or []
    lines = [f"👀 Бот-читатель: {'включён' if reader.RBOT else 'ВЫКЛЮЧЕН (нет READER_BOT_TOKEN)'}",
             f"🚛 Чаты ТК МСК (читаю): {len(ft)} — " + (", ".join(map(str, ft)) or "нет — добавь читателя и подтверди «Да, это ТК МСК»"),
             "📣 Уведомления по грузам → " + (", ".join(f"{t['title']} [{t['chat_id']}" + (f" · тема {t['thread_id']}" if t.get('thread_id') else '') + "]"
                                                      for t in staffnotify.targets())
                                           or "НЕТ чатов — напиши /notify_here в чате склада"),
             "📦 Пакинги → " + (", ".join(t['title'] for t in packing.targets()) or "нет"),
             "", "Последнее, что читатель реально получил из групп:"]
    if reader.SEEN:
        for x in reader.SEEN[-8:]:
            lines.append(f"• {x['at']} «{x['chat']}» от {x['from']}{' (БОТ)' if x['bot'] else ''}: {x['text']}")
    else:
        lines.append("— ничего с последнего перезапуска")
    lines.append("\nℹ️ Telegram не показывает ботам сообщения других ботов. Если FloraMailing — бот, читатель его не видит: "
                 "тогда пересылай мне их сообщения или включим невидимого читателя-аккаунт (Telethon).")
    await m.answer("\n".join(lines)[:4000])


@dp.message(pv, wr, Command("awb"))
async def awb_export_cmd(m: Message):
    from aiogram.types import BufferedInputFile
    from .api import awb_payments_xlsx
    awb = ai.find_mawb(m.text or "")
    if not awb:
        await m.answer("Напиши так: /awb 074-48014901")
        return
    try:
        data, fname = awb_payments_xlsx(awb)
    except Exception as e:
        await m.answer(f"Не получилось: {getattr(e, 'detail', e)}")
        return
    await m.answer_document(BufferedInputFile(data, fname), caption=f"💳 Оплаты по MAWB {awb}")
    from .api import awb_operator_xlsx
    try:
        d2, f2, unpaid = awb_operator_xlsx(awb)
        await m.answer_document(BufferedInputFile(d2, f2), caption=f"📊 Учёт по MAWB {awb} (формат оператора)"
                                + (f"\nНе оплачены, в файл не вошли: {', '.join(unpaid)}" if unpaid else ""))
    except Exception as e:
        await m.answer(f"Учёт в формате оператора: {getattr(e, 'detail', e)}")


@dp.message(pv, sysf, Command("notify_reset"))
async def notify_reset(m: Message):
    from . import staffnotify
    staffnotify.set_targets([])
    await m.answer("Список чатов «Уведомления по грузам» очищен. Напиши /notify_here ОДИН раз в нужном чате (или теме).")


@dp.message(pv, sysf, Command("archive"))
async def archive_cmd(m: Message):
    """Last documents anyone sent — with a button to get each original back."""
    from .api import list_uploads
    ups = list_uploads(m.from_user.id)[:15]
    if not ups:
        await m.answer("Архив пуст.")
        return
    lines, rows = [], []
    for u in ups:
        mark = "🗑" if u["gone"] else "📎"
        lines.append(f"{mark} {u['ts']} · {u['user']} · {u['kind_ru']} · {u['summary'] or u['filename']}"
                     + (f"\n   ↳ {u['status']}" if u["status"] else ""))
        rows.append([(f"{mark} #{u['id']} {(u['summary'] or u['filename'])[:28]}", f"arch:{u['id']}")])
    await m.answer("Архив документов (последние 15). Полный список — «Учёт» → Ещё → Архив.\n\n" + "\n".join(lines)[:3800],
                   reply_markup=_kb(rows))


@dp.callback_query(F.data.startswith("arch:"), sysf)
async def archive_send(c: CallbackQuery):
    from .api import resend_upload
    await c.answer()
    try:
        await resend_upload(int(c.data.split(":")[1]), c.from_user.id)
    except Exception as e:
        await bot.send_message(c.from_user.id, f"Не смог отправить: {e}")


@dp.message(pv, wr, Command("pay"))
async def pay_cmd(m: Message):
    """Show the payment list now (same as the Tue/Wed reminder)."""
    await payment_reminder([m.from_user.id])


REMINDER_SLOTS = [(1, 10), (2, 10)]   # (weekday Mon=0, hour MSK): Tuesday 10:00 and Wednesday 10:00


async def push_packing_all(awb_key: str | None = None, only: list[str] | None = None) -> dict:
    """«Push»: packing lists of ALL goods in transit (with a MAWB) -> every packing chat, again."""
    from aiogram.types import BufferedInputFile
    from . import packing, reader
    tg, sender = packing.targets(), reader.RBOT
    if only:                                   # chats picked in the app: "chat_id:thread_id"
        tg = [t for t in tg if f"{t['chat_id']}:{t.get('thread_id') or ''}" in only]
    if not sender:
        return {"error": "бот-читатель выключен (нет READER_BOT_TOKEN)"}
    if not tg:
        return {"error": "нет чатов для пакингов — добавь @бота-читателя в чат и напиши там /packing_here"}
    items, skipped = packing.in_transit_all(awb_key)
    from . import kbreak
    held = [it for it in items if it[3] and not kbreak.has(it[3])]
    items = [it for it in items if not (it[3] and not kbreak.has(it[3]))]
    skipped += [f"{it[1]} (нет детализации {it[2]})" for it in held]
    sent = await _send_items(items, sender, tg, resend_breakdown=True)
    return {"sent": sent, "chats": len(tg), "skipped": skipped}


def push_text(r: dict) -> str:
    if r.get("error"):
        return "⚠️ " + r["error"]
    txt = f"📦 Отправлено пакингов: {r['sent']} (в {r['chats']} чат(а))"
    no_bd = [x for x in r["skipped"] if "детализации" in x]
    other = [x for x in r["skipped"] if "детализации" not in x and "(" in x]
    no_awb = [x for x in r["skipped"] if "детализации" not in x and "(" not in x]
    if no_awb:
        txt += "\nБез MAWB, не отправлены: " + ", ".join(no_awb)
    if other:
        txt += "\nНе отправлены:\n" + "\n".join("• " + x for x in other)
    if no_bd:
        awbs = sorted({x.split("детализации ")[-1].rstrip(")") for x in no_bd})
        txt += ("\nЖдут консолидационный лист (детализацию по фермам) по MAWB " + ", ".join(awbs)
                + ": " + ", ".join(x.split(" (")[0] for x in no_bd) + ". Пришли файл боту — уйдёт сразу.")
    return txt


@dp.message(pv, wr, Command("push_packing"))
async def push_packing_cmd(m: Message):
    from . import packing
    aw = packing.awbs_in_transit()
    if not aw:
        await m.answer("Грузов в пути с MAWB нет.")
        return
    from .kbreak import flag
    rows = [[(f"{'✅' if a['has_bd'] else '⏳'} {a['awb']} · {flag(a['country'])} · {len(a['farms'])} ферм", f"pp:{a['key']}")] for a in aw]
    rows.append([("📦 Все грузы в пути", "pp:all")])
    await m.answer("Какую поставку отправить в чаты пакингов?\n✅ — есть консолидационный лист, ⏳ — ещё нет (не уйдёт)",
                   reply_markup=_kb(rows))


@dp.callback_query(F.data.startswith("pp:"), wr)
async def push_packing_pick(c: CallbackQuery):
    k = c.data.split(":", 1)[1]
    await c.answer("Отправляю…")
    await c.message.edit_text(push_text(await push_packing_all(None if k == "all" else k)))


_LAST_POSTED: list = []      # messages of the last _post call: [(chat_id, message_id)]


async def _post(sender, tg, data: bytes, fname: str, caption: str) -> bool:
    from aiogram.types import BufferedInputFile
    ok = False
    _LAST_POSTED.clear()
    for t in tg:
        try:
            msg = await sender.send_document(t["chat_id"], BufferedInputFile(data, fname), caption=caption,
                                             message_thread_id=t.get("thread_id"))
            _LAST_POSTED.append((t["chat_id"], getattr(msg, "message_id", None)))
            ok = True
        except Exception as e:
            print(f"[lumen] packing chat {t.get('title')}: {e}", flush=True)
    return ok


def _remember_post(awb_key: str):
    """Shipment file posted: keep where, so a corrected invoice can EDIT this message instead of a new one."""
    from . import kbreak
    st = kbreak._state()
    if awb_key in st and _LAST_POSTED:
        st[awb_key]["msgs"] = [{"chat_id": c, "message_id": m} for c, m in _LAST_POSTED if m]
        kbreak._save_state(st)


async def refresh_shipment(awb_key: str, note: str) -> int:
    """Rebuild the shipment file (детализация + all packing lists) and EDIT the messages already in the chats."""
    from aiogram.types import BufferedInputFile, InputMediaDocument
    from . import kbreak, packing, reader
    from .models import Invoice
    from .calc import norm_awb
    st = kbreak._state().get(awb_key) or {}
    if not st.get("msgs") or not reader.RBOT:
        return 0
    with session() as s:
        invs = [i for i in s.exec(select(Invoice)).all() if norm_awb(i.awb) == awb_key and i.packing_sent]
    ids, farms = [i.id for i in invs], [i.farm for i in invs]
    data = packing.bundle(ids, kbreak.breakdown_bytes(awb_key))
    cap = (f"🔄 Исправлено: {note}\n" + kbreak.caption(awb_key, farms))[:1024]
    fname = "Shipment_" + re.sub(r"[^\w\-]+", "_", st.get("awb") or awb_key) + ".xlsx"
    done = 0
    for m in st["msgs"]:
        try:
            await reader.RBOT.edit_message_media(
                chat_id=m["chat_id"], message_id=m["message_id"],
                media=InputMediaDocument(media=BufferedInputFile(data, fname), caption=cap))
            done += 1
        except Exception as e:
            print(f"[lumen] edit shipment message: {e}", flush=True)
    return done


async def _send_items(items, sender, tg, resend_breakdown=False) -> int:
    """One message per MAWB: Kenya — [детализация + all packing lists] in ONE file with the breakdown as text;
    others — all packing lists of the MAWB in one file."""
    from collections import OrderedDict
    from . import kbreak, packing
    groups = OrderedDict()
    for inv_id, farm, awb, kkey in sorted(items, key=lambda x: (x[2] or "", x[0])):
        if not resend_breakdown and packing.ever_posted(inv_id):
            packing.mark_sent(inv_id)              # this farm's packing for this MAWB is already in the chats
            print(f"[lumen] auto packing: skip {farm} {awb} (уже был в чатах)", flush=True)
            continue
        if not resend_breakdown:
            print(f"[lumen] auto packing: send {farm} {awb} inv#{inv_id}", flush=True)
        groups.setdefault((awb, kkey), []).append((inv_id, farm))
    sent = 0
    for (awb, kkey), invs in groups.items():
        ids, farms = [i for i, _f in invs], [f for _i, f in invs]
        safe = re.sub(r"[^\w\-]+", "_", awb)
        with_bd = bool(kkey) and (resend_breakdown or not kbreak.is_sent(kkey))
        if with_bd:
            data = packing.bundle(ids, kbreak.breakdown_bytes(kkey))
            ok = await _post(sender, tg, data, f"Shipment_{safe}.xlsx", kbreak.caption(kkey, farms))
            if ok:
                kbreak.mark_sent(kkey)
                _remember_post(kkey)
        else:
            data = packing.bundle(ids)
            cap = f"📦 Packing list{'s' if len(ids) > 1 else ''} · MAWB {awb}\n" + "\n".join(f"• {f}" for f in farms)
            ok = await _post(sender, tg, data, f"Packing_{safe}.xlsx", cap)
        if ok:
            for i in ids:
                packing.mark_sent(i)
                packing.remember_posted(i)
            sent += len(ids)
    return sent


async def send_packing_lists():
    """Every invoice that got a MAWB -> its packing list (.xlsx, no prices) to every packing chat.
    Kenya: only after the MAWB box breakdown, which is posted right before."""
    from . import packing, reader
    tg, sender = packing.targets(), reader.RBOT          # the neutral bot posts; the finance bot never shows
    if not tg or not sender:
        return
    from . import kbreak as _kbr
    try:
        _kbr.reconcile()                                  # names / MAWB / kg — the shipment completes by itself
    except Exception as e:
        print(f"[lumen] reconcile: {e}", flush=True)
    while _kbr.MOVED:
        txt = _kbr.MOVED.pop(0)
        for uid in _money_people():
            try:
                await bot.send_message(uid, txt)
            except Exception:
                pass
    items = packing.pending()
    if not items:
        return
    # a shipment goes out automatically only when EVERY farm of its consolidation list has its invoice
    # (or «🚫 Инвойса не будет» was pressed). Late invoices of an already-sent shipment go as before.
    from . import kbreak
    waiting = {x["awb_key"] for x in kbreak.missing_invoices()}
    items = [it for it in items if it[3] not in waiting]       # incomplete shipment: send NOTHING
    if items:
        fresh = sorted({it[2] for it in items if it[3] and not kbreak.is_sent(it[3])})
        await _send_items(items, sender, tg)
        for awb in fresh:
            for uid in _money_people():
                try:
                    await bot.send_message(uid, f"✅ Все инвойсы по MAWB {awb} на месте — поставка ушла в чат пакингов.")
                except Exception:
                    pass


async def missing_reminder(uid_list=None):
    """Every day: farms from consolidation lists without an uploaded invoice — ask until they appear."""
    from . import kbreak
    miss = kbreak.missing_invoices()
    if not miss:
        return False
    by = {}
    for x in miss:
        by.setdefault((x["awb_key"], x["awb"], x["country"]), []).append(x)
    lines, rows = ["❗️ Где инвойсы? По разбивке они есть, а в боте их нет:"], []
    for (k, awb, country), xs in by.items():
        from .kbreak import flag
        lines.append(f"\n✈️ {flag(country)} · " + ("закупки брокера (MAWB ещё нет):" if k == "broker" else f"MAWB {awb}:"))
        for x in xs:
            lines.append((f"• {x['farm']} — {x['packs']} кор." if x.get("packs") else f"• {x['farm']}")
                         + (f" ({x['note']})" if x.get("note") else ""))
            el = kbreak.elsewhere(x["farm"], k) if k != "broker" else None
            if el:
                lines.append(f"   ↳ инвойс {x['farm']} есть, но на MAWB {el['awb']} (ферма написала свой)")
                rows.append([(f"↪️ {x['farm']}: перенести на {awb}", f"mv:{el['id']}:{k}")])
            rows.append([(f"🚫 {x['farm']} ({awb[-4:]}) — инвойса не будет", f"miss:{k}:{x['farm'][:40]}")])
    lines.append("\nКинь инвойсы сюда — сами привяжутся к MAWB. Пока их нет, пакинг этих ферм не уйдёт, а я буду спрашивать каждый день 🙂")
    for uid in uid_list or _money_people():
        try:
            await bot.send_message(uid, "\n".join(lines)[:4000], reply_markup=_kb(rows[:20]))
        except Exception:
            pass
    return True


@dp.callback_query(F.data.startswith("mv:"), wr)
async def move_to_awb(c: CallbackQuery):
    from . import kbreak
    from .models import Invoice
    _, iid, k = c.data.split(":")
    awb = (kbreak._state().get(k) or {}).get("awb") or k
    with session() as s:
        i = s.get(Invoice, int(iid))
        old, i.awb, i.packing_sent = i.awb, awb, False
        s.add(i); s.commit()
        farm = i.farm
    from .backup import mark_dirty
    mark_dirty()
    await c.answer("Перенёс")
    await c.message.answer(f"↪️ {farm}: MAWB {old} → {awb}. Как только все инвойсы поставки будут на месте — она уйдёт сама.")
    await send_packing_lists()


@dp.callback_query(F.data.startswith("miss:"), wr)
async def missing_skip(c: CallbackQuery):
    from . import kbreak
    _, k, farm = c.data.split(":", 2)
    kbreak.skip_farm(k, farm)
    await c.answer("Ок, больше не спрашиваю")
    await c.message.answer(f"Ок, по {farm} инвойса не будет — больше не спрашиваю.")


@dp.message(pv, wr, Command("missing"))
async def missing_cmd(m: Message):
    if not await missing_reminder([m.from_user.id]):
        await m.answer("✅ По всем разбивкам инвойсы на месте.")


def _discrepancies():
    """Arrived invoices with an unresolved mixed-box mismatch."""
    from .models import Invoice
    out = []
    with session() as s:
        for i in s.exec(select(Invoice)).all():
            if i.arrived_at and i.discrepancy_json:
                for n, d in enumerate(json.loads(i.discrepancy_json)):
                    if not d.get("resolved"):
                        out.append((i, n, d))
    return out


async def discrepancy_alert(uid_list=None):
    items = _discrepancies()
    for inv, n, d in items:
        whole = d["box"] == "ИТОГО инвойса"
        txt = (f"🚨 ГРУЗ ПРИЕХАЛ — СРОЧНО УЗНАЙ У КЛАДОВЩИКА!\n"
               + (f"{inv.farm.upper()} · MAWB {inv.awb} · ИНВОЙС {inv.invoice_no or ''}:\n"
                  f"ПО СТРОКАМ {d['units']:g} СТ, А В ИТОГЕ ИНВОЙСА {d['box_stems']:g} СТ.\n" if whole else
                  f"{inv.farm.upper()} · MAWB {inv.awb} · КОРОБКА {d['box'].upper()}:\n"
                  f"ПО СОРТАМ В ИНВОЙСЕ {d['units']:g} СТ, А В КОРОБКЕ {d['box_stems']:g} СТ.\n")
               + f"СКОЛЬКО ПРИЕХАЛО ПО ФАКТУ?\n\nСорта: {', '.join(d.get('varieties') or [])}\n"
               + "Себестоимость пока не трогаю — посчитаю, когда ответишь.")
        kb = _kb([[(f"{d['units']:g} — как по {'строкам' if whole else 'сортам'}", f"dsc:{inv.id}:{n}:u")],
                  [(f"{d['box_stems']:g} — как в {'итоге' if whole else 'коробке'}", f"dsc:{inv.id}:{n}:b")],
                  [("Другое число", f"dsc:{inv.id}:{n}:o")]])
        for uid in uid_list or _money_people():
            try:
                await bot.send_message(uid, txt, reply_markup=kb)
            except Exception:
                pass
    return bool(items)


_DSC_WAIT: dict = {}     # user -> (inv_id, n) waiting for a typed fact


async def _ask_variety(target, inv_id: int, n: int, diff: float):
    """Fact ≠ variety sum: which variety is short / extra?"""
    from .models import Line
    with session() as s:
        ls = s.exec(select(Line).where(Line.invoice_id == inv_id)).all()
    word = "МЕНЬШЕ" if diff < 0 else "БОЛЬШЕ"
    rows = [[(f"{l.name} ({l.stems:g} ст)", f"dsv:{inv_id}:{n}:{l.id}:{diff:g}")] for l in ls]
    await target(f"Какого сорта {word} на {abs(diff):g} ст? Поправлю количество и пересчитаю себестоимость.", _kb(rows))


@dp.callback_query(F.data.startswith("dsc:"), wr)
async def discrepancy_answer(c: CallbackQuery):
    from .models import Invoice
    _, inv_id, n, what = c.data.split(":")
    inv_id, n = int(inv_id), int(n)
    with session() as s:
        d = json.loads(s.get(Invoice, inv_id).discrepancy_json)[n]
    await c.answer()
    if what == "u":                      # fact = varieties: lines are already right
        _resolve(inv_id, n, d["units"])
        await c.message.edit_text(f"✅ По факту {d['units']:g} ст, как по сортам. Сорта и себестоимость оставляю так "
                                  f"(ферма выставила за {d['box_stems']:g}, лишние стебли удешевляют партию).")
        return
    if what == "b":                      # fact = box: some variety is short
        await c.message.edit_text(f"По факту {d['box_stems']:g} ст.")
        await _ask_variety(lambda t, kb: c.message.answer(t, reply_markup=kb), inv_id, n, d["box_stems"] - d["units"])
        return
    _DSC_WAIT[c.from_user.id] = (inv_id, n)
    await c.message.edit_text("Напиши, сколько стеблей приехало по факту в этой коробке, например `245`.", parse_mode="Markdown")


async def send_correction(inv_id: int, before: dict, why: str):
    """«Корректировка себестоимости» + Excel of this farm's invoice (before → after) to owners and the 1С operator."""
    import io as _io
    from aiogram.types import BufferedInputFile
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from .api import invoice_costs
    from .models import Invoice
    after = invoice_costs(inv_id)
    with session() as s:
        inv = s.get(Invoice, inv_id)
        farm, awb, no = inv.farm, inv.awb, inv.invoice_no
    wb = Workbook(); ws = wb.active; ws.title = farm[:28] or "Ферма"
    ws.append([f"Корректировка себестоимости · {farm} · инвойс {no} · MAWB {awb}"]); ws["A1"].font = Font(bold=True, size=12)
    ws.append([why]); ws.append([])
    head = ["Номенклатура", "Стебли было", "Стебли стало", "Себестоимость было, ₽/ст", "Себестоимость стало, ₽/ст", "Разница, ₽/ст"]
    ws.append(head)
    for c in range(1, 7):
        ws.cell(4, c).font = Font(bold=True); ws.cell(4, c).fill = PatternFill("solid", fgColor="D9E1F2")
    for lid, a in after.items():
        b = before.get(lid, a)
        ws.append([a["name"], b["stems"], a["stems"], round(b["total"], 4), round(a["total"], 4), round(a["total"] - b["total"], 4)])
        if abs(a["total"] - b["total"]) > 0.0001 or a["stems"] != b["stems"]:
            for c in range(1, 7):
                ws.cell(ws.max_row, c).fill = PatternFill("solid", fgColor="FFF2CC")
    for col, w in zip("ABCDEF", (36, 12, 12, 22, 22, 14)):
        ws.column_dimensions[col].width = w
    buf = _io.BytesIO(); wb.save(buf)
    cap = f"🔁 Корректировка себестоимости · {farm} · MAWB {awb}\n{why}"
    ids = set(_money_people()) | {u["tg_id"] for u in roles.users() if u["role"] == "viewer"}
    safe_farm = re.sub(r"[^\w\-]+", "_", farm)
    fname = "Корректировка_" + safe_farm + ".xlsx"
    for uid in ids:
        try:
            await bot.send_document(uid, BufferedInputFile(buf.getvalue(), fname),
                                    caption=cap)
        except Exception:
            pass


@dp.callback_query(F.data.startswith("dsv:"), wr)
async def discrepancy_variety(c: CallbackQuery):
    from .models import Line
    from .api import invoice_costs
    _, inv_id, n, line_id, diff = c.data.split(":")
    before = invoice_costs(int(inv_id))
    with session() as s:
        l = s.get(Line, int(line_id))
        l.stems = max(0, l.stems + float(diff))
        s.add(l); s.commit()
        name, st = l.name, l.stems
    with session() as s:
        from .models import Invoice
        d = json.loads(s.get(Invoice, int(inv_id)).discrepancy_json)[int(n)]
    _resolve(int(inv_id), int(n), d["units"] + float(diff))
    await c.answer()
    await c.message.edit_text(f"✅ {name}: теперь {st:g} ст. Себестоимость пересчитана по факту.")
    await send_correction(int(inv_id), before, f"По факту со склада: {name} {'+' if float(diff) > 0 else ''}{float(diff):g} ст "
                                               f"({d['box']}: в инвойсе {d['units']:g}, приехало {d['units'] + float(diff):g}).")


def _resolve(inv_id: int, n: int, fact: float):
    from .models import Invoice
    with session() as s:
        inv = s.get(Invoice, inv_id)
        ds = json.loads(inv.discrepancy_json)
        ds[n]["resolved"], ds[n]["fact"] = True, fact
        inv.discrepancy_json = json.dumps(ds, ensure_ascii=False)
        s.add(inv); s.commit()
    from .backup import mark_dirty
    mark_dirty()


@dp.message(pv, wr, F.from_user.id.func(lambda i: i in _DSC_WAIT), F.text.regexp(r"^\s*\d+([.,]\d+)?\s*$"))
async def discrepancy_typed(m: Message):
    from .models import Invoice
    inv_id, n = _DSC_WAIT.pop(m.from_user.id)
    fact = float(m.text.replace(",", "."))
    with session() as s:
        d = json.loads(s.get(Invoice, inv_id).discrepancy_json)[n]
    diff = fact - d["units"]
    if abs(diff) < 0.5:
        _resolve(inv_id, n, fact)
        await m.answer("✅ Совпадает с сортами — оставляю как есть.")
        return
    await _ask_variety(lambda t, kb: m.answer(t, reply_markup=kb), inv_id, n, diff)


_STARTED = time.time()


# ---------- Logiztik API: consolidation lists for Ecuador / Colombia by themselves ----------
class _LzMsg:
    """Lets the Logiztik sync reuse the «consolidation list» flow: its answers go to the money people."""
    def __init__(self, uid):
        from types import SimpleNamespace
        self.from_user = SimpleNamespace(id=uid)

    async def answer(self, text, **kw):
        for uid in _money_people():
            try:
                await bot.send_message(uid, "🛰 Logiztik — " + text, **kw)
            except Exception:
                pass


async def logiztik_sync(report_to: int | None = None) -> str:
    """New AWBs at Logiztik -> consolidation list (boxes + kg per farm) like a forwarded prealert.
    An AWB that already has a list (prealert sent by hand / already posted) is NOT overwritten — only differences
    are reported."""
    from . import logiztik as lz, kbreak
    from .ai import find_mawb
    from .calc import norm_awb
    from .api import _settings, SETTINGS, resolve_farm
    if not lz.enabled():
        return "Logiztik не настроен: в Bothost нужны переменные LOGIZTIK_TOKEN и LOGIZTIK_CUSTOMER."
    seen = dict(_settings().get("logiztik_seen") or {})
    lines = []
    for sh in await lz.recent(6):
        boxes = await lz.barcodes(sh["awb"])
        if not boxes:
            continue
        awb = find_mawb(sh["awb"]) or sh["awb"]
        key = norm_awb(awb)
        if seen.get(key) == len(boxes):
            continue
        rows = lz.farm_rows(boxes)
        st = kbreak._state().get(key)
        if st and (st.get("src") != "logiztik" or st.get("sent")):
            have = {}
            with session() as s:
                for f in st.get("farms", []):
                    have[f["farm"]] = f.get("packs") or 0
                diff = []
                for r in rows:
                    ff = resolve_farm(s, r["farm_raw"])
                    name = ff.name if ff else _farm_label(r["farm_raw"])
                    if have.get(name) != r["packs"]:
                        diff.append(f"• {name}: Logiztik {r['packs']} кор., в листе {have.get(name, '—')}")
            if diff:
                await _LzMsg(0).answer(f"MAWB {awb}: коробки расходятся с консолидационным листом:\n" + "\n".join(diff))
            lines.append(f"{awb}: лист уже был — сверил ({'расхождения отправил' if diff else 'всё сходится'})")
        else:
            from datetime import datetime as _d
            try:
                eta = _d.strptime(sh["date"], "%Y-%m-%d")
            except ValueError:
                eta = None
            await _consolidation(_LzMsg(report_to or (_money_people() or [0])[0]), awb, sh["country"], rows, eta)
            stt = kbreak._state()
            if key in stt:
                stt[key]["src"] = "logiztik"
                kbreak._save_state(stt)
            lines.append(f"{awb}: {len(boxes)} кор. — внёс как консолидационный лист")
        seen[key] = len(boxes)
    SETTINGS.write_text(json.dumps({**_settings(), "logiztik_seen": seen}, ensure_ascii=False))
    return "🛰 Logiztik: " + ("\n".join(lines) if lines else "новых поставок нет")


@dp.message(pv, sysf, Command("logiztik"))
async def logiztik_cmd(m: Message):
    """/logiztik — проверка связи; /logiztik sync — забрать новые поставки сейчас; /logiztik 2026-10-15 — AWB за дату;
    /logiztik awb 369-12345678 — коробки по AWB."""
    from datetime import date as _date
    from . import logiztik as lz
    arg = (m.text or "").split()[1:]
    note = await m.answer("🛰 Logiztik: запрашиваю…")
    try:
        if not lz.enabled():
            raise RuntimeError("не настроен: добавь в Bothost переменные LOGIZTIK_TOKEN и LOGIZTIK_CUSTOMER и перезапусти.")
        if arg and arg[0] == "sync":
            txt = await logiztik_sync(m.from_user.id)
        elif arg and arg[0] == "awb" and len(arg) > 1:
            boxes = await lz.barcodes(arg[1])
            rows = lz.farm_rows(boxes)
            txt = (f"🛰 AWB {arg[1]}: {len(boxes)} кор.\n" + "\n".join(
                f"• {r['farm_raw']}: {r['packs']} кор. · {r['kg']:g} кг"
                + (f" · {', '.join(f'{k}×{v}' for k, v in r['sizes'].items())}" if r["sizes"] else "")
                + (f" · HAWB {r['hawb']}" if r["hawb"] else "") + f" · принято {r['received']}/{r['packs']}"
                for r in rows))
        else:
            day = _date.fromisoformat(arg[0]) if arg else None
            if day:
                sh = await lz.shipments(day)
                txt = f"🛰 {day:%d.%m.%Y}: AWB {len(sh)}\n" + "\n".join(
                    f"• {x.get('awb')} · {x.get('origin')} · мест {x.get('pieces')}" for x in sh)
            else:
                rec = await lz.recent(4)
                from .api import _settings
                txt = (f"🛰 Logiztik: связь есть ✅ (ключ передаю как «{_settings().get('logiztik_auth')}»).\n"
                       f"AWB за 4 дня: {len(rec)}\n" + "\n".join(f"• {x['awb']} · {x['date']} · {x['country']}" for x in rec)
                       + "\n\n/logiztik sync — внести новые поставки сейчас (дальше бот делает это сам 5 раз в день).")
    except Exception as e:
        txt = f"🛰 Logiztik: ❌ {e}"
    await note.edit_text(txt[:4000])


def _stale_transit() -> list:
    """MAWBs still «в пути» with no arrival time from Floratrack and older than a normal trip -> ask."""
    from datetime import datetime as _d, timedelta as _td
    from .api import _settings, _dt
    from .calc import norm_awb
    from .models import Invoice
    from . import kbreak
    snooze = _settings().get("arr_snooze") or {}
    lists = kbreak._state()
    groups = {}
    with session() as s:
        for i in s.exec(select(Invoice)).all():
            if i.arrived_at or i.eta or not (i.awb or "").strip():
                continue
            groups.setdefault(norm_awb(i.awb), []).append(i)
    out = []
    for k, invs in groups.items():
        if snooze.get(k, "") > _d.now().strftime("%Y-%m-%d"):
            continue
        st = lists.get(k) or {}
        start = _dt(st.get("ts") or "") or max((_dt(i.invoice_date or "") for i in invs if _dt(i.invoice_date or "")),
                                                default=None)
        if not start:
            continue
        country = st.get("country") or invs[0].country or ""
        days = (_d.now() - start).days
        if days >= (7 if country == "Кения" else 10):
            out.append({"key": k, "awb": invs[0].awb, "days": days, "country": country,
                        "farms": sorted({i.farm for i in invs})})
    return out


async def arrival_check(uid_only: int | None = None):
    rows = _stale_transit()
    if not rows and uid_only:
        await bot.send_message(uid_only, "Все поставки «в пути» в норме — зависших без отметки о прибытии нет.")
    for r in rows:
        txt = (f"🛬 MAWB {r['awb']} ({r['country'] or '?'}: {', '.join(r['farms'])}) — в пути уже {r['days']} дн., "
               "а сообщения Floratrack о прибытии не было. Прибыл?")
        kb = _kb([[("✅ Да, прибыл", f"arq:{r['key']}:yes"), ("🚚 Ещё в пути", f"arq:{r['key']}:no")]])
        for uid in ([uid_only] if uid_only else _money_people()):
            try:
                await bot.send_message(uid, txt, reply_markup=kb)
            except Exception:
                pass


async def pickup_check(uid_only: int | None = None, min_hours: float = 20):
    """Arrived MAWBs nobody confirmed «забрал со склада ТК» -> ask (from the stock screen as well)."""
    from .api import awaiting_pickup
    rows = [r for r in awaiting_pickup() if r["since_h"] >= min_hours]
    if not rows and uid_only:
        await bot.send_message(uid_only, "На складе ТК ничего не ждёт — всё прибывшее отмечено как забранное.")
    for r in rows:
        txt = (f"🏬 MAWB {r['awb']} ({', '.join(r['farms'])}) прибыл {r['arrived']}. Груз со склада ТК забрали?")
        kb = _kb([[("✅ Забрал", f"pk:{r['key']}"), ("⏳ Ещё на складе", f"pkn:{r['key']}")]])
        for uid in ([uid_only] if uid_only else _money_people()):
            try:
                await bot.send_message(uid, txt, reply_markup=kb)
            except Exception:
                pass


@dp.callback_query(F.data.startswith("pk:"), wr)
async def picked_answer(c: CallbackQuery):
    from .api import mark_picked
    k = c.data.split(":", 1)[1]
    farms = mark_picked(k)
    await c.answer("Отмечено")
    try:
        await c.message.edit_text((c.message.text or "") + f"\n\n✅ MAWB {k[:3]}-{k[3:]} забран со склада ТК "
                                  f"({', '.join(farms)}) — убрал из товарных остатков.")
    except Exception:
        pass


@dp.callback_query(F.data.startswith("pkn:"), wr)
async def picked_not_yet(c: CallbackQuery):
    await c.answer("Ок, спрошу завтра")
    try:
        await c.message.edit_text((c.message.text or "") + "\n⏳ Ок, ещё на складе — спрошу завтра.")
    except Exception:
        pass


@dp.callback_query(F.data.startswith("arq:"), wr)
async def arrival_answer(c: CallbackQuery):
    from datetime import datetime as _d, timedelta as _td
    from .api import _settings, SETTINGS
    from .calc import norm_awb
    from .models import Invoice
    _, k, ans = c.data.split(":")
    await c.answer()
    if ans == "yes":
        farms = []
        with session() as s:
            for i in s.exec(select(Invoice)).all():
                if norm_awb(i.awb) == k and not i.arrived_at:
                    i.arrived_at = _d.now().strftime("%d.%m.%Y %H:%M") + " (вручную)"
                    s.add(i); farms.append(i.farm)
            s.commit()
        from .backup import mark_dirty
        mark_dirty()
        await c.message.edit_text(c.message.text.split(" — в пути")[0] + f" — ✅ отмечено прибывшим ({', '.join(farms)}). "
                                  "MAWB уйдёт из «В пути» целиком через неделю после того, как оплачены все его инвойсы.")
        await bot.send_message(c.from_user.id, f"Груз MAWB {k[:3]}-{k[3:]} со склада ТК уже забрали?",
                               reply_markup=_kb([[("✅ Забрал", f"pk:{k}"), ("⏳ Ещё на складе", f"pkn:{k}")]]))
    else:
        st = _settings()
        sn = st.get("arr_snooze") or {}
        sn[k] = (_d.now() + _td(days=3)).strftime("%Y-%m-%d")
        SETTINGS.write_text(json.dumps({**st, "arr_snooze": sn}))
        await c.message.edit_text(c.message.text.split(" Прибыл?")[0] + " — ок, спрошу через 3 дня.")


async def scheduler_loop():
    from .api import _settings, SETTINGS, arrive_due
    while True:
        try:
            now = _msk_now()
            if time.time() - _STARTED > 120:          # no automatic posting right after a restart
                await send_packing_lists()
            done = arrive_due(now.strftime("%Y-%m-%dT%H:%M"))
            from . import clients, reader
            from .api import client_arrivals_due
            cdone = client_arrivals_due(now.strftime("%Y-%m-%dT%H:%M"))     # TK time + 6 h
            if cdone:
                await clients.send(reader.RBOT, clients.arrived_messages(cdone))
            slot3 = now.strftime("%Y-%m-%d-%H") if now.hour in (9, 12, 15, 18, 21) else None
            if done or (slot3 and _settings().get("last_dsc") != slot3):
                if slot3:
                    SETTINGS.write_text(json.dumps({**_settings(), "last_dsc": slot3}))
                await discrepancy_alert()
            if done:
                txt = "📦 Прибыло на склад (по сообщению Floratrack +1 ч):\n" + "\n".join(
                    f"• {d['farm']} · MAWB {d['awb']} · {'оплачен' if d['paid'] else 'НЕ оплачен'}" for d in done)
                from .calc import norm_awb
                awbs = sorted({d["awb"] for d in done if d.get("awb")})
                txt += "\n\nКогда заберёте со склада ТК — нажми, и груз уйдёт из товарных остатков."
                kb = _kb([[(f"✅ Забрал {a}", f"pk:{norm_awb(a)}")] for a in awbs]) if awbs else None
                for uid in _money_people():
                    try:
                        await bot.send_message(uid, txt, reply_markup=kb)
                    except Exception:
                        pass
            slot = None
            for wd, hh in REMINDER_SLOTS:
                if now.weekday() == wd and now.hour == hh and now.minute < 15:
                    slot = now.strftime(f"%Y-%m-%d-{hh}")
            st = _settings()
            if now.hour == 11 and now.minute < 15 and st.get("last_missing") != now.strftime("%Y-%m-%d"):
                SETTINGS.write_text(json.dumps({**st, "last_missing": now.strftime("%Y-%m-%d")}))
                st = _settings()
                await missing_reminder()
                await arrival_check()
                await pickup_check()
            lz_slot = now.strftime("%Y-%m-%d-%H") if now.hour in (9, 12, 15, 18, 21) and now.minute < 15 else None
            if lz_slot and st.get("last_lz") != lz_slot:
                SETTINGS.write_text(json.dumps({**st, "last_lz": lz_slot}))
                st = _settings()
                from . import logiztik as _lz
                if _lz.enabled():
                    try:
                        await logiztik_sync()
                    except Exception as e:
                        print(f"[lumen] logiztik: {e}", flush=True)
            if now.weekday() == 1 and now.hour == 10 and now.minute < 15 and st.get("last_bf") != now.strftime("%Y-%m-%d"):
                SETTINGS.write_text(json.dumps({**st, "last_bf": now.strftime("%Y-%m-%d")}))
                st = _settings()
                for uid in _money_people():
                    try:
                        await bot.send_message(uid, "Привет, дорогой! 🌷\nПожалуйста, прогрузи выписку баланса брокера BiFlorica "
                                                    "по Plazoleta и Tessa (Excel из личного кабинета) — просто кинь файл сюда.\n"
                                                    "Я сам внесу закупки и пополнения брокера и пересчитаю себестоимость. "
                                                    "Проверь только, из каких пополнений ушли доллары на биржу — я подскажу.")
                    except Exception:
                        pass
            if slot and st.get("last_reminder") != slot:
                SETTINGS.write_text(json.dumps({**st, "last_reminder": slot}))
                await payment_reminder()
        except Exception as e:
            print(f"[lumen] scheduler: {e}", flush=True)
        await asyncio.sleep(30)


# ---------- Floratrack group chat: truck on the way -> arrival = their time + 1 h ------------
grp = Router()
grp.message.filter(F.chat.type.in_({"group", "supergroup"}))
dp.include_router(grp)
_FT_PENDING: dict[int, list] = {}


_TRUCK_BUF: dict[int, list] = {}


async def _truck_event(uid_list, text: str, sent_msk, notify_uid=None):
    """Parse + apply one FLORA TRUCK message; collect summaries and send one message per batch."""
    from . import truck
    ev = truck.parse(text, sent_msk)
    if not ev:
        return False
    from . import clients, reader
    fresh = clients.is_fresh(sent_msk, _msk_now())
    per_client = clients.messages_for(ev) if fresh else {}
    staff_ok = fresh or (notify_uid is not None and _msk_now() - sent_msk <= timedelta(days=3))
    if staff_ok:                                # staff / warehouse chat: the message itself, in our wording
        from . import staffnotify
        await staffnotify.post(reader.RBOT, text)
    line = truck.apply(ev, _msk_now(), fresh=fresh)
    await clients.send(reader.RBOT, per_client)
    targets = [notify_uid] if notify_uid else uid_list
    for uid in targets:
        buf = _TRUCK_BUF.setdefault(uid, [])
        buf.append((sent_msk, line))
        if len(buf) == 1:
            asyncio.create_task(_flush_truck(uid))
    return True


async def _flush_truck(uid: int):
    await asyncio.sleep(4)                       # forwarded messages arrive one by one — wait for the batch
    items = sorted(_TRUCK_BUF.pop(uid, []), key=lambda x: x[0])
    text = "\n\n".join(l for _d, l in items if l)
    if text:
        try:
            await bot.send_message(uid, text[:4000])
        except Exception:
            pass


async def _handle_truck(chat_title: str, text: str, sent_msk):
    await _truck_event(_money_people(), text, sent_msk)


async def process_group_text(chat_id: int, chat_title: str, text: str, sent_msk):
    """A message from a group/channel (seen by the main bot OR by the reader bot).
    Only an approved TK MSK chat is processed; the first time the system admin is asked in private."""
    from .api import _settings
    if not re.search(r"товар\s+забран|едет на склад|время прибытия|прошла границу", text or "", re.I):
        return
    if chat_id in (_settings().get("ft_chats") or []):
        await _handle_truck(chat_title or "", text, sent_msk)
        return
    pending = _FT_PENDING.setdefault(chat_id, [])
    pending.append((text, sent_msk))
    if len(pending) > 1:
        return                                   # already asked; keep collecting until approved
    for uid in roles.sys_ids():
        try:
            await bot.send_message(uid, f"Вижу сообщения о машинах в чате «{chat_title}». Это чат ТК МСК (Floratrack)?",
                                   reply_markup=_kb([[("Да, это ТК МСК", f"ftc:{chat_id}:0"), ("Нет", f"ftn:{chat_id}")]]))
        except Exception:
            pass


_READER_SEEN: dict = {}


_JOIN: dict = {}     # chat_id -> (title, type) for the «what is this chat for?» buttons


async def reader_joined(chat_id: int, title: str, chat_type: str, status: str, reads_all: bool, username: str):
    """The reader bot was added to (or removed from) a chat: ask what the chat is for (TK MSK / packing / client)."""
    from .api import _settings
    key = (chat_id, "out" if status in ("left", "kicked") else "in")
    if _READER_SEEN.get(chat_id) == key[1]:
        return                                   # same state already reported
    _READER_SEEN[chat_id] = key[1]
    if status in ("left", "kicked"):
        text, kb = f"⚠️ Читатель @{username} удалён из чата «{title}».", None
    else:
        _JOIN[chat_id] = (title, chat_type)
        ok_privacy = reads_all or chat_type == "channel"
        ok_admin = chat_type != "channel" or status == "administrator"
        warn = ([] if ok_privacy else ["❌ не видит обычные сообщения — в @BotFather: /setprivacy → Disable, потом удали и добавь бота заново"]) + \
               ([] if ok_admin else ["❌ в канале бот должен быть админом"])
        text = "\n".join([f"👀 Читатель @{username} добавлен в «{title}»."] + warn + ["Для чего этот чат?"])
        kb = _kb([[("🚛 Чат ТК МСК (машины)", f"rj:ft:{chat_id}")],
                  [("📦 Сюда пакинг-листы", f"rj:pk:{chat_id}")],
                  [("👥 Чат клиента (статусы груза)", f"rj:cl:{chat_id}")],
                  [("🚛 Уведомления по грузам (наш склад/сотрудники)", f"rj:nt:{chat_id}")],
                  [("Ничего, просто так", f"rj:no:{chat_id}")]])
    for uid in roles.sys_ids():
        try:
            await bot.send_message(uid, text, reply_markup=kb)
        except Exception:
            pass


@dp.callback_query(F.data.startswith("rj:"), sysf)
async def reader_join_choice(c: CallbackQuery):
    from .api import _settings, SETTINGS
    from . import packing, clients
    _, what, cid = c.data.split(":", 2)
    cid = int(cid)
    title, ctype = _JOIN.get(cid, ("", ""))
    await c.answer()
    if what == "ft":
        st = _settings()
        SETTINGS.write_text(json.dumps({**st, "ft_chats": sorted(set((st.get("ft_chats") or []) + [cid]))}))
        txt = f"✅ «{title}» — чат ТК МСК. Сообщения о машинах буду разбирать сами."
    elif what == "pk":
        if ctype == "supergroup":
            txt = (f"«{title}» — супергруппа с темами. Открой нужную тему (например «Packing листы») "
                   "и напиши там /packing_here — пакинги будут приходить только в неё, не в General.")
        else:
            ts = [x for x in packing.targets() if not (x["chat_id"] == cid and x.get("thread_id") is None)]
            packing.set_targets(ts + [{"chat_id": cid, "thread_id": None, "title": title}])
            txt = f"✅ В «{title}» будут приходить пакинг-листы."
    elif what == "nt":
        from . import staffnotify
        if ctype == "supergroup":
            txt = (f"«{title}» — супергруппа с темами. Открой нужную тему и напиши там /notify_here — "
                   "уведомления по грузам будут приходить в неё.")
        else:
            staffnotify.add(cid, None, title)
            txt = f"✅ В «{title}» будут приходить уведомления по грузам (сообщения ТК МСК в нашем формате)."
    elif what == "cl":
        marks = sorted(clients.registry().keys())
        rows = [[(m, f"rjm:{cid}:{m}")] for m in marks[:20]]
        await c.message.edit_text(f"«{title}» — чат какого клиента? Выбери маркировку"
                                  + ("" if marks else " (маркировок пока нет)") +
                                  ".\nНовой нет в списке — добавь её в «Ещё» → «Маркировки клиентов» "
                                  "или напиши в том чате /marking_here КОД.",
                                  reply_markup=_kb(rows) if rows else None)
        return
    else:
        txt = f"Ок, «{title}» ни для чего не использую."
    await c.message.edit_text(txt)


@dp.callback_query(F.data.startswith("rjm:"), sysf)
async def reader_join_marking(c: CallbackQuery):
    from . import clients
    _, cid, mk = c.data.split(":", 2)
    title, _t = _JOIN.get(int(cid), ("", ""))
    clients.add_chat(mk, int(cid), None, title)
    await c.answer()
    await c.message.edit_text(f"✅ «{title}» — чат клиента {mk}. Статусы его грузов будут приходить туда.")


@grp.message(F.text)
async def group_text(m: Message):
    from datetime import timedelta
    await process_group_text(m.chat.id, m.chat.title or "", m.text or "", m.date.replace(tzinfo=None) + timedelta(hours=3))


@dp.callback_query(F.data.startswith("ftc:"), sysf)
async def ft_chat_yes(c: CallbackQuery):
    from datetime import datetime
    from .api import _settings, SETTINGS
    _, cid, ts = c.data.split(":")
    cid = int(cid)
    st = _settings()
    SETTINGS.write_text(json.dumps({**st, "ft_chats": sorted(set((st.get("ft_chats") or []) + [cid]))}))
    await c.message.edit_text("✅ Запомнил чат Floratrack — машины из него буду закрывать как «прибыл» сами.")
    await c.answer()
    for text, sent in _FT_PENDING.pop(cid, []):
        await _handle_truck("", text, sent)


@dp.callback_query(F.data.startswith("ftn:"), sysf)
async def ft_chat_no(c: CallbackQuery):
    await c.message.edit_text("Ок, этот чат игнорирую.")
    await c.answer()


@dp.message(pv, wr, Command("arrivals"))
async def arrivals_cmd(m: Message):
    await arrival_check(m.from_user.id)


@dp.message(pv, wr, Command("pickup"))
async def pickup_cmd(m: Message):
    """Что прибыло и ждёт «забрал со склада ТК»."""
    await pickup_check(m.from_user.id, min_hours=0)


@dp.message(pv, wr, Command("log"))
async def autolog_cmd(m: Message):
    """Последние изменения, которые бот сделал сам."""
    from .autolog import entries
    rows = entries(20)
    if not rows:
        await m.answer("🤖 Бот пока ничего сам не менял.")
        return
    txt = "🤖 Что бот поменял сам (последние 20). Отменить — в приложении: Ещё → «Что бот поменял».\n\n" + "\n".join(
        f"{'↩️ ' if r['undone'] else ''}{r['ts']} · {r['farm'] or '—'}"
        + (f" · {r['label']}: {r['old'] or 'пусто'} → {r['new'] or 'пусто'}" if r['field'] != 'merge' else f" · {r['new']}")
        + f"\n   ↳ {r['reason']}" for r in rows)
    await m.answer(txt[:4000])


# catch-all goes LAST so /pay and other commands above get a chance first
@dp.message(pv, ops)
async def other(m: Message):
    await m.answer("Жду PDF/фото инвойса или .xlsx мастер-файла. Всё остальное — в «Учёт».")
