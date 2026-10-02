import hashlib
import re
import hmac
import io
import json
import shutil
import time
import uuid
from collections import defaultdict
from statistics import median
from urllib.parse import parse_qsl

from fastapi import APIRouter, Depends, File, Header, HTTPException, UploadFile
from pydantic import BaseModel
from sqlmodel import select

from . import ai, excel
from .calc import compute, norm_awb, rate_of
from .config import ALLOWED_IDS, BOT_TOKEN, DATA_DIR, DEV_NO_AUTH, MASTER_XLSX
from .models import AwbWeights, Farm, Invoice, Line, Logistics, TopUp, session

router = APIRouter(prefix="/api")
BOT = None  # set by main.py so export can send the file into the chat


# ---------- auth: Telegram Mini App initData --------------------------------------------
def user_id(x_init_data: str = Header(default="")) -> int:
    from .roles import role_of, sys_ids
    if DEV_NO_AUTH:
        return (sys_ids() or [next(iter(ALLOWED_IDS), 0)])[0]
    if not x_init_data:
        print("[lumen] auth: пустой initData", flush=True)
        raise HTTPException(401, "Нет данных входа от Telegram. Открой приложение кнопкой «📒 Открыть учёт» "
                                 "под сообщением /start или кнопкой «Учёт» у поля ввода — не старой кнопкой вместо клавиатуры и не в браузере.")
    pairs = dict(parse_qsl(x_init_data, keep_blank_values=True))
    h = pairs.pop("hash", "")
    check = "\n".join(f"{k}={v}" for k, v in sorted(pairs.items()))
    secret = hmac.new(b"WebAppData", BOT_TOKEN.strip().encode(), hashlib.sha256).digest()
    calc = hmac.new(secret, check.encode(), hashlib.sha256).hexdigest()
    if not h or not hmac.compare_digest(calc, h):
        print(f"[lumen] auth: подпись не совпала. поля={sorted(pairs)} токен=...{BOT_TOKEN.strip()[-4:]}", flush=True)
        raise HTTPException(401, f"Подпись Telegram не совпала: BOT_TOKEN на сервере (…{BOT_TOKEN.strip()[-4:]}) "
                                 "не от этого бота. Проверь BOT_TOKEN в .env и в переменных Bothost.")
    if time.time() - int(pairs.get("auth_date", 0)) > 7 * 86400:
        raise HTTPException(401, "Вход устарел — закрой и открой приложение заново")
    uid = json.loads(pairs.get("user", "{}")).get("id")
    if not role_of(uid):
        raise HTTPException(403, f"Нет доступа (ID {uid}). Напиши боту /start — системный админ получит запрос и выдаст роль.")
    return uid


def writer(uid: int = Depends(user_id)) -> int:
    from .roles import can_write
    if not can_write(role_of_(uid)):
        raise HTTPException(403, "Роль «1С оператор» — только просмотр")
    return uid


def sysadmin(uid: int = Depends(user_id)) -> int:
    if role_of_(uid) != "sys":
        raise HTTPException(403, "Только для системного супер-админа")
    return uid


def role_of_(uid):
    from .roles import role_of
    return role_of(uid)


# ---------- schemas -------------------------------------------------------------------
class TopUpIn(BaseModel):
    date: str
    rub: float
    usd: float
    prev_balance_rub: float | None = None
    note: str = ""


class LineIn(BaseModel):
    name: str
    boxes: float | None = None
    stems: float = 0
    weight_kg: float | None = None
    price_usd: float = 0
    mrc: float | None = None


class InvoiceIn(BaseModel):
    topup_id: int = 0                  # 0 = not paid yet (груз в пути)
    est_usd: float | None = None       # unpaid: approximate $
    farm_usd: float | None = None      # $ that reached the farm (None = exactly the invoice total)
    country: str = ""
    client_code: str = ""              # empty -> default marking (LUMEN)
    invoice_no: str = ""
    invoice_date: str = ""
    awb: str = ""
    farm: str = ""
    weight_kg: float | None = None
    invoice_total_usd: float | None = None
    usd_paid: float = 0
    rub_paid_override: float | None = None
    alloc_mode: str = "value"
    paid: bool = True
    paid_date: str = ""
    note: str = ""
    source_file: str | None = None
    lines: list[LineIn] = []


class LogisticsIn(BaseModel):
    topup_id: int | None = None
    awb: str
    leg: str = "air"
    provider: str = ""
    invoice_no: str = ""
    usd: float | None = None
    rub: float | None = None
    basis: str = "auto"
    paid_date: str = ""
    note: str = ""
    source_file: str | None = None
    farm_kg: dict[str, float] = {}     # optional: set weight_kg on farm invoices of this AWB
    weight_kg: float | None = None     # weight on the bill
    paid: bool = True                  # False = deferred (Expolanka)


class FarmIn(BaseModel):
    name: str
    country: str
    aliases: str = ""
    is_forwarder: bool = False
    notes: str = ""
    box_kg_json: str = "{}"
    box_dims_json: str = "{}"


# ---------- helpers ------------------------------------------------------------------
def _all(s):
    from . import calc
    farms = s.exec(select(Farm)).all()
    calc.RULES = {f.name.strip().lower(): {"account": (getattr(f, "account", "") or "").strip(),
                                           "in_fee": getattr(f, "in_fee_pct", 0) or 0, "markup": getattr(f, "markup_pct", 0) or 0}
                  for f in farms if (getattr(f, "account", "") or getattr(f, "in_fee_pct", 0) or getattr(f, "markup_pct", 0))}
    calc.OPENINGS = {((getattr(f, "account", "") or f.name).strip().lower()): (f.opening_usd or 0, f.opening_rate, f.name)
                     for f in farms if f.opening_usd}
    from .models import BrokerDeposit
    calc.DEPOSITS = s.exec(select(BrokerDeposit)).all()
    bo = _settings().get("broker_opening")
    if bo and "брокер" not in calc.OPENINGS:
        calc.OPENINGS["брокер"] = (bo.get("usd") or 0, bo.get("rate"), "Брокер")
    return (s.exec(select(TopUp)).all(), s.exec(select(Invoice)).all(),
            s.exec(select(Line)).all(), s.exec(select(Logistics)).all())


def _norm_name(x: str) -> str:
    return re.sub(r"[^a-zа-я0-9]", "", (x or "").lower())


def _farm_keys(s, farm: str) -> set:
    """All spellings of a farm: its name + aliases from the directory."""
    keys = {_norm_name(farm)}
    for f in s.exec(select(Farm)).all():
        names = [f.name] + [a for a in f.aliases.split(",") if a.strip()]
        if _norm_name(farm) in {_norm_name(n) for n in names}:
            keys |= {_norm_name(n) for n in names}
    return keys


def _kg_for(s, farm: str, farm_kg: dict):
    keys = _farm_keys(s, farm)
    for k, v in farm_kg.items():
        nk = _norm_name(k)
        if nk in keys or any(nk.startswith(x) or x.startswith(nk) for x in keys if len(x) >= 4 and len(nk) >= 4):
            return v
    return None


# ---------- active top-up (where documents from the chat go) -------------------------------
SETTINGS = DATA_DIR / "settings.json"


def _settings() -> dict:
    try:
        return json.loads(SETTINGS.read_text())
    except (FileNotFoundError, ValueError):
        return {}


def set_active_topup(tid: int):
    SETTINGS.write_text(json.dumps({**_settings(), "active_topup": tid}))
    from .backup import mark_dirty
    mark_dirty()


def active_topup():
    """The chosen top-up, else the newest one."""
    with session() as s:
        t = s.get(TopUp, _settings().get("active_topup") or 0)
        return t or s.exec(select(TopUp).order_by(TopUp.id.desc())).first()


NUM = r"\d{1,3}(?:[ \u00a0]\d{3})+(?:[.,]\d+)?|\d+(?:[.,]\d+)?"
USD_RE = re.compile(rf"\$\s*({NUM})|(?<![\d.,])({NUM})\s*(?:\$|usd\b|долл)", re.I)
RUB_RE = re.compile(rf"(?<![\d.,])({NUM})\s*(?:₽|р\b|р\.|руб|rub\b)", re.I)


def _num(x: str) -> float:
    return float(x.replace(" ", "").replace("\u00a0", "").replace(",", "."))


def money_from_text(text: str):
    """'1198$ 105 472₽' -> (1198.0, 105472.0). Either may be None."""
    text = text or ""
    u, r = USD_RE.search(text), RUB_RE.search(text)
    usd = _num(u.group(1) or u.group(2)) if u else None
    rub = _num(r.group(1)) if r else None
    return usd, rub


def topup_from_text(text: str):
    """A date like 23.09 in the caption picks that top-up."""
    text = RUB_RE.sub(" ", USD_RE.sub(" ", text or ""))
    with session() as s:
        tops = s.exec(select(TopUp).order_by(TopUp.id.desc())).all()
        for m in re.finditer(r"(?<![\d.])(\d{1,2})\.(\d{2})(?:\.(\d{2,4}))?(?![\d])", text):
            d, mth = int(m.group(1)), int(m.group(2))
            if 1 <= d <= 31 and 1 <= mth <= 12:
                for t in tops:
                    if t.date.startswith(f"{d:02d}.{mth:02d}"):
                        return t
    return None


def leg_of(provider: str, fallback: str | None = None) -> str:
    p = (provider or "").lower()
    if "flora" in p or "флора" in p:
        return "msk"       # Floratrack
    if "expolanka" in p or "экспо" in p:
        return "air"       # Expolanka
    return fallback or "air"


def _apply_box_rules(body):
    """Farm rule 'these items come in a 25 kg box' -> kg on those lines (shared box split by stems)."""
    with session() as s:
        keys = _farm_keys(s, body.farm)
        rules = {}
        for f in s.exec(select(Farm)).all():
            if _norm_name(f.name) in keys:
                rules = json.loads(f.box_kg_json or "{}")
    if not rules:
        return
    def with_boxmates(first):
        """Invoice convention: a mixed box shows its box count on the first line only;
        the following lines with no box count are in that same box (Fire Works + Salinero)."""
        out, on = [], False
        for l in body.lines:
            if l in first:
                out.append(l); on = True
            elif on and not l.boxes and l.stems:
                out.append(l)
            else:
                on = False
        return out

    for item, kg in rules.items():
        ni = _norm_name(item)
        hit = [l for l in body.lines if l.stems and (ni in _norm_name(l.name) or _norm_name(l.name) in ni)]
        hit = with_boxmates(hit)
        if not hit or any(l.weight_kg for l in hit):
            continue
        rules[item] = (kg, hit)
    groups = {}                     # lines of different rules with the same kg share boxes (one mixed box)
    for item, v in rules.items():
        if isinstance(v, tuple):
            kg, hit = v
            groups.setdefault(kg, []).extend(l for l in hit if l not in groups.get(kg, []))
    for kg, hit in groups.items():
        boxes = max(1.0, sum(l.boxes or 0 for l in hit))
        total = boxes * kg
        st = sum(l.stems for l in hit) or 1
        for l in hit:
            l.weight_kg = round(total * l.stems / st, 3)


def split_by_farm(out: dict) -> list[dict]:
    """One trader invoice (NextWave) covering several farms -> one sub-invoice per farm."""
    lines = out.get("lines") or []
    farms = []
    for l in lines:
        f = (l.get("farm") or out.get("farm") or "").strip()
        if f not in farms:
            farms.append(f)
    if len(farms) <= 1:
        if farms and farms[0] and not out.get("farm"):
            out["farm"] = farms[0]
        return [out]
    subs = []
    for f in farms:
        ls = [l for l in lines if (l.get("farm") or "").strip() == f]
        d = {**out, "farm": f or out.get("farm"), "lines": ls,
             "invoice_total_usd": round(sum((l.get("stems") or 0) * (l.get("price_usd") or 0) for l in ls), 2),
             "subtotal_usd": None, "fees_usd": None, "weight_kg": None, "mawb_note": None,
             "note": f"общий инвойс {out.get('invoice_no') or ''} ({', '.join(x for x in farms if x)})".strip(),
             "warnings": [w for w in out.get("warnings", []) if "плантац" not in w.lower()]}
        subs.append(d)
    return subs


def payments_for(text: str, subs: list[dict]):
    """'Agriflora 301$ 26500₽\nMassai 306$ 26940₽' (or just two lines in order) -> [(usd, rub)] per sub.
    One pair for several farms -> split by invoice value."""
    segs = [s for s in re.split(r"[\n;]+", text or "") if any(money_from_text(s))]
    pairs = [money_from_text(s) for s in segs]
    if not pairs:
        return None
    if len(subs) == 1:
        return [pairs[0]]
    if len(pairs) == 1:          # one payment for all of them
        usd, rub = pairs[0]
        tot = sum(d.get("invoice_total_usd") or 0 for d in subs) or 1
        return [(round(usd * (d.get("invoice_total_usd") or 0) / tot, 2) if usd else None,
                 round(rub * (d.get("invoice_total_usd") or 0) / tot, 2) if rub else None) for d in subs]
    out = [None] * len(subs)
    rest = []
    for seg, pr in zip(segs, pairs):
        low = _norm_name(seg)
        with session() as s:
            hit = next((i for i, d in enumerate(subs) if out[i] is None and d.get("farm")
                        and any(k[:4] in low for k in _farm_keys(s, d["farm"]) if len(k) >= 4)), None)
        if hit is None:
            rest.append(pr)
        else:
            out[hit] = pr
    for i in range(len(out)):    # unnamed lines go in order
        if out[i] is None and rest:
            out[i] = rest.pop(0)
    return out


def book_document(out: dict, topup_id: int, usd, rub, uid: int, paid: bool = True, farm_usd: float | None = None):
    """Farm invoice / freight bill -> straight into the books, no Mini App.
    paid=True: into top-up `topup_id` with the $ (and ₽, else $ × rate) actually paid.
    paid=False: груз в пути — farm invoice with ≈$ (usd), freight bill with its own $.
    Returns (snapshot, what) or (None, reason)."""
    from datetime import date
    today = date.today().strftime("%d.%m.%Y")
    if not paid:
        topup_id, rub = 0, None
    if paid and usd and not rub:
        with session() as s:
            rub = round(usd * rate_of(s.get(TopUp, topup_id)))   # ₽ not given: $ × top-up rate
        out["_rub_auto"] = rub
    if out.get("doc_type") == "farm_invoice":
        lines = [LineIn(name=l["name"], boxes=l.get("boxes"), stems=l.get("stems") or 0,
                        weight_kg=None, price_usd=l.get("price_usd") or 0)       # kg on farm invoices is never right
                 for l in out.get("lines", []) if l.get("name") and l.get("stems")]
        if not lines or not out.get("farm"):
            return None, "не распознаны строки или плантация"
        with session() as s:
            f = s.exec(select(Farm)).all()
            country = out.get("country") or next((x.country for x in f if x.name.lower() == out["farm"].lower()), "")
        body = InvoiceIn(topup_id=topup_id, country=country, invoice_no=out.get("invoice_no") or "",
                         invoice_date=out.get("invoice_date") or "", awb=out.get("awb") or "", farm=out["farm"],
                         weight_kg=out.get("weight_kg") if out.get("mawb_note") else None,   # only from a breakdown
                         invoice_total_usd=out.get("invoice_total_usd"),
                         usd_paid=usd if paid else 0, rub_paid_override=rub, paid_date=today if paid else "",
                         paid=paid, est_usd=None if paid else usd, farm_usd=farm_usd if paid else None,
                         note=out.get("mawb_note") or "", source_file=out.get("source_file"), lines=lines)
        return save_invoice(body, None, uid), "invoice"
    if out.get("doc_type") == "freight_invoice":
        fr = out.get("freight") or {}
        prov = (fr.get("provider") or out.get("farm") or "")
        body = LogisticsIn(topup_id=topup_id or None, paid=paid, awb=out.get("awb") or "", leg=leg_of(prov, fr.get("leg")),
                           weight_kg=out.get("total_weight_kg"),
                           provider=fr.get("provider") or out.get("farm") or "", invoice_no=out.get("invoice_no") or "",
                           usd=usd, rub=rub, paid_date=today if paid else "", source_file=out.get("source_file"),
                           farm_kg={x["farm"]: x["kg"] for x in fr.get("per_farm_kg") or [] if x.get("kg")})
        if not body.awb:
            return None, "нет MAWB"
        return save_logistics(body, None, topup_id or None, uid), "freight"
    return None, "этот тип документа сразу не вносится"


TEXT_KG_RE = re.compile(r"([A-Za-zА-Яа-яЁё][A-Za-zА-Яа-яЁё .\-]{1,40}?)\s*[:\-–]?\s*(\d+(?:[.,]\d+)?)\s*(?:кг|kg)?(?=\s|$|[,;\n])", re.I)


def breakdown_from_text(text: str):
    """'706-52324005 Agriflora 45 Massai 57' (any layout) -> (mawb, [{farm, kg}])."""
    from .ai import find_mawb, MAWB_RE
    awb = find_mawb(text)
    if not awb:
        return None, []
    rest = MAWB_RE.sub(" ", text)
    pairs = [{"farm": n.strip(" -:"), "kg": float(k.replace(",", "."))} for n, k in TEXT_KG_RE.findall(rest)
             if n.strip(" -:").lower() not in ("mawb", "awb", "авб")]
    return awb, pairs


# ---------- volumetric weights (Expolanka) ------------------------------------------------
def resolve_farm(s, name: str, create_country: str | None = None):
    """'Zee flora' / 'Maasai' / 'Redlands' -> the Farm row (alias / prefix match); optionally create."""
    n = _norm_name(name)
    if not n:
        return None
    farms = s.exec(select(Farm)).all()
    for f in farms:
        keys = {_norm_name(x) for x in [f.name] + (f.aliases or "").split(",") if x.strip()}
        if n in keys or any(len(k) >= 4 and len(n) >= 4 and (k.startswith(n) or n.startswith(k)) for k in keys):
            return f
    if n.startswith("zee"):
        return next((f for f in farms if f.name == "Zeeflora"), None)
    if create_country:
        f = Farm(name=name.strip().title(), country=create_country, aliases=name.strip().upper())
        s.add(f); s.flush()
        return f
    return None


def learn_dims(entries: list, create_country: str | None = "Кения"):
    """Remember which box sizes each farm ships (registry of usual boxes)."""
    from .volumetric import REGISTRY_ENABLED
    if not REGISTRY_ENABLED:
        return
    with session() as s:
        for e in entries:
            f = resolve_farm(s, e["farm"], create_country)
            if not f:
                continue
            d = json.loads(f.box_dims_json or "{}")
            for n, key, *_ in e["dims"]:
                if len(_) and _[-1] == "обычная коробка":
                    continue
                d[key] = d.get(key, 0) + n
            f.box_dims_json = json.dumps(d)
            s.add(f)
        s.commit()


def usual_dims(name: str):
    from .volumetric import usual
    with session() as s:
        f = resolve_farm(s, name)
        return usual(json.loads(f.box_dims_json or "{}")) if f else None


def volumetric_breakdown(text: str) -> list:
    """WhatsApp weights message -> [{farm (canonical), boxes, kg, dims}] by L×W×H/6000."""
    from .volumetric import parse_message
    from .volumetric import REGISTRY_ENABLED
    out = parse_message(text, usual_dims=usual_dims if REGISTRY_ENABLED else None)
    with session() as s:
        for e in out:
            f = resolve_farm(s, e["farm"])
            e["raw"] = e["farm"]
            if f:
                e["farm"] = f.name
    return out


def import_floratrack(data: bytes) -> dict:
    """Floratrack xlsx -> one Floratrack logistics record per AWB, ₽ at the rate we actually paid.
    Re-uploading the next report updates the same records (rates firm up once payments arrive)."""
    from . import floratrack as ft
    from .ai import find_mawb
    from .cbr import floratrack_rate
    rep = ft.parse(data, floratrack_rate())
    added = updated = 0
    matched, unmatched, ambiguous = [], [], []
    with session() as s:
        invs = s.exec(select(Invoice)).all()
        weights = {w.awb: json.loads(w.farm_kg_json or "{}") for w in s.exec(select(AwbWeights)).all()}
        expo = {}
        for lg in s.exec(select(Logistics)).all():
            if lg.leg == "air" and lg.weight_kg:
                expo[norm_awb(lg.awb)] = expo.get(norm_awb(lg.awb), 0) + lg.weight_kg
        known = {}
        for i in invs:
            k = norm_awb(i.awb)
            if k:
                known.setdefault(k, {"countries": set(), "kg": 0})["countries"].add(i.country)
        for k, bd in weights.items():
            known.setdefault(k, {"countries": {"Кения"}, "kg": 0})["kg"] = sum(float(v) for v in bd.values() if v)
        for k, kg in expo.items():
            known.setdefault(k, {"countries": {"Кения"}, "kg": 0})
            known[k]["kg"] = known[k]["kg"] or kg
        want = {"import": "Кения", "ecuador": "Эквадор", "colombia": "Колумбия"}
        for ch in rep.charges:
            cands = [k for k in known if k.isdigit() and k.endswith(ch.last4)]
            same = [k for k in cands if want[ch.kind] in known[k]["countries"] or not any(known[k]["countries"])]
            cands = same or cands
            if len(cands) > 1:   # prefer the one whose kg agrees
                cands.sort(key=lambda k: abs((known[k]["kg"] or 1e9) - ch.kg))
                ambiguous.append(f"…{ch.last4} ({ch.sheet})")
            if not cands:
                unmatched.append(ch)
                continue
            awb = find_mawb(cands[0]) or cands[0]
            if ch.kind == "import" and known[cands[0]]["kg"] and abs(known[cands[0]]["kg"] - ch.kg) > max(2, 0.05 * ch.kg):
                rep.warnings.append(f"MAWB {awb}: у Floratrack {ch.kg:g} кг, у нас {known[cands[0]]['kg']:g} кг — проверь")
            key = f"ft:{ch.sheet}:{ch.last4}"
            lg = s.exec(select(Logistics).where(Logistics.ext_key == key)).first()
            if lg:
                updated += 1
            else:
                lg = Logistics(ext_key=key, awb=awb)
                added += 1
            lg.awb, lg.leg, lg.provider = awb, "msk", "Floratrack"
            lg.topup_id, lg.usd, lg.rub, lg.weight_kg = None, round(ch.usd, 2), ch.rub, ch.kg
            lg.invoice_no = ch.sheet
            lg.paid_date = ch.date.strftime("%d.%m.%Y") if ch.date else ""
            lg.note = ch.rate_note
            s.add(lg)
            matched.append((awb, ch))
        s.commit()
    from .backup import mark_dirty
    mark_dirty()
    return {"added": added, "updated": updated, "matched": matched, "unmatched": unmatched,
            "ambiguous": ambiguous, "warnings": rep.warnings, "balance_usd": rep.balance_usd}


def infer_mawb(farm: str, topup_id: int | None = None):
    """Farm invoices from Kenya carry no MAWB. The Expolanka breakdown does: find the MAWB whose
    breakdown lists this farm and that doesn't have this farm's invoice yet. Newest first.
    Returns (mawb, kg, n_candidates) or (None, None, 0)."""
    from sqlalchemy import text
    if not farm:
        return None, None, 0
    with session() as s:
        rows = s.exec(text("SELECT awb, farm_kg_json FROM awbweights ORDER BY rowid DESC")).all()
        taken = {(norm_awb(i.awb), _norm_name(i.farm)) for i in s.exec(select(Invoice)).all()}
        keys = _farm_keys(s, farm)
        cands = []
        for awb, js in rows:
            kg = _kg_for(s, farm, json.loads(js or "{}"))
            if kg and not any((awb, k) in taken for k in keys):
                cands.append((awb, kg))
        if topup_id and len(cands) > 1:
            # MAWBs whose freight is paid from this top-up go first
            paid_here = {norm_awb(l.awb) for l in s.exec(select(Logistics)).all() if l.topup_id == topup_id}
            cands.sort(key=lambda c: c[0] not in paid_here)
    if not cands:
        return None, None, 0
    a, kg = cands[0]
    return f"{a[:3]}-{a[3:]}", kg, len(cands)


def _apply_farm_kg(s, awb, farm_kg):
    """Forwarder breakdown -> weight_kg of every farm invoice already on this AWB."""
    if not farm_kg:
        return
    key = norm_awb(awb)
    for inv in s.exec(select(Invoice)).all():
        if norm_awb(inv.awb) == key:
            v = _kg_for(s, inv.farm, farm_kg)
            if v:
                inv.weight_kg = v
                s.add(inv)


def _awb_breakdown(s) -> dict:
    """{MAWB digits: {farm: kg}} — for the Excel: ТК weight broken down by plantation."""
    return {w.awb: {k: float(v) for k, v in json.loads(w.farm_kg_json or "{}").items() if v}
            for w in s.exec(select(AwbWeights)).all()}


def marking() -> str:
    """Default marking (код клиента) for new invoices. LUMEN until client markings are sorted out."""
    return (_settings().get("marking") or "LUMEN").strip()


def _awb_kg(s) -> dict:
    """{MAWB: total kg of its breakdown} for the calc engine."""
    out = {}
    for w in s.exec(select(AwbWeights)).all():
        out[w.awb] = sum(float(v) for v in json.loads(w.farm_kg_json or "{}").values() if v)
    return out


def _store_weights(s, awb, farm_kg, source_file=None, replace=False):
    """Save a per-farm kg breakdown for a MAWB and push it onto the farm invoices.
    replace=True: a full breakdown document replaces the old one; False: merge (manual edits)."""
    farm_kg = {k: float(v) for k, v in (farm_kg or {}).items() if v not in (None, "", 0)}
    if not awb or not farm_kg:
        return
    key = norm_awb(awb)
    row = s.get(AwbWeights, key) or AwbWeights(awb=key)
    base = {} if replace else json.loads(row.farm_kg_json or "{}")
    row.farm_kg_json = json.dumps({**base, **farm_kg}, ensure_ascii=False)
    if source_file:
        row.source_file = source_file
    s.add(row)
    from .ai import find_mawb
    for inv in s.exec(select(Invoice).order_by(Invoice.id.desc())).all():
        if not inv.awb and _kg_for(s, inv.farm, farm_kg):
            taken = any(norm_awb(i.awb) == key and _norm_name(i.farm) == _norm_name(inv.farm)
                        for i in s.exec(select(Invoice)).all())
            if not taken:
                inv.awb = find_mawb(awb) or awb
                s.add(inv)
    s.flush()
    _apply_farm_kg(s, awb, farm_kg)


def _fill_weight_from_logistics(s, inv):
    """Invoice added AFTER the Expolanka bill: take its kg from the stored breakdown."""
    if inv.weight_kg or not inv.awb:
        return
    row = s.get(AwbWeights, norm_awb(inv.awb))
    sources = [json.loads(row.farm_kg_json or "{}")] if row else []
    sources += [json.loads(lg.farm_kg_json or "{}") for lg in s.exec(select(Logistics)).all()
                if norm_awb(lg.awb) == norm_awb(inv.awb)]          # older data kept on the freight record
    for bd in sources:
        v = _kg_for(s, inv.farm, bd)
        if v:
            inv.weight_kg = v
            s.add(inv)
            return


def snapshot(s, topup_id):
    if not topup_id:
        return transit_view(s)
    tops, invs, lines, logs = _all(s)
    t = next((x for x in tops if x.id == topup_id), None)
    if not t:
        raise HTTPException(404)
    res = compute(topup_id, tops, invs, lines, logs, _awb_kg(s))
    lines_by = defaultdict(list)
    for l in lines:
        lines_by[l.invoice_id].append(l)
    out_inv = []
    for inv in (i for i in invs if i.topup_id == topup_id):
        ls = []
        for l in lines_by[inv.id]:
            c = res.lines[l.id]
            ls.append({**l.model_dump(), "price_rub": round(c.price_rub, 2), "air_rub": round(c.air_rub_stem, 2),
                       "msk_rub": round(c.msk_rub_stem, 2), "total_rub": round(c.total_rub_stem, 2)})
        out_inv.append({**inv.model_dump(), "rub_paid": res.invoice_rub[inv.id], "lines": ls,
                        "true_rate": round(res.true_rate.get(inv.id, 0), 4), "cost_pct": round(res.cost_pct.get(inv.id, 0), 1),
                        "rub_cost": round(res.invoice_cost.get(inv.id, 0)), "ledger": res.inv_ledger.get(inv.id)})
    related = {norm_awb(i.awb) for i in invs if i.topup_id == topup_id}
    kg = _awb_kg(s)
    rpk = {k: g.get("rub_per_kg") for k, g in res.legs.items()}
    tdate = {x.id: x.date for x in tops}
    goods = {}
    for i in invs:
        goods.setdefault(norm_awb(i.awb), set()).add(tdate.get(i.topup_id, "?"))
    out_log = [{**lg.model_dump(), "kg_total": kg.get(norm_awb(lg.awb), 0), "rub_per_kg": rpk.get((norm_awb(lg.awb), lg.leg)),
                "goods_topups": sorted(goods.get(norm_awb(lg.awb), set())),
                "provisional": "предварительн" in (lg.note or "")}
               for lg in logs if lg.topup_id == topup_id or norm_awb(lg.awb) in related]
    act = _settings().get("active_topup")
    return {"topup": {**t.model_dump(), "rate": rate_of(t), "active": t.id == act}, "invoices": out_inv, "logistics": out_log,
            "warnings": res.warnings}


def history(s):
    tops, invs, lines, logs = _all(s)
    rates = [{"date": t.date, "rate": round(rate_of(t), 4)} for t in tops]
    per = defaultdict(list)
    for t in tops:
        res = compute(t.id, tops, invs, lines, logs, _awb_kg(s))
        inv_by = {i.id: i for i in invs if i.topup_id == t.id}
        for l in lines:
            if l.invoice_id in inv_by:
                per[f"{inv_by[l.invoice_id].farm} | {l.name}"].append(res.lines[l.id].total_rub_stem)
    typical = {k: round(median(v), 2) for k, v in per.items() if v}
    return {"rates": rates[-15:], "typical_cost_per_stem_rub": dict(list(typical.items())[-300:])}


# ---------- routes -------------------------------------------------------------------
@router.get("/topups")
def list_topups(uid: int = Depends(user_id)):
    from .roles import can_write
    if not can_write(role_of_(uid)):                 # 1С operator: only dates, no money
        with session() as s:
            return [{"id": t.id, "date": t.date} for t in s.exec(select(TopUp).order_by(TopUp.id.desc())).all()]
    active = active_topup()
    with session() as s:
        tops, invs, lines, logs = _all(s)
        out = []
        for t in sorted(tops, key=lambda x: x.id, reverse=True):
            spent = sum(i.usd_paid for i in invs if i.topup_id == t.id) + sum(l.usd or 0 for l in logs if l.topup_id == t.id)
            out.append({**t.model_dump(), "rate": rate_of(t),
                        "active": bool(active) and t.id == active.id,
                        "n_invoices": sum(1 for i in invs if i.topup_id == t.id)})
        return out


def money_of(tid: int) -> dict:
    with session() as s:
        tops, invs, lines, logs = _all(s)
        t = s.get(TopUp, tid)
        res = compute(tid, tops, invs, lines, logs, _awb_kg(s))
        return {"usd_spent": round(res.usd_spent, 2), "usd_left": round(t.usd - res.usd_spent, 2)}


@router.get("/money")
def money(uid: int = Depends(sysadmin)):
    """Остатки — только системному супер-админу."""
    with session() as s:
        ids = [t.id for t in s.exec(select(TopUp)).all()]
    return {tid: money_of(tid) for tid in ids}


@router.get("/me")
def me(uid: int = Depends(user_id)):
    from .roles import ROLE_NAMES, can_write, sees_money
    r = role_of_(uid)
    return {"id": uid, "role": r, "role_name": ROLE_NAMES.get(r, r), "can_write": can_write(r), "sees_money": sees_money(r)}


class UserIn(BaseModel):
    tg_id: int
    name: str = ""
    role: str


@router.get("/users")
def list_users(uid: int = Depends(sysadmin)):
    from .roles import users
    return users()


@router.post("/users")
def save_user(body: UserIn, uid: int = Depends(sysadmin)):
    from .roles import set_role, users
    try:
        set_role(body.tg_id, body.role, body.name)
    except ValueError as e:
        raise HTTPException(400, str(e))
    from .backup import mark_dirty
    mark_dirty()
    return users()


@router.delete("/users/{tg}")
def drop_user(tg: int, uid: int = Depends(sysadmin)):
    from .roles import remove, users
    try:
        remove(tg)
    except ValueError as e:
        raise HTTPException(400, str(e))
    from .backup import mark_dirty
    mark_dirty()
    return users()


@router.post("/topups")
def create_topup(body: TopUpIn, uid: int = Depends(writer)):
    with session() as s:
        t = TopUp(**body.model_dump())
        s.add(t); s.commit(); s.refresh(t)
    set_active_topup(t.id)            # a new top-up becomes the one chat documents go to
    return t


@router.post("/topups/{tid}/activate")
def activate_topup(tid: int, uid: int = Depends(writer)):
    set_active_topup(tid)
    return {"ok": True}


@router.delete("/topups/{tid}")
def delete_topup(tid: int, uid: int = Depends(writer)):
    """Delete a top-up with everything booked into it (its invoices + lines, freight paid from it)."""
    with session() as s:
        for inv in s.exec(select(Invoice).where(Invoice.topup_id == tid)).all():
            for l in s.exec(select(Line).where(Line.invoice_id == inv.id)).all():
                s.delete(l)
            s.delete(inv)
        for lg in s.exec(select(Logistics).where(Logistics.topup_id == tid)).all():
            s.delete(lg)
        t = s.get(TopUp, tid)
        if t:
            s.delete(t)
        s.commit()
    if _settings().get("active_topup") == tid:
        SETTINGS.write_text(json.dumps({**_settings(), "active_topup": None}))   # falls back to the newest
    return {"ok": True}


@router.put("/topups/{tid}")
def update_topup(tid: int, body: TopUpIn, uid: int = Depends(writer)):
    with session() as s:
        t = s.get(TopUp, tid)
        for k, v in body.model_dump().items():
            setattr(t, k, v)
        s.add(t); s.commit()
        return snapshot(s, tid)


@router.get("/topups/{tid}")
def get_topup(tid: int, uid: int = Depends(user_id)):
    from .roles import can_write
    with session() as s:
        if not can_write(role_of_(uid)):             # 1С operator: date only, export is allowed
            t = s.get(TopUp, tid)
            if not t:
                raise HTTPException(404)
            return {"topup": {"id": t.id, "date": t.date}, "viewer": True}
        return snapshot(s, tid)


@router.post("/invoices")
def save_invoice(body: InvoiceIn, inv_id: int | None = None, uid: int = Depends(writer)):
    from .ai import find_mawb
    body.client_code = (body.client_code or "").strip() or marking()
    with session() as s0:                                # one name per farm: «Positano» -> «Tessa», «ZEEFLORA LTD» -> «Zeeflora»
        fm = resolve_farm(s0, body.farm)
        if fm:
            body.farm = fm.name
    body.awb = find_mawb(body.awb) or body.awb.strip()   # one spelling everywhere: 065-40538245
    _apply_box_rules(body)
    prices = {round(l.price_usd, 4) for l in body.lines if l.stems}
    if not inv_id and len(prices) == 1 and len(body.lines) > 1:
        body.alloc_mode = "stems"    # one price for everything (hydrangeas): ₽ paid / stems
    if not body.awb:
        body.awb = infer_mawb(body.farm)[0] or ""      # typed by hand without MAWB -> take it from the breakdown
    with session() as s:
        data = body.model_dump(exclude={"lines"})
        inv = s.get(Invoice, inv_id) if inv_id else Invoice(**data)
        if inv_id:
            for k, v in data.items():
                setattr(inv, k, v)
            for l in s.exec(select(Line).where(Line.invoice_id == inv_id)).all():
                s.delete(l)
        s.add(inv); s.commit(); s.refresh(inv)
        for l in body.lines:
            s.add(Line(invoice_id=inv.id, **l.model_dump()))
        _fill_weight_from_logistics(s, inv)
        s.commit()
        if not inv.topup_id:
            inv.paid = False
            s.add(inv); s.commit()
        return snapshot(s, inv.topup_id)


@router.delete("/invoices/{inv_id}")
def delete_invoice(inv_id: int, uid: int = Depends(writer)):
    with session() as s:
        inv = s.get(Invoice, inv_id)
        for l in s.exec(select(Line).where(Line.invoice_id == inv_id)).all():
            s.delete(l)
        tid = inv.topup_id
        s.delete(inv); s.commit()
        return snapshot(s, tid)


@router.post("/logistics")
def save_logistics(body: LogisticsIn, log_id: int | None = None, view_topup: int | None = None,
                   uid: int = Depends(writer)):
    with session() as s:
        from .ai import find_mawb
        body.awb = find_mawb(body.awb) or body.awb.strip()
        body.leg = leg_of(body.provider, body.leg)      # Expolanka -> air, Floratrack -> msk
        data = body.model_dump(exclude={"farm_kg"})
        lg = s.get(Logistics, log_id) if log_id else Logistics(**data)
        if log_id:
            for k, v in data.items():
                setattr(lg, k, v)
        s.add(lg)
        _store_weights(s, body.awb, body.farm_kg, replace=True)
        s.commit()
        return snapshot(s, view_topup or body.topup_id or 0)


@router.delete("/logistics/{log_id}")
def delete_logistics(log_id: int, view_topup: int, uid: int = Depends(writer)):
    with session() as s:
        s.delete(s.get(Logistics, log_id)); s.commit()
        return snapshot(s, view_topup)


@router.get("/awb/{awb}")
def awb_invoices(awb: str, uid: int = Depends(user_id)):
    with session() as s:
        k = norm_awb(awb)
        return [{"id": i.id, "farm": i.farm, "weight_kg": i.weight_kg, "topup_id": i.topup_id}
                for i in s.exec(select(Invoice)).all() if norm_awb(i.awb) == k]


class WeightsIn(BaseModel):
    awb: str
    farm_kg: dict[str, float] = {}
    source_file: str | None = None


def weights_now(awb: str) -> dict:
    with session() as s:
        row = s.get(AwbWeights, norm_awb(awb))
        k = norm_awb(awb)
        invs = [{"id": i.id, "farm": i.farm, "weight_kg": i.weight_kg, "topup_id": i.topup_id}
                for i in s.exec(select(Invoice)).all() if norm_awb(i.awb) == k]
        return {"awb": awb, "farm_kg": json.loads(row.farm_kg_json) if row else {}, "invoices": invs}


def store_breakdown(awb: str, per_farm_kg: list, source_file=None) -> dict:
    """Used by the bot: a breakdown with a MAWB is applied straight away, no draft needed."""
    farm_kg = {x["farm"]: x["kg"] for x in per_farm_kg or [] if x.get("farm") and x.get("kg")}
    with session() as s:
        _store_weights(s, awb, farm_kg, source_file, replace=True)
        s.commit()
    from .backup import mark_dirty
    mark_dirty()
    return weights_now(awb)


@router.get("/weights/{awb}")
def _get_weights_alias(awb: str, uid: int = Depends(user_id)):
    return weights_now(awb)


def get_weights(awb: str, uid: int = Depends(user_id)):
    """Breakdown for a MAWB + the farm invoices already on it."""
    with session() as s:
        row = s.get(AwbWeights, norm_awb(awb))
        k = norm_awb(awb)
        invs = [{"id": i.id, "farm": i.farm, "weight_kg": i.weight_kg, "topup_id": i.topup_id}
                for i in s.exec(select(Invoice)).all() if norm_awb(i.awb) == k]
        return {"awb": awb, "farm_kg": json.loads(row.farm_kg_json) if row else {}, "invoices": invs}


@router.post("/weights")
def save_weights(body: WeightsIn, uid: int = Depends(writer)):
    with session() as s:
        _store_weights(s, body.awb, body.farm_kg, body.source_file)
        s.commit()
    return get_weights(body.awb, uid)


class MarkingIn(BaseModel):
    marking: str


@router.get("/settings/marking")
def get_marking(uid: int = Depends(user_id)):
    return {"marking": marking()}


@router.post("/settings/marking")
def set_marking(body: MarkingIn, uid: int = Depends(writer)):
    m = body.marking.strip().upper()
    if not m:
        raise HTTPException(400, "пустая маркировка")
    SETTINGS.write_text(json.dumps({**_settings(), "marking": m}))
    from .backup import mark_dirty
    mark_dirty()
    return {"marking": m}


# ---------- settlements with farms (advances / debts) ----------
def ledger_view() -> dict:
    with session() as s:
        tops, invs, lines, logs = _all(s)
        res = compute(0, tops, invs, lines, logs, _awb_kg(s))
    farms = sorted(res.ledger.values(), key=lambda f: (-abs(f["balance_usd"]), f["farm"]))
    return {"farms": [f for f in farms if abs(f["balance_usd"]) > 0.01 or f["advances"] or f["debts"]]}


def calc_rule_account(farm: str) -> str:
    from . import calc
    return (calc.RULES.get((farm or "").strip().lower()) or {}).get("account") or ""


def farm_balance_text(farm: str) -> str:
    """«Kikwetu: аванс $100 (по 95.34 ₽/$, оплата 23.09)» / «долг $150» / «расчёты закрыты»."""
    with session() as s:
        tops, invs, lines, logs = _all(s)
        res = compute(0, tops, invs, lines, logs, _awb_kg(s))
    acc = calc_rule_account(farm)
    f = next((v for k, v in res.ledger.items() if k.lower() == (farm or "").lower()
              or (acc and k.lower().startswith(acc.lower()))), None)
    if acc and f:
        farm = f["farm"]
    if not f or (abs(f["balance_usd"]) < 0.01 and not f["debts"]):
        return f"⚖️ {farm}: расчёты закрыты, аванса и долга нет"
    parts = []
    if f["advance_usd"] > 0.01:
        parts.append(f"аванс ${f['advance_usd']:,.2f}".replace(",", " ") + " (" + ", ".join(
            f"${a['usd']:g} по {a['rate']:.2f} ₽/$ от {a['from']}" for a in f["advances"]) + ")")
    if f["debt_usd"] > 0.01:
        parts.append(f"долг ${f['debt_usd']:,.2f}".replace(",", " ") + " — ≈ по последнему курсу, уточнится после оплаты")
    return f"⚖️ {farm}: " + "; ".join(parts)


def _farm_now(farm: str):
    v = ledger_view()
    acc = calc_rule_account(farm)
    return next((f for f in v["farms"] if f["farm"].lower() == (farm or "").lower()
                 or (acc and f["farm"].lower().startswith(acc.lower()))), None)


def farm_balance_projection(farm: str, total: float) -> str:
    """Before booking an unpaid invoice: «Kikwetu: сейчас аванс $48 → после поставки долг $512.20»."""
    f = _farm_now(farm)
    bal = f["balance_usd"] if f else 0.0
    now = "расчёты закрыты" if abs(bal) < 0.01 else (f"аванс ${bal:,.2f}" if bal > 0 else f"долг ${-bal:,.2f}")
    after = bal - (total or 0)
    aft = "ровно" if abs(after) < 0.01 else (f"аванс ${after:,.2f}" if after > 0 else f"долг ${-after:,.2f}")
    return f"⚖️ {farm}: сейчас {now} → после этой поставки {aft} (инвойс ${total:,.2f})".replace(",", " ")


def adjust_farm_balance(farm: str, target_usd: float):
    """«У фермы другой баланс»: make the balance BEFORE the new invoice equal target (via the opening balance)."""
    f = _farm_now(farm)
    cur = f["balance_usd"] if f else 0.0
    delta = target_usd - cur
    if abs(delta) < 0.005:
        return
    with session() as s:
        fr = resolve_farm(s, farm, create_country="Кения")
        fr.opening_usd = (fr.opening_usd or 0) + delta
        s.add(fr); s.commit()
    from .backup import mark_dirty
    mark_dirty()


@router.get("/ledger")
def get_ledger(uid: int = Depends(writer)):
    return ledger_view()


class OpeningIn(BaseModel):
    farm: str
    usd: float                 # + advance at the farm / − our debt
    rate: float | None = None


@router.post("/ledger/opening")
def set_opening(body: OpeningIn, uid: int = Depends(writer)):
    """Balance with a farm that existed BEFORE the bot (e.g. a $48 advance)."""
    with session() as s:
        f = resolve_farm(s, body.farm, create_country="Кения")
        f.opening_usd, f.opening_rate = body.usd, body.rate
        s.add(f); s.commit()
    from .backup import mark_dirty
    mark_dirty()
    return ledger_view()


# ---------- archive of every uploaded document ----------
def record_upload(uid: int, user: str, filename: str, mime: str, tg_file_id=None, local_path=None, kind="") -> int:
    from datetime import datetime, timedelta, timezone
    from .models import Upload
    ts = (datetime.now(timezone.utc) + timedelta(hours=3)).strftime("%d.%m.%Y %H:%M")
    with session() as s:
        u = Upload(ts=ts, uid=uid, user=user, filename=filename, mime=mime, tg_file_id=tg_file_id,
                   local_path=local_path, kind=kind)
        s.add(u); s.commit(); s.refresh(u)
        return u.id


def update_upload(up_id: int, kind: str = None, summary: str = None):
    from .models import Upload
    if not up_id:
        return
    with session() as s:
        u = s.get(Upload, up_id)
        if u:
            if kind is not None:
                u.kind = kind
            if summary is not None:
                u.summary = summary[:300]
            s.add(u); s.commit()
    from .backup import mark_dirty
    mark_dirty()


KIND_RU = {"farm_invoice": "инвойс фермы", "freight_invoice": "счёт ТК", "kg_breakdown": "разбивка кг",
           "topup_receipt": "пополнение", "kenya_breakdown": "детализация Кении", "floratrack": "отчёт ТК МСК",
           "master": "мастер-файл учёта", "other": "другое"}


@router.get("/uploads")
def list_uploads(uid: int = Depends(sysadmin)):
    """Archive for the system admin: who sent what, and what is left of it in the books."""
    from .models import Upload
    with session() as s:
        ups = s.exec(select(Upload).order_by(Upload.id.desc()).limit(300)).all()
        invs = s.exec(select(Invoice)).all()
        tdate = {t.id: t.date for t in s.exec(select(TopUp)).all()}
    drafts = {}
    for p in DRAFTS.glob("*.json"):
        try:
            d = json.loads(p.read_text())
            if d.get("upload_id"):
                drafts.setdefault(d["upload_id"], []).append(d.get("farm") or "")
        except ValueError:
            pass
    out = []
    for u in ups:
        booked = [i for i in invs if i.source_file == f"up:{u.id}"]
        if booked:
            status = "в учёте: " + ", ".join(f"{i.farm} ({tdate.get(i.topup_id, 'в пути')})" for i in booked)
        elif u.id in drafts:
            status = "черновик"
        elif u.kind in ("farm_invoice", "freight_invoice"):
            status = "нет в учёте — удалён или не внесён"
        else:
            status = ""
        out.append({**u.model_dump(), "kind_ru": KIND_RU.get(u.kind, u.kind or "—"), "status": status,
                    "gone": status.startswith("нет в учёте")})
    return out


@router.post("/uploads/{up_id}/send")
async def resend_upload(up_id: int, uid: int = Depends(sysadmin)):
    """Send the ORIGINAL file back into the admin's chat with the finance bot."""
    from aiogram.types import FSInputFile
    from .models import Upload
    with session() as s:
        u = s.get(Upload, up_id)
    if not u or not BOT:
        raise HTTPException(404)
    cap = f"📎 {u.filename or 'документ'} · {u.user} · {u.ts}"
    if u.tg_file_id:
        if u.mime.startswith("image/") and not u.filename:
            await BOT.send_photo(uid, u.tg_file_id, caption=cap)
        else:
            await BOT.send_document(uid, u.tg_file_id, caption=cap)
    elif u.local_path and (DATA_DIR / "files" / u.local_path).exists():
        await BOT.send_document(uid, FSInputFile(DATA_DIR / "files" / u.local_path, filename=u.filename), caption=cap)
    else:
        raise HTTPException(410, "Файл не сохранился")
    return {"ok": True}


# ---------- client markings & chats ----------
@router.get("/markings")
def list_markings(uid: int = Depends(writer)):
    from .clients import registry
    return {"markings": registry(), "default": marking()}


class MarkIn(BaseModel):
    name: str


@router.post("/markings")
def add_marking(body: MarkIn, uid: int = Depends(writer)):
    from .clients import registry, save_registry
    reg = registry()
    reg.setdefault(body.name.strip().upper(), [])
    save_registry(reg)
    return {"markings": reg}


@router.delete("/markings/{name}")
def drop_marking(name: str, chat_id: int | None = None, thread_id: int | None = None, uid: int = Depends(writer)):
    from .clients import registry, save_registry
    reg = registry()
    n = name.strip().upper()
    if chat_id is None:
        reg.pop(n, None)
    else:
        reg[n] = [c for c in reg.get(n, []) if not (c["chat_id"] == chat_id and c.get("thread_id") == thread_id)]
    save_registry(reg)
    return {"markings": reg}


class AwbMarksIn(BaseModel):
    markings: list[str]


@router.post("/awb/{awb}/markings")
def set_awb_markings(awb: str, body: AwbMarksIn, uid: int = Depends(writer)):
    """Markings present in a shipment: set on every invoice of this MAWB."""
    k = norm_awb(awb)
    val = ", ".join(dict.fromkeys(m.strip().upper() for m in body.markings if m.strip())) or marking()
    with session() as s:
        for i in s.exec(select(Invoice)).all():
            if norm_awb(i.awb) == k:
                i.client_code = val
                s.add(i)
        s.commit()
    from .backup import mark_dirty
    mark_dirty()
    return transit_view()


@router.delete("/drafts")
def drop_all_drafts(uid: int = Depends(writer)):
    for p in DRAFTS.glob("*.json"):
        p.unlink(missing_ok=True)
    from .backup import mark_dirty
    mark_dirty()
    return {"ok": True}


@router.get("/farms")
def farms(uid: int = Depends(user_id)):
    with session() as s:
        return s.exec(select(Farm).order_by(Farm.country, Farm.name)).all()


@router.post("/farms")
def save_farm(body: FarmIn, farm_id: int | None = None, uid: int = Depends(writer)):
    with session() as s:
        f = s.get(Farm, farm_id) if farm_id else Farm(**body.model_dump())
        if farm_id:
            for k, v in body.model_dump().items():
                setattr(f, k, v)
        s.add(f); s.commit(); s.refresh(f)
        return f


@router.post("/parse")
async def parse(file: UploadFile = File(...), uid: int = Depends(writer)):
    data = await file.read()
    mime = file.content_type or "application/pdf"
    if mime not in ("application/pdf", "image/jpeg", "image/png", "image/webp"):
        raise HTTPException(400, "Нужен PDF или фото (jpg/png)")
    name = f"{uuid.uuid4().hex[:10]}_{file.filename}"
    (DATA_DIR / "files" / name).write_bytes(data)
    with session() as s:
        fs = [f.model_dump() for f in s.exec(select(Farm)).all()]
        catalog = sorted({l.name for l in s.exec(select(Line)).all()})
    out = await ai.parse_document(data, mime, fs, catalog)
    up = record_upload(uid, "приложение", file.filename or name, mime, local_path=name, kind=out.get("doc_type", ""))
    update_upload(up, summary=f"{out.get('farm') or ''} {out.get('invoice_no') or ''} ${out.get('invoice_total_usd') or ''}".strip())
    out["source_file"] = f"up:{up}"
    out["upload_id"] = up
    fill_mawb(out)
    return out


@router.post("/topups/{tid}/audit")
async def audit(tid: int, uid: int = Depends(writer)):
    with session() as s:
        snap, hist = snapshot(s, tid), history(s)
    return await ai.audit_topup(snap, hist)


def _audience(uid: int, wanted: str | None) -> str:
    """viewer (1С) always gets the operator version; others choose."""
    from .roles import can_write
    if not can_write(role_of_(uid)):
        return "operator"
    return "operator" if wanted == "operator" else "owner"


@router.post("/topups/{tid}/export")
async def export(tid: int, audience: str | None = None, uid: int = Depends(user_id)):
    """Only this top-up: the file contains just its sheet."""
    with session() as s:
        t = s.get(TopUp, tid)
        if not t:
            raise HTTPException(404)
    names = await export_topups([tid], uid, "", audience=_audience(uid, audience), only=True)
    return {"ok": True, "sheet": names[0] if names else None}


def awb_spread(awb: str) -> list[dict]:
    """Farms on this MAWB (any top-up) and their logistics per stem after the latest freight."""
    k = norm_awb(awb)
    with session() as s:
        tops, invs, lines, logs = _all(s)
        kg = _awb_kg(s)
        tmap = {t.id: t for t in tops}
        out = []
        for tid in sorted({i.topup_id for i in invs if norm_awb(i.awb) == k}):
            res = compute(tid, tops, invs, lines, logs, kg)
            for i in invs:
                if i.topup_id == tid and norm_awb(i.awb) == k:
                    ls = [l for l in lines if l.invoice_id == i.id]
                    st = sum(l.stems for l in ls) or 1
                    out.append({"farm": i.farm, "topup_id": tid, "topup": tmap[tid].date, "stems": st,
                                "air": sum(res.lines[l.id].air_rub_stem * l.stems for l in ls) / st,
                                "msk": sum(res.lines[l.id].msk_rub_stem * l.stems for l in ls) / st})
        return out


def _ledger_sheet(path):
    """Owner file: a sheet with advances / debts per farm (never in the operator version)."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    try:
        wb = load_workbook(path)
    except Exception:
        return
    name = "Расчёты с фермами"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    hdr = ["Ферма", "Аванс у фермы, $", "Долг ферме, $", "Баланс, $", "Аванс: курс ₽/$ и откуда"]
    for c, h in enumerate(hdr, 1):
        x = ws.cell(1, c, h); x.font = Font(name="Arial", bold=True); x.fill = PatternFill("solid", fgColor="D9E1F2")
    for r, f in enumerate(ledger_view()["farms"], 2):
        ws.cell(r, 1, f["farm"]); ws.cell(r, 2, f["advance_usd"] or None); ws.cell(r, 3, f["debt_usd"] or None)
        ws.cell(r, 4, f["balance_usd"]).font = Font(name="Arial", bold=True, color="C00000" if f["balance_usd"] < 0 else "008000")
        ws.cell(r, 5, "; ".join(f"${a['usd']:g} по {a['rate']:.2f} от {a['from']}" for a in f["advances"]))
    for col, w in zip("ABCDE", (24, 16, 16, 14, 50)):
        ws.column_dimensions[col].width = w
    wb.save(path)


async def export_topups(ids: list[int], uid: int, caption: str, audience: str = "owner", only: bool = False):
    """Rebuild sheets and send ONE file.
    owner: sheets are rebuilt inside the master учет.xlsx (formulas live); if only=True the file sent
           contains just these sheets (+ «Пополнения», which their formulas use).
    operator: a separate file with plain numbers — no «Пополнения», no ТК share columns."""
    from aiogram.types import FSInputFile
    from openpyxl import load_workbook
    names = []
    with session() as s:
        tops, invs, lines, logs = _all(s)
        kg, bd = _awb_kg(s), _awb_breakdown(s)
        tmap = {t.id: t for t in tops}
        # master always gets the fresh formulas version
        tmp = DATA_DIR / "export_multi.xlsx"
        if MASTER_XLSX.exists():
            shutil.copy(MASTER_XLSX, tmp)
        else:
            tmp.unlink(missing_ok=True)
        for tid in ids:
            t = s.get(TopUp, tid)
            if not t:
                continue
            name, _ = excel.build(tmp, t, tops, invs, lines, logs, awb_kg=kg, awb_breakdown=bd)
            names.append(name)
            s.add(t)
        s.commit()
        _ledger_sheet(tmp)
        shutil.copy(tmp, MASTER_XLSX)
        out = tmp
        if audience == "operator":
            out = DATA_DIR / f"export_operator_{uuid.uuid4().hex[:6]}.xlsx"
            out.unlink(missing_ok=True)
            for tid in ids:
                if tid in tmap:
                    t = s.get(TopUp, tid)
                    sheet = t.sheet_name
                    excel.build(out, t, tops, invs, lines, logs, awb_kg=kg, awb_breakdown=bd, operator=True)
                    t.sheet_name = sheet               # don't let the operator copy rename the master sheet
                    s.add(t)
            s.commit()
        elif only:
            out = DATA_DIR / f"export_only_{uuid.uuid4().hex[:6]}.xlsx"
            wb = load_workbook(tmp)
            keep = set(names) | {"Пополнения", "Расчёты с фермами"}
            for n in list(wb.sheetnames):
                if n not in keep:
                    del wb[n]
            wb.save(out)
    if BOT and names:
        fname = f"учет_{names[0].replace('Пополнение ', '')}.xlsx" if len(names) == 1 else "учет.xlsx"
        if audience == "operator":
            fname = fname.replace("учет", "учет_оператор")
        cap = caption or (f"Лист «{names[0]}»" if len(names) == 1 else f"Пополнений: {len(names)}")
        if audience == "operator":
            cap += "\n👤 Версия для оператора: без пополнений и долей ТК"
        await BOT.send_document(uid, FSInputFile(out, filename=fname),
                                caption=cap + ("" if len(names) == 1 else "\nЛисты: " + ", ".join(f"«{n}»" for n in names)))
    if out != tmp:
        try:
            out.unlink()
        except Exception:
            pass
    return names


@router.post("/export_all")
async def export_all(audience: str | None = None, uid: int = Depends(user_id)):
    """Every top-up rebuilt into one file and sent to the chat."""
    with session() as s:
        ids = [t.id for t in s.exec(select(TopUp).order_by(TopUp.id)).all()]
    names = await export_topups(ids, uid, "📊 Учёт: все пополнения", audience=_audience(uid, audience))
    if BOT and names:
        from .backup import backup_now
        try:
            await backup_now(BOT, "выгрузка всех пополнений")
        except Exception as e:
            print(f"[lumen] backup after export failed: {e}", flush=True)
    return {"ok": True, "sheets": names}


# ---------- груз в пути ------------------------------------------------------------------
FT_TARIFF = {"Эквадор": 7.51, "Колумбия": 7.11}   # $/kg, consolidation door-to-Moscow


def _country_of(s, inv) -> str:
    if inv.country:
        return inv.country
    keys = _farm_keys(s, inv.farm)
    for f in s.exec(select(Farm)).all():
        if _norm_name(f.name) in keys:
            return f.country
    return ""


def _prev_leg(invs, lines, res, inv, line, leg):
    """Kenya without a bill yet: logistics per stem of the previous shipment of the SAME farm and SAME item."""
    key = "air_rub_stem" if leg == "air" else "msk_rub_stem"
    name = _norm_name(line.name)
    for p in sorted((i for i in invs if i.farm.lower() == inv.farm.lower() and i.id < inv.id), key=lambda x: -x.id):
        for l in lines:
            if l.invoice_id == p.id and _norm_name(l.name) == name and getattr(res.lines[l.id], key):
                return getattr(res.lines[l.id], key), p
    return None, None


def transit_view(s=None) -> dict:
    """Everything not arrived yet: paid (in a top-up) and unpaid, with ≈ cost per stem.
    Missing freight is estimated per leg:
      Kenya  — copy of the previous shipment of this farm & this item (Expolanka and/or Floratrack)
      Ecuador/Colombia — Floratrack tariff 7.51 / 7.11 $/kg × farm kg × (ЦБ today + 3) / 0.96"""
    from .cbr import floratrack_rate
    own = s is None
    s = s or session()
    try:
        tops, invs, lines, logs = _all(s)
        kg = _awb_kg(s)
        res = compute(0, tops, invs, lines, logs, kg)
        tdate = {t.id: t.date for t in tops}
        by_inv = defaultdict(list)
        for l in lines:
            by_inv[l.invoice_id].append(l)
        moving = [i for i in invs if not i.arrived_at]
        ft_rate = None
        out = []
        for i in sorted(moving, key=lambda x: (norm_awb(x.awb), x.id)):
            ls = by_inv[i.id]
            st = sum(l.stems for l in ls) or 1
            country = _country_of(s, i)
            has_air = any(res.lines[l.id].air_rub_stem for l in ls)
            has_msk = any(res.lines[l.id].msk_rub_stem for l in ls)
            notes, approx_legs = [], False
            est = {l.id: [res.lines[l.id].air_rub_stem, res.lines[l.id].msk_rub_stem] for l in ls}
            k = norm_awb(i.awb)
            src = {"air": ("счёт ТК" + (" (≈ по последнему курсу)" if (k, "air") in res.estimated_legs else "")) if has_air else "нет данных",
                   "msk": ("счёт ТК" + (" (предварительный курс)" if (k, "msk") in res.estimated_legs else "")) if has_msk else "нет данных"}
            if country == "Кения":
                for leg, has, label in (("air", has_air, "Expolanka"), ("msk", has_msk, "Floratrack")):
                    if has:
                        continue
                    got = 0
                    for l in ls:
                        v, p = _prev_leg(invs, lines, res, i, l, leg)
                        if v:
                            est[l.id][0 if leg == "air" else 1] = v
                            got += 1
                    notes.append(f"{label}: копия прошлой поставки ({got}/{len(ls)} поз.)" if got
                                 else f"{label}: нет прошлой поставки этих позиций")
                    src[leg] = (f"копия прошлой поставки ({got}/{len(ls)} поз.)" if got
                                else "нет прошлой поставки этих позиций")
                    approx_legs = True
            elif country in FT_TARIFF and not has_msk:
                if ft_rate is None:
                    ft_rate = floratrack_rate()
                w = i.weight_kg
                if w and ft_rate[0]:
                    per = FT_TARIFF[country] * w * ft_rate[0] / st
                    for l in ls:
                        est[l.id][1] = per
                    notes.append(f"Floratrack: тариф {FT_TARIFF[country]} $/кг × {w:g} кг × {ft_rate[0]:.2f}")
                    src["msk"] = f"тариф {FT_TARIFF[country]} $/кг × {w:g} кг × {ft_rate[0]:.2f}"
                else:
                    src["msk"] = "нет кг фермы — пришли разбивку" if not w else "нет курса ЦБ"
                    notes.append("Floratrack: нет кг фермы — пришли разбивку или впиши кг" if not w else "Floratrack: нет курса ЦБ")
                approx_legs = True
            legs_est = any((norm_awb(i.awb), leg) in res.estimated_legs for leg in ("air", "msk"))
            if (has_air or has_msk) and not notes:
                notes.append("по счетам" + (" (часть приблизительно)" if legs_est else ""))
            rows = []
            flower_rub = logi_rub = air_rub = msk_rub = 0.0
            for l in ls:
                c = res.lines[l.id]
                lg = sum(est[l.id])
                air_rub += est[l.id][0] * l.stems
                msk_rub += est[l.id][1] * l.stems
                rows.append({**l.model_dump(), "price_rub": round(c.price_rub, 2), "logi_rub": round(lg, 2),
                             "air_rub": round(est[l.id][0], 2), "msk_rub": round(est[l.id][1], 2),
                             "total_rub": round(c.price_rub + lg, 2)})
                flower_rub += c.price_rub * l.stems
                logi_rub += lg * l.stems
            approx = (not i.topup_id) or legs_est or approx_legs
            out.append({**i.model_dump(), "lines": rows, "stems": st, "country": country,
                        "paid_state": "paid" if (i.topup_id or getattr(i, "via_broker", False)) else "unpaid",
                        "topup": tdate.get(i.topup_id) or ("брокер" if getattr(i, "via_broker", False) else None),
                        "rub_goods": round(res.invoice_rub.get(i.id) or flower_rub), "rub_logi": round(logi_rub),
                        "avg_rub": round((flower_rub + logi_rub) / st, 2), "approx": approx,
                        "logi_source": ("; ".join(notes) or "нет данных").replace("Expolanka", "ТК Кения").replace("Floratrack", "ТК МСК"),
                        "legs": {"air": {"name": "ТК Кения", "rub": round(air_rub), "per": round(air_rub / st, 4), "src": src["air"]},
                                 "msk": {"name": "ТК МСК", "rub": round(msk_rub), "per": round(msk_rub / st, 4), "src": src["msk"]}},
                        "rub_paid": res.invoice_rub.get(i.id)})
        unpaid_logs = [{**lg.model_dump(), "rub_est": round((lg.usd or 0) * res.est_rate)}
                       for lg in logs if not lg.paid]
        return {"invoices": out, "unpaid_logistics": unpaid_logs, "est_rate": res.est_rate,
                "est_topup": max(tops, key=lambda t: t.id).date if tops else None}
    finally:
        if own:
            s.close()


@router.post("/packing/push")
async def api_push_packing(uid: int = Depends(writer)):
    from .bot import push_packing_all, push_text
    r = await push_packing_all()
    return {"ok": not r.get("error"), "text": push_text(r)}


@router.post("/transit/export")
async def export_transit(uid: int = Depends(user_id)):
    """Excel of all goods in transit: sheets «Не оплачены» / «Оплачены», a row per item, status & cost."""
    from aiogram.types import BufferedInputFile
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side
    from .roles import can_write
    full = can_write(role_of_(uid))                  # 1С operator: no paid sums
    v = transit_view()
    F_, thin = "Arial", Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = ["MAWB", "Ферма", "Маркировка", "Номенклатура", "Стебли", "Цена $", "Цветок ₽/ст", "ТК Кения ₽/ст",
            "ТК МСК ₽/ст", "Итого ₽/ст", "Сумма ₽", "Статус оплаты", "Пополнение", "Оплачено $", "Машина", "Прибытие", "Точность"]
    if not full:
        head = [h for h in head if h not in ("Пополнение", "Оплачено $")]
    wb = Workbook()
    wb.remove(wb.active)
    for title, want in (("Не оплачены", "unpaid"), ("Оплачены", "paid")):
        ws = wb.create_sheet(title)
        for c, h in enumerate(head, 1):
            x = ws.cell(1, c, h); x.font = Font(name=F_, bold=True); x.border = box
            x.fill = PatternFill("solid", fgColor="FCE4D6" if want == "unpaid" else "E2EFDA")
        r = 2
        tot_st = tot_rub = 0
        for i in [x for x in v["invoices"] if x["paid_state"] == want]:
            eta = (i.get("eta") or "").replace("T", " ")
            for l in i["lines"]:
                row = {"MAWB": i.get("awb") or "—", "Ферма": i["farm"], "Маркировка": i.get("client_code") or "",
                       "Номенклатура": l["name"], "Стебли": l["stems"], "Цена $": l["price_usd"],
                       "Цветок ₽/ст": round(l["price_rub"], 4), "ТК Кения ₽/ст": round(l.get("air_rub") or 0, 4),
                       "ТК МСК ₽/ст": round(l.get("msk_rub") or 0, 4), "Итого ₽/ст": round(l["total_rub"], 4),
                       "Сумма ₽": round(l["total_rub"] * l["stems"]),
                       "Статус оплаты": "оплачен" if want == "paid" else "НЕ оплачен",
                       "Пополнение": i.get("topup") or "", "Оплачено $": i.get("usd_paid") if want == "paid" else i.get("est_usd"),
                       "Машина": i.get("truck") or "", "Прибытие": eta,
                       "Точность": "≈ ±10%" if i.get("approx") else "точно"}
                for c, h in enumerate(head, 1):
                    x = ws.cell(r, c, row[h]); x.font = Font(name=F_); x.border = box
                    if h.endswith("₽/ст"):
                        x.number_format = "#,##0.0000"
                    elif h == "Сумма ₽":
                        x.number_format = "#,##0"
                tot_st += l["stems"]; tot_rub += row["Сумма ₽"]
                r += 1
        ws.cell(r, 1, "ИТОГО").font = Font(name=F_, bold=True)
        ws.cell(r, head.index("Стебли") + 1, tot_st).font = Font(name=F_, bold=True)
        ws.cell(r, head.index("Сумма ₽") + 1, tot_rub).font = Font(name=F_, bold=True)
        ws.cell(r, head.index("Сумма ₽") + 1).number_format = "#,##0"
        for c, w in enumerate([15, 18, 14, 34, 9, 8, 12, 13, 12, 12, 12, 14, 12, 11, 16, 16, 10][:len(head)], 1):
            ws.column_dimensions[chr(64 + c)].width = w
        ws.freeze_panes = "A2"
    buf = io.BytesIO(); wb.save(buf)
    n_un = sum(1 for x in v["invoices"] if x["paid_state"] != "paid")
    if BOT:
        await BOT.send_document(uid, BufferedInputFile(buf.getvalue(), "Грузы_в_пути.xlsx"),
                                caption=f"🚚 Грузы в пути: не оплачено {n_un}, оплачено {len(v['invoices']) - n_un} инв.")
    return {"ok": True}


@router.get("/transit")
def get_transit(uid: int = Depends(user_id)):
    return transit_view()


class PayIn(BaseModel):
    topup_id: int
    usd: float | None = None
    rub: float | None = None
    farm_usd: float | None = None      # reached the farm (None = exactly the invoice)


def pay_invoice(inv_id: int, topup_id: int, usd: float, rub: float | None = None, farm_usd: float | None = None) -> dict:
    from datetime import date
    with session() as s:
        inv, t = s.get(Invoice, inv_id), s.get(TopUp, topup_id)
        if not inv or not t:
            raise ValueError("нет такого инвойса или пополнения")
        inv.topup_id, inv.usd_paid, inv.paid = topup_id, usd, True
        inv.farm_usd = farm_usd
        inv.rub_paid_override = rub if rub else round(usd * rate_of(t))
        inv.paid_date = date.today().strftime("%d.%m.%Y")
        s.add(inv); s.commit()
        out = {"farm": inv.farm, "usd": usd, "rub": inv.rub_paid_override, "topup": t.date, "auto": not rub}
    from .backup import mark_dirty
    mark_dirty()
    return out


def pay_logistics(log_id: int, topup_id: int, usd: float | None = None, rub: float | None = None) -> dict:
    from datetime import date
    with session() as s:
        lg, t = s.get(Logistics, log_id), s.get(TopUp, topup_id)
        if not lg or not t:
            raise ValueError("нет такой логистики или пополнения")
        lg.topup_id, lg.paid = topup_id, True
        lg.usd = usd or lg.usd
        lg.rub = rub if rub else None          # None -> $ × rate of the top-up in the calc
        lg.paid_date = date.today().strftime("%d.%m.%Y")
        s.add(lg); s.commit()
        out = {"provider": lg.provider, "usd": lg.usd, "rub": rub or round((lg.usd or 0) * rate_of(t)), "topup": t.date}
    from .backup import mark_dirty
    mark_dirty()
    return out


@router.post("/invoices/{inv_id}/pay")
def api_pay_invoice(inv_id: int, body: PayIn, uid: int = Depends(writer)):
    try:
        pay_invoice(inv_id, body.topup_id, body.usd, body.rub, body.farm_usd)
    except ValueError as e:
        raise HTTPException(400, str(e))
    return transit_view()


@router.post("/logistics/{log_id}/pay")
def api_pay_logistics(log_id: int, body: PayIn, uid: int = Depends(writer)):
    try:
        pay_logistics(log_id, body.topup_id, body.usd, body.rub)
    except ValueError as e:
        raise HTTPException(400, str(e))
    return transit_view()


@router.post("/invoices/{inv_id}/arrived")
def api_arrived(inv_id: int, back: bool = False, uid: int = Depends(writer)):
    """Manual: mark arrived (or back to 'в пути')."""
    from datetime import datetime
    with session() as s:
        inv = s.get(Invoice, inv_id)
        inv.arrived_at = None if back else datetime.now().strftime("%d.%m.%Y %H:%M") + " (вручную)"
        s.add(inv); s.commit()
    from .backup import mark_dirty
    mark_dirty()
    return transit_view()


def set_eta(awbs: list[str], eta_iso: str) -> tuple[list, list]:
    """Floratrack chat: these MAWBs arrive at eta -> set eta on their invoices still in transit."""
    matched, unknown = [], []
    with session() as s:
        invs = s.exec(select(Invoice)).all()
        for a in awbs:
            k = norm_awb(a)
            hit = [i for i in invs if norm_awb(i.awb) == k and not i.arrived_at]
            if not hit:
                unknown.append(a)
                continue
            for i in hit:
                i.eta = eta_iso
                s.add(i)
            matched.append((a, [i.farm for i in hit]))
        s.commit()
    from .backup import mark_dirty
    mark_dirty()
    return matched, unknown


def set_client_eta(awbs: list[str], iso: str, done: bool = False):
    """When clients are told «прибыл» (TK time + 6 h). done=True for old history: never notify clients."""
    keys = {norm_awb(a) for a in awbs}
    with session() as s:
        for i in s.exec(select(Invoice)).all():
            if norm_awb(i.awb) in keys and not i.client_done:
                i.client_eta = iso
                i.client_done = done
                s.add(i)
        s.commit()


def client_arrivals_due(now_iso: str) -> list:
    """Invoices whose client ETA has passed and clients weren't told yet -> mark and return them."""
    out = []
    with session() as s:
        for i in s.exec(select(Invoice)).all():
            if i.client_eta and not i.client_done and i.client_eta <= now_iso:
                i.client_done = True
                s.add(i)
                out.append({"awb": i.awb})
        s.commit()
    return out


def assign_awb_by_farms(awb: str, farms: list[str]) -> list[str]:
    """Consolidation list without weights: in-transit invoices of these farms that have no MAWB get this one."""
    from .ai import find_mawb
    awb = find_mawb(awb) or awb
    done = []
    with session() as s:
        for i in sorted(s.exec(select(Invoice)).all(), key=lambda x: -x.id):
            if not (i.awb or "").strip() and not i.arrived_at and i.farm in farms and i.farm not in done:
                i.awb = awb; s.add(i); done.append(i.farm)
        s.commit()
    from .backup import mark_dirty
    mark_dirty()
    return done


def set_truck(awb: str, truck: str) -> list[str]:
    """MAWB loaded into a truck -> remember it on the invoices still in transit. Returns farms."""
    k, farms = norm_awb(awb), []
    with session() as s:
        for i in s.exec(select(Invoice)).all():
            if norm_awb(i.awb) == k and not i.arrived_at:
                i.truck = truck
                s.add(i)
                farms.append(i.farm)
        s.commit()
    return farms


def arrive_due(now_iso: str) -> list[dict]:
    """Scheduler: invoices whose eta has passed become arrived."""
    done = []
    with session() as s:
        for i in s.exec(select(Invoice)).all():
            if not i.arrived_at and i.eta and i.eta <= now_iso:
                i.arrived_at = i.eta.replace("T", " ")[:16] + " (Floratrack)"
                s.add(i)
                done.append({"farm": i.farm, "awb": i.awb, "paid": bool(i.topup_id)})
        s.commit()
    if done:
        from .backup import mark_dirty
        mark_dirty()
    return done


def unpaid_items() -> dict:
    """For the Tuesday/Wednesday payment reminder."""
    v = transit_view()
    invs = [x for x in v["invoices"] if x["paid_state"] == "unpaid"]
    return {"invoices": invs, "logistics": v["unpaid_logistics"], "rate": v["est_rate"], "topup": v["est_topup"]}


# ---------- drafts: invoices sent straight into the bot chat, waiting for the operator ----
DRAFTS = DATA_DIR / "drafts"
DRAFTS.mkdir(exist_ok=True)


def fill_mawb(out: dict, topup_id: int | None = None):
    """Parsed farm invoice without MAWB -> take it from the forwarder breakdown that lists this farm."""
    if out.get("doc_type") != "farm_invoice" or out.get("awb"):
        return
    awb, kg, n = infer_mawb(out.get("farm") or "", topup_id)
    if not awb:
        out.setdefault("warnings", []).append(
            "MAWB в инвойсе нет, и разбивки с этой плантацией пока нет — кинь разбивку Expolanka, MAWB подставится сам")
        return
    out["awb"], out["weight_kg"] = awb, kg
    out["warnings"] = [w for w in out.get("warnings", []) if "MAWB" not in w and "awb" not in w.lower()]
    out["mawb_note"] = f"MAWB взят из разбивки: {out.get('farm')} {kg:g} кг" + (
        f" (подходящих MAWB {n}, взят последний — проверь)" if n > 1 else "")


def _draft_key(d: dict):
    return (d.get("doc_type"), norm_awb(d.get("awb") or ""), (d.get("farm") or "").lower(), d.get("invoice_no") or "")


def save_draft(parsed: dict) -> str:
    """Same document sent twice -> replaces the old draft instead of piling up."""
    key = _draft_key(parsed)
    if key[1] or key[3]:
        for p in DRAFTS.glob("*.json"):
            try:
                if _draft_key(json.loads(p.read_text())) == key:
                    p.unlink()
            except ValueError:
                pass
    did = uuid.uuid4().hex[:8]
    (DRAFTS / f"{did}.json").write_text(json.dumps(parsed, ensure_ascii=False))
    from .backup import mark_dirty
    mark_dirty()                      # drafts from the chat don't go through the API middleware
    return did


@router.get("/drafts")
def list_drafts(uid: int = Depends(user_id)):
    from .ai import find_mawb
    out, seen, changed = [], set(), False
    for p in sorted(DRAFTS.glob("*.json"), key=lambda p: p.stat().st_mtime, reverse=True):
        d = json.loads(p.read_text())
        clean = find_mawb(d.get("awb") or "")               # old drafts: strip "/ HAWB ..."
        if clean and clean != d.get("awb"):
            d["awb"] = clean
            p.write_text(json.dumps(d, ensure_ascii=False))
        k = _draft_key(d)
        if (k[1] or k[3]) and k in seen:                    # duplicate of a newer draft
            p.unlink(); changed = True
            continue
        seen.add(k)
        if d.get("doc_type") == "kg_breakdown" and d.get("awb"):
            store_breakdown(d["awb"], d.get("per_farm_kg"), d.get("source_file"))   # belongs to the MAWB, not a draft
            p.unlink(); changed = True
            continue
        out.append({"id": p.stem, **d})
    kg = {}
    with session() as s:
        kg = _awb_kg(s)
    for d in out:
        d["kg_total"] = kg.get(norm_awb(d.get("awb") or ""), 0)
    if changed:
        from .backup import mark_dirty
        mark_dirty()
    return out


@router.delete("/drafts/{did}")
def drop_draft(did: str, uid: int = Depends(writer)):
    (DRAFTS / f"{did}.json").unlink(missing_ok=True)
    return {"ok": True}
